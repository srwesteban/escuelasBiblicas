"""
Lectura y validación de archivos de seguimiento (p. ej. export de formularios / pagos).
Mantiene la lógica fuera de las vistas para emparejar ofertas con Escuela operativa y persistencia.
"""
from __future__ import annotations

import os
import re
import secrets
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from typing import TYPE_CHECKING, Any

from django.db import transaction

if TYPE_CHECKING:
    from core.models import Sede

DEFAULT_PREVIEW_MAX_ROWS = int(os.getenv('SEGUIMIENTO_IMPORT_PREVIEW_ROWS', '50'))
MAX_UPLOAD_BYTES = int(os.getenv('SEGUIMIENTO_IMPORT_MAX_UPLOAD_BYTES', str(15 * 1024 * 1024)))
ALLOWED_EXTENSIONS = ('.xlsx',)

# Documento: sin alias de 2 letras («ti» matcheaba dentro de «Efectivo»).
# Evitar «documento» suelto (aparece en «tipo de documento», etc.).
_HEADER_ALIASES_DOCUMENTO = (
    'numero de identificacion',
    'número de identificación',
    'numero de documento',
    'número de documento',
    'documento de identidad',
    'doc. identidad',
    'identificacion',
    'identificación',
    'cedula de ciudadania',
    'cédula de ciudadanía',
    'cedula',
    'cédula',
    'cc o ti',
    'nit',
)
_HEADER_ALIASES_NOMBRE = (
    'nombre del estudiante',
    'nombre estudiante',
    'nombre completo',
    'nombres y apellidos',
    'nombre',
)
# Nombre de la escuela operativa (modelo Escuela en BD), no la línea larga de curso/docente.
_HEADER_ALIASES_ESCUELA_INSTITUCION = (
    'nombre de la escuela',
    'escuela a la que asiste',
    'escuela asignada',
    'escuela elegida',
    'unidad academica',
    'unidad académica',
    'institucion',
    'institución',
)
# Texto de oferta: curso, docente, modalidad (lo que se empareja con la edición).
_HEADER_ALIASES_OFERTA_CURSO = (
    'curso ofertado',
    'nombre del curso',
    'opcion de curso',
    'opción de curso',
    'programa elegido',
    'programa',
    'eleccion',
    'elección',
    'texto de la oferta',
    'oferta',
    'linea elegida',
    'línea elegida',
    'modalidad y curso',
)
# Si solo hay una columna «escuela» con el texto largo, se usa como oferta (institución vacía).
_HEADER_ALIASES_ESCUELA_UNICA_LEGACY = ('escuela',)


class SeguimientoImportError(Exception):
    """Error de negocio al interpretar el libro (mensaje listo para mostrar al usuario)."""


@dataclass(frozen=True)
class SeguimientoPreview:
    """Vista previa no destructiva de la primera hoja."""

    headers: list[str]
    rows: list[list[str]]
    sheet_name: str
    has_more_rows: bool
    preview_row_limit: int


@dataclass(frozen=True)
class EdicionMatchDetail:
    """Match de oferta; ``edicion_id`` es el id de ``Escuela`` (nombre histórico)."""

    edicion_id: int
    label: str
    score: float


@dataclass
class RowMatchAnalysis:
    """Análisis de una fila de datos frente a BD (sin persistir)."""

    row_number: int
    doc_excel: str
    nombre_excel: str
    escuela_texto_excel: str
    edicion: EdicionMatchDetail | None
    edicion_alternativas: tuple[EdicionMatchDetail, ...] = field(default_factory=tuple)
    estudiante_id: int | None = None
    estudiante_nombre: str | None = None
    status: str = 'error'  # ok | warning | error
    notes: tuple[str, ...] = field(default_factory=tuple)
    escuela_institucion_excel: str = ''


@dataclass
class SeguimientoMatchReport:
    """Resultado del cruce Excel ↔ ediciones y estudiantes en BD."""

    column_documento: int | None
    column_nombre: int | None
    column_escuela_texto: int | None
    column_escuela_institucional: int | None
    file_notes: tuple[str, ...]
    ediciones_en_catalogo: int
    rows: list[RowMatchAnalysis] = field(default_factory=list)

    @property
    def status_counts(self) -> dict[str, int]:
        out: dict[str, int] = {'ok': 0, 'warning': 0, 'error': 0}
        for r in self.rows:
            out[r.status] = out.get(r.status, 0) + 1
        out['total'] = len(self.rows)
        return out


@dataclass
class ImportRowOutcome:
    """Resultado de aplicar una fila del Excel en base de datos."""

    row_number: int
    doc_excel: str
    nombre_excel: str
    status: str  # created | matricula_ok | skipped | error
    message: str
    username: str = ''
    password_plain: str = ''
    estudiante_id: int | None = None
    edicion_label: str = ''


@dataclass
class SeguimientoImportExecuteReport:
    rows: list[ImportRowOutcome]

    def status_counts(self) -> dict[str, int]:
        out: dict[str, int] = {'created': 0, 'matricula_ok': 0, 'skipped': 0, 'error': 0}
        for r in self.rows:
            out[r.status] = out.get(r.status, 0) + 1
        out['total'] = len(self.rows)
        return out


def _fold_text(value: str) -> str:
    s = unicodedata.normalize('NFKD', (value or '').strip().lower())
    return ''.join(c for c in s if not unicodedata.combining(c))


def _needle_is_whole_token(haystack: str, needle: str) -> bool:
    """Evita que «ti» coincida dentro de «efectivo»."""
    for tok in re.findall(r'[\w]+', haystack, flags=re.UNICODE):
        if tok == needle:
            return True
    return False


def _alias_matches_header(hf: str, aln: str) -> bool:
    if not hf or not aln:
        return False
    if len(aln) <= 3:
        if hf == aln:
            return True
        return _needle_is_whole_token(hf, aln)
    return aln in hf or hf in aln


def _resolve_header_column(
    headers: list[str],
    aliases: tuple[str, ...],
    *,
    skip_indices: set[int] | frozenset[int] | None = None,
) -> int | None:
    """Primera columna cuyo encabezado encaje; alias largos primero; cortos solo como palabra completa."""
    skip = skip_indices or set()
    aliases_sorted = sorted(aliases, key=len, reverse=True)
    for i, h in enumerate(headers):
        if i in skip:
            continue
        hf = _fold_text(h)
        if not hf:
            continue
        for al in aliases_sorted:
            aln = _fold_text(al)
            if _alias_matches_header(hf, aln):
                return i
    return None


def _document_cell_likelihood(cell: str) -> float:
    """Heurística cédula/documento CO (solo dígitos, longitud típica)."""
    raw = ''.join((cell or '').strip().split())
    if not raw:
        return 0.0
    digits = ''.join(c for c in raw if c.isdigit())
    if len(digits) < 5:
        return 0.0
    if raw.replace('.', '').replace('-', '').isdigit() or digits == raw:
        if 6 <= len(digits) <= 12:
            return 1.0
        if 5 <= len(digits) <= 15:
            return 0.5
    return 0.0


def _resolve_documento_column(headers: list[str], rows: list[list[str]]) -> int | None:
    """
    Columna de documento: combina encabezado reconocible con la columna cuyas celdas parezcan
    más cédulas/IDs en la muestra. Si el encabezado apunta a una columna sin números pero otra
    columna sí tiene IDs claros, se prioriza la evidencia de los datos.
    """
    header_idx = _resolve_header_column(headers, _HEADER_ALIASES_DOCUMENTO)
    ncols = len(headers)
    if ncols == 0:
        return None
    if not rows:
        return header_idx

    scores = [
        sum(_document_cell_likelihood(row[i] if i < len(row) else '') for row in rows)
        for i in range(ncols)
    ]
    best_i = max(range(ncols), key=lambda i: scores[i])
    best_score = scores[best_i]
    # Al menos ~1 de cada 3 filas con documento plausible, o suma acumulada clara
    data_confident = best_score >= max(1.0, 0.35 * len(rows))

    if data_confident:
        if header_idx is None:
            return best_i
        if scores[header_idx] < best_score * 0.5:
            return best_i
        if scores[header_idx] >= best_score * 0.85:
            return header_idx
        return best_i

    return header_idx


# Términos muy frecuentes en formularios que inflan el unión del Jaccard sin aportar identidad.
_MATCH_STOPWORDS = frozenset(
    {
        'presencial',
        'virtual',
        'modalidad',
        'horario',
        'lunes',
        'martes',
        'miercoles',
        'jueves',
        'viernes',
        'sabado',
        'domingo',
        'am',
        'pm',
    }
)


def _normalize_dashes_punct(s: str) -> str:
    """Unifica guiones unicode y separadores típicos del Excel de ofertas."""
    t = (s or '').replace('–', ' ').replace('—', ' ').replace('−', ' ').replace('·', ' ')
    t = t.replace('‐', ' ')
    return t


def _unificar_variantes_nombre_oferta_folded(t: str) -> str:
    """
    Tras _fold_text: unifica el lenguaje del Excel con variantes del catálogo interno.
    Ej.: «nuevo lead - grupo hechos» (BD) ≡ «nuevo lead hechos» (Excel / encuesta).
    """
    if not t:
        return t
    t = re.sub(r'\s+', ' ', t).strip()
    t = re.sub(
        r'nuevo\s+lead\s*[\-\u2013]?\s*grupo\s+hechos',
        'nuevo lead hechos',
        t,
    )
    return t


def _prepare_texto_oferta_match(s: str) -> str:
    """Texto listo para comparar oferta Excel ↔ catálogo (misma forma canónica)."""
    return _unificar_variantes_nombre_oferta_folded(
        _fold_text(_normalize_dashes_punct(s))
    )


def _tokenize_prepared_oferta(prepared_folded: str, *, min_len: int = 3) -> set[str]:
    s = re.sub(r'[^\w\s]', ' ', prepared_folded, flags=re.UNICODE)
    out: set[str] = set()
    for t in s.split():
        if len(t) < min_len:
            continue
        if t in _MATCH_STOPWORDS:
            continue
        out.add(t)
    return out


def _visible_nombre_curso_alineado_excel(nucleo: str) -> str:
    """
    Ajusta el nombre visible del curso a la forma típica del Excel cuando es la misma oferta
    en distintas redacciones internas.
    """
    if not (nucleo or '').strip():
        return nucleo
    return re.sub(
        r'(?i)nuevo\s+lead\s*[\u2013\-]?\s*grupo\s+hechos',
        'Nuevo Lead Hechos',
        nucleo,
        count=1,
    )


def _oferta_nucleo(excel_text: str) -> str:
    """
    Quita cola tipo «, Presencial (martes 7:00 pm)» para comparar curso + docente como en formularios reales.
    """
    t = _normalize_dashes_punct((excel_text or '').strip())
    low = t.casefold()
    for needle in (', presencial', ', virtual', ' presencial(', ' presencial ', '('):
        idx = low.find(needle.casefold())
        if idx != -1:
            t = t[:idx].strip()
            break
    return t.strip()


def _tokenize_match(s: str, *, min_len: int = 3) -> set[str]:
    s = _fold_text(re.sub(r'[^\w\s]', ' ', _normalize_dashes_punct(s), flags=re.UNICODE))
    out: set[str] = set()
    for t in s.split():
        if len(t) < min_len:
            continue
        if t in _MATCH_STOPWORDS:
            continue
        out.add(t)
    return out


def _score_texto_vs_catalogo_one(excel_text: str, catalogo_blob: str) -> float:
    """Una variante de texto Excel frente al blob de catálogo."""
    a = _prepare_texto_oferta_match(excel_text)
    b = _prepare_texto_oferta_match(catalogo_blob)
    if not a:
        return 0.0
    if not b:
        return 0.0
    if a in b or b in a:
        return 0.92
    ta = _tokenize_prepared_oferta(a)
    tb = _tokenize_prepared_oferta(b)
    if not ta or not tb:
        return 0.0
    inter = len(ta & tb)
    if inter == 0:
        hit_tokens = [tok for tok in ta if len(tok) >= 4 and tok in b]
        hit = len(hit_tokens)
        if hit >= 2:
            return min(0.48, 0.14 + 0.09 * hit)
        if hit == 1 and len(hit_tokens[0]) >= 6:
            return 0.2
        return 0.0
    union = len(ta | tb) or 1
    jacc = inter / union
    cov_a = inter / len(ta)
    cov_b = inter / len(tb)
    blended = 0.32 * jacc + 0.38 * cov_a + 0.30 * cov_b
    if inter >= 3:
        blended = max(blended, 0.82 * cov_a)
    if inter >= 2 and cov_a >= 0.4:
        blended = max(blended, 0.72 * cov_a)
    return float(min(1.0, max(jacc, blended)))


def _score_texto_vs_catalogo(excel_text: str, catalogo_blob: str) -> float:
    """
    Compara texto de oferta del Excel con el blob enriquecido de una edición.
    Usa el núcleo (sin modalidad/horario) y el texto completo; se queda con el mejor puntaje.
    """
    raw = (excel_text or '').strip()
    if not raw:
        return 0.0
    nucleo = _oferta_nucleo(raw)
    scores = [_score_texto_vs_catalogo_one(raw, catalogo_blob)]
    if nucleo and nucleo.casefold() != raw.casefold():
        scores.append(_score_texto_vs_catalogo_one(nucleo, catalogo_blob))
    return max(scores)


def _match_escuela_ids_en_sede(celda: str, sede: Sede) -> list[int]:
    """Escuelas operativas de la sede cuyo nombre hace match con la celda del Excel."""
    from hechos.models import Escuela

    hits: list[int] = []
    folded_cell = _fold_text(_normalize_dashes_punct(celda))
    if not folded_cell:
        return hits
    qs = Escuela.objects.filter(sede=sede, is_active=True)
    for esc in qs:
        en = _fold_text(_normalize_dashes_punct(esc.nombre))
        if not en:
            continue
        if folded_cell == en or en in folded_cell or folded_cell in en:
            hits.append(esc.pk)
            continue
        tt = _tokenize_match(celda)
        et = _tokenize_match(esc.nombre)
        if not et:
            continue
        inter = len(tt & et)
        if inter >= max(1, min(2, len(et))):
            hits.append(esc.pk)
    return hits


def _infer_escuela_ids_desde_oferta(oferta: str, sede: Sede) -> frozenset[int] | None:
    """
    Si el nombre de una sola Escuela (operativa) aparece como subcadena en el texto de oferta,
    restringe a esa escuela. Si hay 0 o varias coincidencias ambiguas, devuelve None.
    """
    from hechos.models import Escuela

    if not (oferta or '').strip():
        return None
    a = _fold_text(_normalize_dashes_punct(oferta))
    hits: list[int] = []
    for esc in Escuela.objects.filter(sede=sede, is_active=True):
        en = _fold_text(_normalize_dashes_punct(esc.nombre))
        if len(en) >= 5 and en in a:
            hits.append(esc.pk)
    if len(hits) == 1:
        return frozenset(hits)
    return None


def _escuela_ids_restriccion_fila(
    sede: Sede,
    celda_institucion: str,
    celda_oferta: str,
) -> tuple[frozenset[int] | None, tuple[str, ...]]:
    """
    None = sin filtro (todas las ediciones de la sede en el catálogo).
    frozenset() = institución informada pero sin match con Escuela activa.
    frozenset({...}) = solo ediciones ligadas a esas escuelas operativas.
    """
    ci = (celda_institucion or '').strip()
    co = (celda_oferta or '').strip()
    if ci:
        matched = _match_escuela_ids_en_sede(ci, sede)
        if not matched:
            return frozenset(), (
                'El nombre de escuela en el Excel no coincide con ninguna escuela operativa activa de esta sede en la base de datos.',
            )
        return frozenset(matched), ()
    inferred = _infer_escuela_ids_desde_oferta(co, sede)
    if inferred is not None:
        return inferred, ()
    return None, ()


def _normalize_documento(raw: str) -> str:
    return ''.join((raw or '').split())


def _find_estudiante_por_documento(sede_id: int, doc_excel: str):
    from hechos.models import Estudiante

    doc = _normalize_documento(doc_excel)
    if len(doc) < 4:
        return None
    qs = Estudiante.objects.filter(sede_id=sede_id, is_active=True).select_related('user')
    est = qs.filter(numero_documento=doc).first()
    if est:
        return est
    est = qs.filter(numero_documento__iexact=doc).first()
    if est:
        return est
    digits = ''.join(c for c in doc if c.isdigit())
    if len(digits) < 5:
        return None
    for e in qs:
        ed = ''.join(c for c in (e.numero_documento or '') if c.isdigit())
        if ed == digits:
            return e
    return None


def _find_estudiante_en_sede_por_documento(sede_id: int, doc_excel: str):
    """Igual que _find_estudiante_por_documento pero incluye inactivos (reactivación en import)."""
    from hechos.models import Estudiante

    doc = _normalize_documento(doc_excel)
    if len(doc) < 4:
        return None
    qs = Estudiante.objects.filter(sede_id=sede_id).select_related('user')
    est = qs.filter(numero_documento=doc).first()
    if est:
        return est
    est = qs.filter(numero_documento__iexact=doc).first()
    if est:
        return est
    digits = ''.join(c for c in doc if c.isdigit())
    if len(digits) < 5:
        return None
    for e in qs:
        ed = ''.join(c for c in (e.numero_documento or '') if c.isdigit())
        if ed == digits:
            return e
    return None


def _normalize_doc_username(raw: str) -> str | None:
    """
    Valor de username / documento_identidad coherente con registro de estudiantes (cédula/TI numérica).
    """
    s = _normalize_documento(raw)
    if not s:
        return None
    digits = re.sub(r'[\s.\-]', '', s)
    if digits.isdigit() and 6 <= len(digits) <= 10:
        if len(set(digits)) == 1:
            return None
        return digits
    return None


def _split_nombre_import(nombre_excel: str) -> tuple[str, str]:
    s = (nombre_excel or '').strip()
    if not s:
        return 'Estudiante', 'Importado'
    parts = s.split(None, 1)
    if len(parts) == 1:
        return parts[0], 'Importado'
    return (parts[0][:150], parts[1][:150])


def _periodo_desde_escuela(escuela) -> str:
    d = getattr(escuela, 'fecha_inicio', None)
    if d:
        return f'{d.year}-1'
    return f'{date.today().year}-1'


def _grant_hechos_permission(user) -> None:
    from core.models import AppModule, UserAppPermission

    hechos_module = AppModule.objects.filter(name='hechos').first()
    if hechos_module:
        UserAppPermission.objects.get_or_create(
            user=user,
            app_module=hechos_module,
            defaults={
                'can_view': True,
                'can_edit': False,
                'can_delete': False,
                'can_manage': False,
            },
        )


def _horario_desde_parentesis_oferta(texto: str) -> str:
    m = re.search(r'\(([^)]{3,100})\)', texto or '')
    if m:
        return m.group(1).strip()[:100]
    return 'Por definir'


def _ancla_ruta_estudio_provision(sede: Sede):
    """Primera escuela operativa de la sede + primera ruta activa (ancla estable para nuevas ofertas)."""
    from hechos.models import Escuela, RutaEstudio

    esc = Escuela.objects.filter(sede=sede, is_active=True).order_by('id').first()
    if not esc:
        return None
    return RutaEstudio.objects.filter(escuela=esc, is_active=True).order_by('id').first()


def provisionar_curso_edicion_desde_texto_excel(sede: Sede, texto_oferta: str):
    """
    Alinea la BD con el texto del Excel: crea o reutiliza Curso (nombre = núcleo de la oferta)
    bajo la primera ruta activa de la sede y actualiza horario en la escuela operativa.
    """
    from django.db.models import Max

    from hechos.models import Curso, Escuela

    nucleo = (_oferta_nucleo(texto_oferta) or '').strip() or (texto_oferta or '').strip()
    nucleo = _visible_nombre_curso_alineado_excel(nucleo)
    nucleo = nucleo[:200]
    if not nucleo:
        return None
    ruta = _ancla_ruta_estudio_provision(sede)
    if not ruta or not ruta.escuela_id:
        return None
    escuela: Escuela = ruta.escuela
    max_o = Curso.objects.filter(ruta_estudio=ruta).aggregate(n=Max('orden'))['n'] or 0
    curso, created = Curso.objects.get_or_create(
        ruta_estudio=ruta,
        nombre=nucleo,
        defaults={
            'sede': sede,
            'orden': max_o + 1,
            'duracion_semanas': 4,
            'descripcion': '',
            'is_active': True,
        },
    )
    if not created and curso.sede_id != sede.pk:
        curso.sede = sede
        curso.save(update_fields=['sede', 'updated_at'])
    horario = _horario_desde_parentesis_oferta(texto_oferta)
    today = date.today()
    fin = date(today.year + 1, 6, 30)
    if escuela.fecha_inicio is None:
        escuela.fecha_inicio = today
    if escuela.fecha_fin is None:
        escuela.fecha_fin = fin
    if horario != 'Por definir':
        if not escuela.horario or escuela.horario == 'Por definir':
            escuela.horario = horario
    escuela.save(
        update_fields=['fecha_inicio', 'fecha_fin', 'horario', 'updated_at'],
    )
    return escuela


def execute_seguimiento_import(
    preview: SeguimientoPreview,
    sede: Sede,
    *,
    match_report: SeguimientoMatchReport | None = None,
    crear_ediciones_faltantes: bool = False,
) -> SeguimientoImportExecuteReport:
    """
    Crea usuarios (usuario = documento, contraseña aleatoria), perfiles Estudiante, permiso Hechos
    y matrícula en la escuela operativa detectada por fila. Idempotente en matrícula (get_or_create).

    Si ``crear_ediciones_faltantes`` y no hubo match, intenta crear curso / actualizar escuela desde el texto
    de oferta del Excel (ancla: primera escuela operativa y primera ruta de la sede).
    """
    from django.contrib.auth import get_user_model

    from hechos.models import Escuela, Estudiante, Matricula

    User = get_user_model()
    if match_report is None:
        match_report = analyze_seguimiento_preview_against_bd(preview, sede)

    outcomes: list[ImportRowOutcome] = []

    for analysis in match_report.rows:
        base = dict(
            row_number=analysis.row_number,
            doc_excel=analysis.doc_excel,
            nombre_excel=analysis.nombre_excel,
            username='',
            password_plain='',
            estudiante_id=None,
        )

        escuela_oferta: Escuela | None = None
        ed_label = ''
        provision_note = ''

        if analysis.edicion is not None:
            try:
                escuela_oferta = Escuela.objects.select_related('sede').get(
                    pk=analysis.edicion.edicion_id,
                    sede_id=sede.pk,
                    is_active=True,
                )
            except Escuela.DoesNotExist:
                outcomes.append(
                    ImportRowOutcome(
                        **base,
                        status='error',
                        message='La escuela ya no está disponible o no pertenece a esta sede.',
                        edicion_label=analysis.edicion.label,
                    )
                )
                continue
            ed_label = _label_escuela_para_ui(escuela_oferta)
        elif crear_ediciones_faltantes and (analysis.escuela_texto_excel or '').strip():
            try:
                escuela_oferta = provisionar_curso_edicion_desde_texto_excel(sede, analysis.escuela_texto_excel)
            except Exception as exc:
                outcomes.append(
                    ImportRowOutcome(
                        **base,
                        status='error',
                        message=f'No se pudo crear la oferta desde el Excel: {exc}',
                        edicion_label='',
                    )
                )
                continue
            if escuela_oferta is None:
                outcomes.append(
                    ImportRowOutcome(
                        **base,
                        status='skipped',
                        message='Sin escuela emparejada y no hay escuela/ruta activa en la sede para crear la oferta.',
                        edicion_label='',
                    )
                )
                continue
            ed_label = _label_escuela_para_ui(escuela_oferta)
            provision_note = (
                'Oferta creada o reutilizada desde el texto del Excel (primera escuela y ruta de la sede). '
            )
        else:
            outcomes.append(
                ImportRowOutcome(
                    **base,
                    status='skipped',
                    message='Sin escuela emparejada; no se importó la fila. Marca «Crear ofertas faltantes» si el Excel debe mandar en la BD.',
                    edicion_label='',
                )
            )
            continue

        doc_norm = _normalize_doc_username(analysis.doc_excel)
        if not doc_norm:
            outcomes.append(
                ImportRowOutcome(
                    **base,
                    status='error',
                    message='Documento vacío o no válido para crear usuario (use 6–10 dígitos, cédula o TI).',
                    edicion_label=ed_label,
                )
            )
            continue

        first_name, last_name = _split_nombre_import(analysis.nombre_excel)

        existing_user = User.objects.filter(username=doc_norm).first()
        if existing_user is None:
            existing_user = User.objects.filter(documento_identidad__iexact=doc_norm).first()
        if existing_user is not None and hasattr(existing_user, 'estudiante_profile'):
            other = existing_user.estudiante_profile
            if other.sede_id is not None and other.sede_id != sede.pk:
                outcomes.append(
                    ImportRowOutcome(
                        **base,
                        status='error',
                        edicion_label=ed_label,
                        message='Ese documento ya pertenece a un estudiante de otra sede.',
                    )
                )
                continue

        password_plain = ''
        msg_tail = ''
        mat_created = False

        try:
            with transaction.atomic():
                est = _find_estudiante_en_sede_por_documento(sede.pk, analysis.doc_excel)

                if est:
                    if not est.is_active:
                        est.is_active = True
                        est.save(update_fields=['is_active'])
                    user = est.user
                    if user.sede_id != sede.pk:
                        user.sede = sede
                        user.save(update_fields=['sede'])
                    if (analysis.nombre_excel or '').strip():
                        user.first_name = first_name
                        user.last_name = last_name
                        user.save(update_fields=['first_name', 'last_name'])
                    msg_tail = 'Ya estaba en la sede; se aseguró la matrícula en la edición.'
                    _grant_hechos_permission(user)
                else:
                    ex_user = existing_user
                    if ex_user is not None:
                        if hasattr(ex_user, 'estudiante_profile'):
                            est = ex_user.estudiante_profile
                            est.sede_id = sede.pk
                            est.numero_documento = doc_norm
                            est.is_active = True
                            est.save(update_fields=['sede', 'numero_documento', 'is_active'])
                            user = ex_user
                            user.sede = sede
                            if not user.documento_identidad:
                                user.documento_identidad = doc_norm
                            if (analysis.nombre_excel or '').strip():
                                user.first_name = first_name
                                user.last_name = last_name
                                user.save(
                                    update_fields=['sede', 'documento_identidad', 'first_name', 'last_name']
                                )
                            else:
                                user.save(update_fields=['sede', 'documento_identidad'])
                            msg_tail = 'Estudiante existente vinculado a esta sede; matrícula asegurada.'
                        else:
                            user = ex_user
                            user.sede = sede
                            if not user.documento_identidad:
                                user.documento_identidad = doc_norm
                            if (analysis.nombre_excel or '').strip():
                                user.first_name = first_name
                                user.last_name = last_name
                                user.save(
                                    update_fields=['sede', 'documento_identidad', 'first_name', 'last_name']
                                )
                            else:
                                user.save(update_fields=['sede', 'documento_identidad'])
                            est = Estudiante.objects.create(
                                user=user,
                                sede=sede,
                                tipo_documento=Estudiante.TipoDocumento.CC,
                                numero_documento=doc_norm,
                                is_active=True,
                            )
                            password_plain = secrets.token_urlsafe(9)
                            user.set_password(password_plain)
                            user.save(update_fields=['password'])
                            msg_tail = (
                                'Había usuario sin perfil de estudiante; se creó el perfil y una contraseña nueva.'
                            )
                        _grant_hechos_permission(user)
                    else:
                        password_plain = secrets.token_urlsafe(9)
                        user = User(
                            username=doc_norm,
                            first_name=first_name,
                            last_name=last_name,
                            email=None,
                            role='user',
                            sede=sede,
                            tipo_documento=User.TipoDocumento.CC,
                            documento_identidad=doc_norm,
                            is_active=True,
                        )
                        user.set_password(password_plain)
                        user.save()
                        est = Estudiante.objects.create(
                            user=user,
                            sede=sede,
                            tipo_documento=Estudiante.TipoDocumento.CC,
                            numero_documento=doc_norm,
                            is_active=True,
                        )
                        _grant_hechos_permission(user)
                        msg_tail = 'Usuario y estudiante creados; entrega el usuario (documento) y la contraseña mostrada.'

                mat, mat_created = Matricula.objects.get_or_create(
                    estudiante=est,
                    escuela=escuela_oferta,
                    defaults={
                        'sede': sede,
                        'estado': 'activa',
                        'is_active': True,
                        'periodo': _periodo_desde_escuela(escuela_oferta),
                    },
                )
                if not mat_created and not mat.is_active:
                    mat.is_active = True
                    mat.estado = 'activa'
                    mat.save(update_fields=['is_active', 'estado'])

        except Exception as exc:
            outcomes.append(
                ImportRowOutcome(
                    **base,
                    status='error',
                    edicion_label=ed_label,
                    message=f'Error al guardar: {exc}',
                )
            )
            continue

        if provision_note:
            msg_tail = f'{provision_note}{msg_tail}'
        if not mat_created:
            msg_tail = f'{msg_tail} La matrícula en esta escuela ya existía.'
        st = 'created' if password_plain else 'matricula_ok'
        outcomes.append(
            ImportRowOutcome(
                row_number=analysis.row_number,
                doc_excel=analysis.doc_excel,
                nombre_excel=analysis.nombre_excel,
                status=st,
                message=msg_tail,
                username=doc_norm,
                password_plain=password_plain,
                estudiante_id=est.pk,
                edicion_label=ed_label,
            )
        )

    return SeguimientoImportExecuteReport(rows=outcomes)


def _label_escuela_para_ui(esc) -> str:
    from hechos.models import Curso

    nom = (esc.nombre or '').strip() or '—'
    curso = (
        Curso.objects.filter(ruta_estudio__escuela=esc, is_active=True).order_by('orden', 'id').first()
    )
    cx = f' · {curso.nombre}' if curso else ''
    return f'{nom}{cx} (escuela id {esc.pk})'


def _ediciones_catalogo_para_sede(sede: Sede) -> list[tuple[Any, str, float]]:
    """
    Escuelas operativas activas de la sede para hacer match fila a fila con el texto de oferta del Excel.
    """
    from hechos.models import Curso, Escuela

    escuelas = (
        Escuela.objects.filter(sede=sede, is_active=True)
        .select_related('maestro__user')
        .order_by('nombre', 'id')
    )
    out: list[tuple[Any, str, float]] = []
    for escuela_op in escuelas:
        curso = (
            Curso.objects.filter(ruta_estudio__escuela=escuela_op, is_active=True)
            .select_related('ruta_estudio')
            .order_by('orden', 'id')
            .first()
        )
        ruta_nombre = ''
        if curso and curso.ruta_estudio_id:
            ruta_nombre = (curso.ruta_estudio.nombre or '').strip()
        partes: list[str] = [
            escuela_op.nombre or '',
            ruta_nombre,
            escuela_op.linea_identificacion_pedagogica(),
            (curso.nombre or '').strip() if curso else '',
            escuela_op.horario or '',
            escuela_op.aula or '',
        ]
        if escuela_op.maestro and escuela_op.maestro.user_id:
            partes.append(escuela_op.maestro.user.get_full_name() or '')
            if escuela_op.maestro.user.first_name:
                partes.append(escuela_op.maestro.user.first_name)
            if escuela_op.maestro.user.last_name:
                partes.append(escuela_op.maestro.user.last_name)
        blob = ' '.join(dict.fromkeys(p for p in partes if p))
        out.append((escuela_op, blob, escuela_op.pk))
    return out


def analyze_seguimiento_preview_against_bd(
    preview: SeguimientoPreview,
    sede: Sede,
) -> SeguimientoMatchReport:
    """
    Cruza filas del preview con:
    - todas las escuelas operativas activas de la sede (el Excel trae el texto de la oferta por fila);
    - estudiantes activos de la sede por número de documento.

    La persistencia (usuario, contraseña, matrícula) la realiza aparte ``execute_seguimiento_import``.
    """
    headers = preview.headers
    idx_doc = _resolve_documento_column(headers, preview.rows)
    idx_nom = _resolve_header_column(headers, _HEADER_ALIASES_NOMBRE)
    idx_inst = _resolve_header_column(headers, _HEADER_ALIASES_ESCUELA_INSTITUCION)
    skip_for_oferta = {idx_inst} if idx_inst is not None else set()
    idx_oferta = _resolve_header_column(
        headers, _HEADER_ALIASES_OFERTA_CURSO, skip_indices=skip_for_oferta
    )
    if idx_oferta is None:
        idx_oferta = _resolve_header_column(
            headers, _HEADER_ALIASES_ESCUELA_UNICA_LEGACY, skip_indices=skip_for_oferta
        )
    if idx_oferta is None:
        idx_oferta = _resolve_header_column(
            headers, _HEADER_ALIASES_OFERTA_CURSO + _HEADER_ALIASES_ESCUELA_UNICA_LEGACY
        )

    file_notes: list[str] = []
    if idx_doc is None:
        file_notes.append(
            'No se detectó columna de documento (encabezado ni valores típicos de cédula en la muestra); no se puede buscar estudiantes en la base de datos.'
        )
    if idx_nom is None:
        file_notes.append(
            'No se detectó una columna de nombre reconocible; conviene revisar el archivo manualmente.'
        )
    if idx_oferta is None:
        file_notes.append(
            'No se detectó columna de oferta/curso ni «escuela» con texto de inscripción; el emparejado con ediciones será débil o vacío.'
        )
    if idx_inst is not None:
        file_notes.append(
            'Columna de nombre de escuela detectada: cada fila se cruza primero con Escuela (operativa) de la sede; solo se consideran ediciones de esa escuela.'
        )

    catalogo = _ediciones_catalogo_para_sede(sede)
    if not catalogo:
        file_notes.append(
            'No hay escuelas operativas activas en la base para esta sede; no se puede emparejar el texto del Excel con ofertas hasta que existan escuelas.'
        )

    SCORE_MIN_ACEPTAR = 0.10
    SCORE_AMBIGUO_DIF = 0.06

    rows_out: list[RowMatchAnalysis] = []
    for i, row in enumerate(preview.rows, start=1):

        def cell(col: int | None) -> str:
            if col is None or col >= len(row):
                return ''
            return (row[col] or '').strip()

        doc_excel = cell(idx_doc)
        nom_excel = cell(idx_nom)
        esc_inst_excel = cell(idx_inst) if idx_inst is not None else ''
        esc_oferta_excel = cell(idx_oferta) if idx_oferta is not None else ''

        notes: list[str] = []
        edicion_pick: EdicionMatchDetail | None = None
        alternativas: list[tuple[float, EdicionMatchDetail]] = []

        escuela_ids_filtro, notas_filtro = _escuela_ids_restriccion_fila(sede, esc_inst_excel, esc_oferta_excel)
        for nf in notas_filtro:
            notes.append(nf)

        if catalogo and idx_oferta is not None and esc_oferta_excel:
            scored: list[tuple[float, Any]] = []
            for esc_row, blob, escuela_pk in catalogo:
                if escuela_ids_filtro is not None:
                    if escuela_pk is None or escuela_pk not in escuela_ids_filtro:
                        continue
                sc = _score_texto_vs_catalogo(esc_oferta_excel, blob)
                scored.append((sc, esc_row))
            scored.sort(key=lambda x: -x[0])
            if scored:
                best_sc, best_ed = scored[0]
                det = EdicionMatchDetail(
                    edicion_id=best_ed.pk,
                    label=_label_escuela_para_ui(best_ed),
                    score=round(best_sc, 3),
                )
                for sc, esc_cand in scored[1:4]:
                    if sc <= 0:
                        continue
                    alternativas.append(
                        (
                            sc,
                            EdicionMatchDetail(
                                edicion_id=esc_cand.pk,
                                label=_label_escuela_para_ui(esc_cand),
                                score=round(sc, 3),
                            ),
                        )
                    )
                second = scored[1][0] if len(scored) > 1 else 0.0
                if best_sc >= SCORE_MIN_ACEPTAR:
                    edicion_pick = det
                    if second >= best_sc - SCORE_AMBIGUO_DIF and second >= SCORE_MIN_ACEPTAR * 0.85:
                        notes.append(
                            'Varias ediciones tienen puntuación parecida; revisa cuál corresponde al texto del Excel.'
                        )
                elif best_sc > 0:
                    notes.append(
                        f'La mejor coincidencia con una edición es débil ({best_sc:.2f}); puede no ser la oferta correcta.'
                    )
                    edicion_pick = det
                else:
                    if escuela_ids_filtro == frozenset():
                        notes.append(
                            'No se emparejó edición: el nombre de escuela del Excel no corresponde a ninguna escuela en sistema.'
                        )
                    elif escuela_ids_filtro is not None and len(scored) == 0:
                        notes.append(
                            'No hay ediciones activas en la escuela indicada del Excel para esta sede.'
                        )
                    else:
                        notes.append(
                            'No hubo coincidencia razonable con ninguna edición activa de esta sede (revisa curso, docente y horario en BD).'
                        )
            elif escuela_ids_filtro == frozenset():
                pass
            elif escuela_ids_filtro is not None:
                notes.append(
                    'No hay ediciones en catálogo para la escuela filtrada por esta fila (comprueba Escuela y ediciones activas en Hechos).'
                )
            else:
                notes.append('No hubo candidatos de edición para comparar.')
        elif catalogo and esc_oferta_excel and idx_oferta is None:
            notes.append('Falta columna de oferta/curso en el archivo para emparejar con ediciones.')
        elif not esc_oferta_excel and idx_oferta is not None:
            notes.append('Celda de oferta/curso vacía en esta fila.')

        est_id: int | None = None
        est_nombre: str | None = None
        if idx_doc is not None and doc_excel:
            est = _find_estudiante_por_documento(sede.pk, doc_excel)
            if est:
                est_id = est.pk
                est_nombre = est.user.get_full_name() or est.user.username
            else:
                notes.append(
                    'Sin estudiante activo en esta sede con ese documento; la importación puede crear cuenta y matrícula si activas esa opción al subir el archivo.'
                )

        status = 'ok'
        if edicion_pick is None and catalogo:
            status = 'error'
        elif edicion_pick is None and not catalogo:
            status = 'error'
        elif 'débil' in ' '.join(notes) or 'parecida' in ' '.join(notes):
            status = 'warning'
        elif idx_doc is None or not doc_excel:
            status = 'warning'

        alt_tuple = tuple(x[1] for x in sorted(alternativas, key=lambda t: -t[0])[:3])

        rows_out.append(
            RowMatchAnalysis(
                row_number=i,
                doc_excel=doc_excel,
                nombre_excel=nom_excel,
                escuela_texto_excel=esc_oferta_excel,
                escuela_institucion_excel=esc_inst_excel,
                edicion=edicion_pick,
                edicion_alternativas=alt_tuple,
                estudiante_id=est_id,
                estudiante_nombre=est_nombre,
                status=status,
                notes=tuple(notes),
            )
        )

    return SeguimientoMatchReport(
        column_documento=idx_doc,
        column_nombre=idx_nom,
        column_escuela_texto=idx_oferta,
        column_escuela_institucional=idx_inst,
        file_notes=tuple(file_notes),
        ediciones_en_catalogo=len(catalogo),
        rows=rows_out,
    )


def _cell_display(value: Any) -> str:
    if value is None:
        return ''
    if isinstance(value, float):
        if value == int(value):
            return str(int(value))
        return str(value)
    if isinstance(value, Decimal):
        return format(value, 'f').rstrip('0').rstrip('.') or '0'
    if isinstance(value, datetime):
        return value.isoformat(sep=' ', timespec='minutes')
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def build_seguimiento_preview(
    file_obj,
    *,
    max_rows: int = DEFAULT_PREVIEW_MAX_ROWS,
) -> SeguimientoPreview:
    """
    Lee hasta ``max_rows`` filas de datos (sin contar encabezados) de la primera hoja.
    No modifica la base de datos. ``file_obj`` debe ser un UploadedFile u objeto compatible con openpyxl.
    """
    from openpyxl import load_workbook

    if hasattr(file_obj, 'seek'):
        file_obj.seek(0)

    wb = None
    try:
        try:
            wb = load_workbook(file_obj, read_only=True, data_only=True)
        except Exception as exc:
            raise SeguimientoImportError(
                'No se pudo abrir el archivo. Comprueba que sea un Excel .xlsx válido y no esté corrupto.'
            ) from exc

        if not wb.sheetnames:
            raise SeguimientoImportError('El archivo no contiene ninguna hoja.')
        ws = wb[wb.sheetnames[0]]
        sheet_title = ws.title or 'Hoja 1'
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration:
            raise SeguimientoImportError('La primera hoja está vacía.')

        headers = [_cell_display(c) for c in header_row]

        data_rows: list[list[str]] = []
        has_more = False
        for i, row in enumerate(rows_iter):
            if i >= max_rows:
                has_more = True
                break
            data_rows.append([_cell_display(v) for v in row])
    finally:
        if wb is not None:
            wb.close()

    return SeguimientoPreview(
        headers=headers,
        rows=data_rows,
        sheet_name=sheet_title,
        has_more_rows=has_more,
        preview_row_limit=max_rows,
    )

from django.shortcuts import render, get_object_or_404, redirect
from django.urls import reverse, resolve
from django.views.decorators.http import require_POST
from django.contrib.auth.decorators import login_required
from django.contrib import messages
import os

from django.http import FileResponse, Http404, JsonResponse, HttpResponse
from django.db.models import Q, Count, Avg, Exists, OuterRef, Sum, Max
from django.db import IntegrityError
from django.db import models
from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError as DjangoValidationError
from django.utils import timezone
from django.core.mail import EmailMessage
from django.conf import settings
from .models import (
    Estudiante, Profesor, AdminEscuela, Escuela, EscuelaHorario, EscuelaPrograma, RutaEstudio, Curso,
    Matricula, Clase, Asistencia, AsistenciaGrillaCelda, EscuelaGrillaAsistenciaConfig,
    EscuelaGrillaNotasConfig, NotasGrillaCelda, NotasGrillaEstadoFinal, Nota,
    SolicitudMatricula, NotificacionEstudiante,
    SolicitudEspecialEstudiante, QuejaReclamo, Ofrenda, PresupuestoEvento,
    RecaudoOcasional, Salon, ActividadEscuela, EntregaActividad, ArchivoRecursoEscuela,
    NivelProgramaPlantilla, EscuelaProgramaPlantilla, default_anio_escuela,
)
from .utils import (
    get_current_period,
    get_period_from_date,
    get_period_display,
    get_period_stats,
    sesiones_agrupadas_por_mes,
)
from core.models import Sede, generate_unique_username
from hechos import coordinador_access as ca
from .user_sede import get_user_sede
from .escuela_operativa import edicion_prioritaria_para_matricula, ediciones_activas_en_sede
from .portal_entry import home_redirect_response
from .profesor_entregas import entregas_pendientes_calificacion_qs
from .forms import (
    ActividadEscuelaForm,
    ArchivoRecursoEscuelaForm,
    EntregaActividadTextoForm,
    entrega_evidencia_archivo_validator,
    normalizar_archivo_entrega,
    PresupuestoEventoForm,
    ProfesorEvaluaEntregaForm,
    QuejaReclamoForm,
    RecaudoOcasionalForm,
    SalonForm,
    ingresos_totales_sede,
)

from collections import Counter, defaultdict, deque
from datetime import datetime, timedelta, date, time as dt_time
from decimal import Decimal
from urllib.parse import quote, urlencode
from django.utils.text import slugify
import re
import secrets
from django.db import transaction
import math

User = get_user_model()


def _whatsapp_me_url(phone_raw):
    """Devuelve URL base https://wa.me/... o None si no hay número válido."""
    if not phone_raw:
        return None
    digits = "".join(c for c in str(phone_raw) if c.isdigit())
    if not digits:
        return None
    if len(digits) == 10:
        digits = "57" + digits
    return f"https://wa.me/{digits}"


def _redirect_coordinador_logistica(*, tab='salones', escuela_id=None):
    base = reverse('hechos:coordinador_logistica')
    q = [f'tab={tab}']
    if escuela_id is not None:
        q.append(f'escuela={int(escuela_id)}')
    return redirect(f'{base}?{"&".join(q)}')


def _norm_catalogo_escuela_key(s: str) -> str:
    return (s or '').strip().casefold()


def _build_escuela_programa_lookup_por_nombre(sede_ids) -> dict:
    """
    (sede_id, nombre de escuela operativa normalizado) → EscuelaPrograma.
    Cruza por nombre local y por nombre_para_catalogo (plantilla global).
    Si hay colisión, gana la fila con menor jerarquía de nivel (entrada más baja al programa).
    """
    if not sede_ids:
        return {}
    out = {}
    qs = EscuelaPrograma.objects.filter(
        sede_id__in=sede_ids,
        is_active=True,
    ).select_related('nivel_programa', 'plantilla', 'nivel_programa__plantilla')
    for ep in qs:
        labels = {_norm_catalogo_escuela_key(ep.nombre)}
        labels.add(_norm_catalogo_escuela_key(ep.nombre_para_catalogo))
        labels.discard('')
        for lbl in labels:
            k = (ep.sede_id, lbl)
            prev = out.get(k)
            if prev is None:
                out[k] = ep
            elif ep.nivel_programa.jerarquia < prev.nivel_programa.jerarquia:
                out[k] = ep
    return out


def _escuela_programa_para_escuela_operativa(escuela: Escuela, sede_id: int | None, lookup: dict):
    if not escuela or not sede_id:
        return None
    return lookup.get((sede_id, _norm_catalogo_escuela_key(escuela.nombre)))


# Alta mínima de una sola edición por escuela (nombre interno estable, sin «Principal» genérico).
_EDICION_INICIAL_UNICA = 'Única'


@login_required
def hechos_dashboard(request):
    """
    Punto de entrada de /hechos/ (y /hechos/dashboard/): mismo criterio de destino
    que el portal (``hechos.portal_entry.home_redirect_response``), con avisos si
    aplica. El registro de ofrendas vive en ``ofrendas`` (/ofrendas/).
    """
    u = request.user
    has_hechos = bool(
        hasattr(u, 'estudiante_profile')
        or hasattr(u, 'profesor_profile')
        or hasattr(u, 'admin_escuela_profile')
    )
    if not has_hechos:
        messages.warning(request, 'No tienes acceso al módulo Hechos.')
        return home_redirect_response(u)

    if not get_user_sede(u) and not hasattr(u, 'estudiante_profile'):
        messages.warning(
            request,
            'No tienes una sede asignada. Contacta al administrador.',
        )
        return home_redirect_response(u)

    return home_redirect_response(u)


@login_required
def seleccionar_sede_estudiante(request):
    if not hasattr(request.user, 'estudiante_profile'):
        messages.error(request, 'Solo los estudiantes pueden seleccionar sede.')
        return redirect('core:dashboard')

    estudiante = request.user.estudiante_profile
    sedes_qs = Sede.objects.filter(is_active=True).order_by('ciudad', 'nombre')
    next_url = request.GET.get('next') or request.POST.get('next') or ''

    if request.method == 'POST':
        sede_id = request.POST.get('sede_id')
        sede = get_object_or_404(sedes_qs, id=sede_id)
        estudiante.sede = sede
        estudiante.save(update_fields=['sede'])
        if request.user.sede_id != sede.id:
            request.user.sede = sede
            request.user.save(update_fields=['sede'])
        messages.success(request, f'Sede asignada: {sede.titulo_con_ciudad()}.')

        if next_url and next_url.startswith('/'):
            return redirect(next_url)
        return redirect('hechos:escuelas_disponibles')

    sedes_con_ciudad = (
        Sede.objects.filter(is_active=True)
        .exclude(ciudad__exact='')
        .order_by('ciudad', 'nombre')
    )
    por_ciudad = {}
    for s in sedes_con_ciudad:
        c = (s.ciudad or '').strip()
        if not c:
            continue
        por_ciudad.setdefault(c, []).append(
            {'id': s.id, 'label': s.titulo_con_ciudad()}
        )
    ciudades = sorted(por_ciudad.keys(), key=str.casefold)

    sedes_sin_ciudad = [
        {'id': s.id, 'label': s.nombre}
        for s in Sede.objects.filter(is_active=True)
        .filter(Q(ciudad='') | Q(ciudad__isnull=True))
        .order_by('nombre')
    ]

    sede_actual = estudiante.sede if estudiante.sede_id else None
    ciudad_inicial = ''
    usar_bucket_otras = False
    if sede_actual:
        c0 = (sede_actual.ciudad or '').strip()
        if c0:
            ciudad_inicial = c0
        elif sedes_sin_ciudad:
            usar_bucket_otras = True

    return render(
        request,
        'hechos/seleccionar_sede_estudiante.html',
        {
            'sedes_todas': sedes_qs,
            'por_ciudad': por_ciudad,
            'sedes_sin_ciudad_data': sedes_sin_ciudad,
            'ciudades': ciudades,
            'tiene_cascada_ciudad': bool(ciudades),
            'sede_actual_id': estudiante.sede_id or request.user.sede_id,
            'sede_actual': sede_actual,
            'ciudad_inicial': ciudad_inicial,
            'usar_bucket_otras': usar_bucket_otras,
            'is_change': bool(estudiante.sede_id),
            'next_url': next_url,
        },
    )


@login_required
def estudiantes_list(request):
    """
    Lista de estudiantes (para admins y profesores)
    """
    allowed = (
        request.user.is_super_admin()
        or hasattr(request.user, 'profesor_profile')
        or (
            hasattr(request.user, 'admin_escuela_profile')
            and ca.has_capacidad(request.user, 'academico')
        )
    )
    if not allowed:
        messages.error(request, 'No tienes permisos para ver esta página.')
        return redirect('core:dashboard')
    
    # Obtener la sede del usuario
    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada. Contacta al administrador.')
        return redirect('core:dashboard')
    
    # Filtrar estudiantes según el perfil del usuario
    if hasattr(request.user, 'profesor_profile'):
        # Profesor ve solo estudiantes de sus cursos
        estudiantes = Estudiante.objects.filter(
            matriculas__escuela__maestro=request.user.profesor_profile,
            matriculas__sede=user_sede,
            matriculas__is_active=True,
            sede=user_sede,
            is_active=True
        ).distinct().select_related('user').prefetch_related('matriculas')
    else:
        # Admin ve todos los estudiantes
        estudiantes = Estudiante.objects.filter(sede=user_sede, is_active=True).select_related('user').prefetch_related('matriculas')

    busqueda = (request.GET.get('busqueda') or '').strip()
    filtro_estado = (request.GET.get('estado') or '').strip()
    if busqueda:
        estudiantes = estudiantes.filter(
            Q(user__first_name__icontains=busqueda)
            | Q(user__last_name__icontains=busqueda)
            | Q(user__email__icontains=busqueda)
            | Q(codigo_estudiante__icontains=busqueda)
            | Q(iglesia__icontains=busqueda)
        ).distinct()
    if filtro_estado == 'matriculados':
        estudiantes = estudiantes.filter(
            matriculas__is_active=True,
            matriculas__sede=user_sede,
        ).distinct()
    elif filtro_estado == 'no_matriculados':
        matricula_activa = Matricula.objects.filter(
            estudiante_id=OuterRef('pk'),
            is_active=True,
            sede=user_sede,
        )
        estudiantes = estudiantes.filter(~Exists(matricula_activa))

    # Estadísticas
    estudiantes_activos = estudiantes.filter(is_active=True)
    estudiantes_matriculados = estudiantes.filter(
        matriculas__is_active=True
    ).distinct()
    
    # Calcular promedio general
    from django.db.models import Avg
    promedio_general = estudiantes.aggregate(
        promedio=Avg('notas__puntaje_obtenido')
    )['promedio'] or 0
    
    # Calcular estadísticas por estudiante
    estudiantes_con_estadisticas = []
    for estudiante in estudiantes:
        matriculas_activas = estudiante.matriculas.filter(is_active=True).count()
        total_matriculas = estudiante.matriculas.count()
        estudiantes_con_estadisticas.append({
            'estudiante': estudiante,
            'matriculas_activas': matriculas_activas,
            'total_matriculas': total_matriculas
        })
    
    context = {
        'estudiantes_con_estadisticas': estudiantes_con_estadisticas,
        'estudiantes_activos': estudiantes_activos,
        'estudiantes_matriculados': estudiantes_matriculados,
        'promedio_general': round(promedio_general, 1),
        'user_sede': user_sede,
        'busqueda': busqueda,
        'filtro_estado': filtro_estado,
    }
    
    return render(request, 'hechos/estudiantes_list_modern.html', context)


@login_required
def estudiantes_listado_por_sede(request):
    """
    Directorio global: todos los estudiantes activos de todas las sedes, con búsqueda transversal
    (nombre, correo, documento, código, nombre de sede). Para investigación y localizar dónde cursa.
    """
    allowed = (
        request.user.is_super_admin()
        or hasattr(request.user, 'profesor_profile')
        or (
            hasattr(request.user, 'admin_escuela_profile')
            and ca.has_capacidad(request.user, 'academico')
        )
    )
    if not allowed:
        messages.error(request, 'No tienes permisos para ver esta página.')
        return redirect('core:dashboard')

    user_sede = get_user_sede(request.user)

    base_qs = (
        Estudiante.objects.filter(is_active=True)
        .select_related('user', 'sede')
        .order_by('sede__nombre', 'user__last_name', 'user__first_name', 'user__email')
    )
    total_plataforma = base_qs.count()

    q = (request.GET.get('q') or '').strip()
    qs = base_qs
    if q:
        qs = qs.filter(
            Q(user__first_name__icontains=q)
            | Q(user__last_name__icontains=q)
            | Q(user__email__icontains=q)
            | Q(numero_documento__icontains=q)
            | Q(codigo_estudiante__icontains=q)
            | Q(sede__nombre__icontains=q)
        ).distinct()

    total_coincidencias = qs.count()
    estudiantes = list(qs)

    puede_ver_detalle_estudiante = allowed

    return render(
        request,
        'hechos/estudiantes_listado_por_sede.html',
        {
            'user_sede': user_sede,
            'estudiantes': estudiantes,
            'busqueda_q': q,
            'total_plataforma': total_plataforma,
            'total_coincidencias': total_coincidencias,
            'puede_ver_detalle_estudiante': puede_ver_detalle_estudiante,
        },
    )


def _time_from_hora_12_am_pm(hour_1_12, am_pm: str) -> dt_time | None:
    """Convierte hora 1–12 y 'am'/'pm' a datetime.time (minutos 0)."""
    ap = (am_pm or '').strip().lower()
    if ap not in ('am', 'pm'):
        return None
    try:
        h = int(hour_1_12)
    except (TypeError, ValueError):
        return None
    if h < 1 or h > 12:
        return None
    if ap == 'am':
        h24 = 0 if h == 12 else h
    else:
        h24 = 12 if h == 12 else h + 12
    return dt_time(hour=h24, minute=0)


def _horario_edicion_a_texto_ui(escuela) -> str:
    """Texto de horario para UI (español, 12 h) a partir de filas o del campo libre."""
    horarios = list(escuela.estructura_horarios.all())
    if horarios:
        horarios.sort(key=lambda h: (h.dia_semana, h.hora))
        return ' · '.join(
            _horario_texto_dia_hora_12(h.dia_semana, h.hora) for h in horarios
        )
    return (escuela.horario or '').strip()


def _modalidad_edicion_label(escuela) -> str:
    if not escuela:
        return 'Presencial'
    try:
        return escuela.get_modalidad_display() or 'Presencial'
    except Exception:
        return 'Presencial'


def _horario_texto_dia_hora_12(dia_idx: int, t: dt_time) -> str:
    """Texto legible en español para día de semana + hora (grilla de estructura)."""
    dias = ['Lunes', 'Martes', 'Miércoles', 'Jueves', 'Viernes', 'Sábado', 'Domingo']
    if dia_idx < 0 or dia_idx > 6:
        return ''
    h24 = t.hour
    h12 = h24 % 12 or 12
    suf = 'a. m.' if h24 < 12 else 'p. m.'
    return f'{dias[dia_idx]}, {h12}:00 {suf}'


def _parse_horario_filas_crear_instancia(request) -> list[tuple[int, dt_time]]:
    """
    Lee filas paralelas dia_semana / hora_12 / hora_am_pm (getlist).
    Obligatorio al menos una fila completa; no se admite el mismo día dos veces.
    """
    dias = request.POST.getlist('dia_semana')
    horas = request.POST.getlist('hora_12')
    amps = request.POST.getlist('hora_am_pm')
    if len(dias) != len(horas) or len(dias) != len(amps):
        raise ValueError('Los datos del horario no coinciden. Recarga la página e intenta de nuevo.')
    rows: list[tuple[int, dt_time]] = []
    for i in range(len(dias)):
        dr = (dias[i] or '').strip()
        hr = (horas[i] or '').strip()
        ar = (amps[i] or '').strip()
        nfill = sum(1 for x in (dr, hr, ar) if x)
        if nfill == 0:
            continue
        if nfill != 3:
            raise ValueError(
                'En cada fila debes elegir día, hora y a. m. o p. m., o pulsa «Quitar» en la fila que no uses.'
            )
        try:
            d_int = int(dr)
        except ValueError as e:
            raise ValueError('Día de la semana no válido.') from e
        if not (0 <= d_int <= 6):
            raise ValueError('Día de la semana no válido.')
        hora_t = _time_from_hora_12_am_pm(hr, ar)
        if hora_t is None:
            raise ValueError('Hora no válida.')
        rows.append((d_int, hora_t))
    if not rows:
        raise ValueError('El horario es obligatorio: indica al menos un día con su hora.')
    seen: set[int] = set()
    for d_int, _ in rows:
        if d_int in seen:
            raise ValueError('No puedes repetir el mismo día en dos filas; un solo horario por día.')
        seen.add(d_int)
    rows.sort(key=lambda x: x[0])
    return rows


def _escuela_nombre_disponible(sede, nombre: str, anio: int, ciclo: str) -> bool:
    """Indica si no existe una escuela activa con misma sede+nombre+año+ciclo."""
    base = (nombre or '').strip()
    if not base:
        return False
    return not Escuela.objects.filter(
        sede=sede,
        nombre__iexact=base,
        anio=anio,
        ciclo=ciclo,
        is_active=True,
    ).exists()


@login_required
def profesor_crear_instancia_escuela(request):
    """
    El docente crea una Escuela operativa en su sede a partir del catálogo global
    (nivel → escuela plantilla), con horario obligatorio (uno o más días) y salón opcional.
    """
    if not hasattr(request.user, 'profesor_profile'):
        messages.error(request, 'Solo los profesores pueden crear una instancia de escuela desde aquí.')
        return redirect('core:dashboard')

    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada.')
        return redirect('core:dashboard')

    prof = request.user.profesor_profile
    niveles = list(NivelProgramaPlantilla.objects.order_by('jerarquia', 'nombre'))
    plantillas_qs = EscuelaProgramaPlantilla.objects.select_related('nivel_plantilla').order_by(
        'nivel_plantilla__jerarquia', 'nombre'
    )
    plantillas_data = [
        {
            'id': p.id,
            'nivel_id': p.nivel_plantilla_id,
            'nombre': p.nombre,
            'nivel_label': p.nivel_plantilla.titulo_acordeon(),
        }
        for p in plantillas_qs
    ]
    salones_libres = list(
        Salon.objects.filter(sede=user_sede, is_active=True, escuela__isnull=True).order_by('nombre')
    )

    if request.method == 'POST':
        try:
            pid = int(request.POST.get('plantilla_id') or '0')
        except ValueError:
            pid = 0
        plantilla = EscuelaProgramaPlantilla.objects.filter(id=pid).select_related('nivel_plantilla').first()
        if not plantilla:
            messages.error(request, 'Elige una escuela del catálogo global.')
            return redirect('hechos:profesor_crear_instancia_escuela')

        nombre_esc = (plantilla.nombre or '').strip()
        if not nombre_esc:
            messages.error(request, 'La plantilla no tiene nombre válido.')
            return redirect('hechos:profesor_crear_instancia_escuela')

        try:
            anio = int(request.POST.get('anio_escuela') or '0')
        except ValueError:
            anio = 0
        if anio < 2016 or anio > 2030:
            messages.error(request, 'Elige un año académico entre 2016 y 2030.')
            return redirect('hechos:profesor_crear_instancia_escuela')
        ciclo_post = (request.POST.get('ciclo') or Escuela.Ciclo.A).strip().upper()
        if ciclo_post not in (Escuela.Ciclo.A, Escuela.Ciclo.B):
            ciclo_post = Escuela.Ciclo.A
        ciclo = ciclo_post
        if not _escuela_nombre_disponible(user_sede, nombre_esc, anio, ciclo):
            messages.error(
                request,
                'Ya existe una escuela activa con ese nombre para el mismo año y ciclo en la sede.',
            )
            return redirect('hechos:profesor_crear_instancia_escuela')

        try:
            horario_rows = _parse_horario_filas_crear_instancia(request)
        except ValueError as e:
            messages.error(request, str(e))
            return redirect('hechos:profesor_crear_instancia_escuela')

        # Regla operativa: un profesor no puede estar asignado al mismo día y hora en dos escuelas activas.
        # (La diferencia real entre instancias es el horario.)
        q_conf = Q()
        for d_int, hora_t in horario_rows:
            q_conf |= Q(dia_semana=d_int, hora=hora_t)
        if q_conf:
            choque = (
                EscuelaHorario.objects.filter(q_conf)
                .filter(
                    escuela__maestro=prof,
                    escuela__is_active=True,
                    escuela__sede=user_sede,
                )
                .select_related('escuela')
                .order_by('dia_semana', 'hora')
                .first()
            )
            if choque:
                day_labels = {int(k): v for k, v in EscuelaHorario.DiaSemana.choices}
                dia_txt = day_labels.get(int(choque.dia_semana), str(choque.dia_semana))
                hora_txt = choque.hora.strftime('%I:%M %p').lstrip('0').lower()
                esc_choque = choque.escuela
                messages.error(
                    request,
                    f'Conflicto de horario: ya tienes una escuela asignada el {dia_txt} a las {hora_txt} '
                    f'({esc_choque.titulo_tarjeta()}). Elige otra hora.',
                )
                return redirect('hechos:profesor_crear_instancia_escuela')

        fecha_inicio = _parse_iso_date(request.POST.get('fecha_inicio'))
        fecha_fin = _parse_iso_date(request.POST.get('fecha_fin'))
        if not fecha_inicio:
            messages.error(request, 'Indica la fecha de inicio del ciclo de clases.')
            return redirect('hechos:profesor_crear_instancia_escuela')
        if not fecha_fin:
            messages.error(request, 'Indica la fecha de fin del ciclo de clases.')
            return redirect('hechos:profesor_crear_instancia_escuela')
        if fecha_fin < fecha_inicio:
            messages.error(request, 'La fecha de fin debe ser igual o posterior a la de inicio.')
            return redirect('hechos:profesor_crear_instancia_escuela')

        delta_days = (fecha_fin - fecha_inicio).days
        curso_duracion_semanas = max(1, (delta_days + 6) // 7)

        horario_txt = ' · '.join(
            _horario_texto_dia_hora_12(d_int, hora_t) for d_int, hora_t in horario_rows
        )

        salon_id = (request.POST.get('salon_id') or '').strip()
        modalidad = (request.POST.get('modalidad') or Escuela.Modalidad.PRESENCIAL).strip().lower()
        modalidades_validas = {
            Escuela.Modalidad.PRESENCIAL,
            Escuela.Modalidad.VIRTUAL,
            Escuela.Modalidad.HIBRIDA,
        }
        if modalidad not in modalidades_validas:
            modalidad = Escuela.Modalidad.PRESENCIAL

        creadas = 0
        try:
            with transaction.atomic():
                escuela = Escuela.objects.create(
                    sede=user_sede,
                    nombre=nombre_esc,
                    descripcion=(plantilla.descripcion or '').strip(),
                    anio=anio,
                    ciclo=ciclo,
                    maestro=prof,
                    fecha_inicio=fecha_inicio,
                    fecha_fin=fecha_fin,
                    modalidad=modalidad,
                    horario=horario_txt,
                    aula='',
                    cupo_maximo=30,
                    is_active=True,
                )

                ruta = RutaEstudio.objects.create(
                    sede=user_sede,
                    escuela=escuela,
                    nombre=(nombre_esc or '').strip() or escuela.nombre,
                    descripcion='',
                    duracion_semanas=12,
                    nivel='basico',
                    is_active=True,
                )
                Curso.objects.create(
                    sede=user_sede,
                    ruta_estudio=ruta,
                    nombre=(nombre_esc or '').strip() or escuela.nombre,
                    descripcion='',
                    orden=1,
                    duracion_semanas=curso_duracion_semanas,
                    is_active=True,
                )
                salon_pk = None
                if salon_id:
                    try:
                        salon_pk = int(salon_id)
                    except ValueError:
                        salon_pk = None
                if salon_pk:
                    sal = Salon.objects.select_for_update().filter(
                        id=salon_pk, sede=user_sede, is_active=True
                    ).first()
                    if sal and sal.escuela_id is None:
                        sal.escuela = escuela
                        sal.save(update_fields=['escuela', 'updated_at'])
                        escuela.aula = sal.nombre
                        escuela.cupo_maximo = max(1, int(sal.capacidad_plazas or 30))
                        escuela.save(update_fields=['aula', 'cupo_maximo', 'updated_at'])
                    elif sal and sal.escuela_id != escuela.id:
                        raise ValueError('El salón elegido ya está asignado a otra escuela.')

                for d_int, hora_t in horario_rows:
                    EscuelaHorario.objects.get_or_create(
                        escuela=escuela,
                        dia_semana=d_int,
                        defaults={'hora': hora_t},
                    )

                creadas = _crear_clases_si_vacias_desde_estructura(escuela, user_sede, prof)
                if creadas > 0:
                    cfg_grilla, _ = EscuelaGrillaAsistenciaConfig.objects.get_or_create(
                        escuela=escuela,
                        sede=user_sede,
                        defaults={'num_columnas': GRILLA_ASISTENCIA_COLS_DEFECTO},
                    )
                    want_cols = min(
                        max(creadas, cfg_grilla.num_columnas),
                        GRILLA_ASISTENCIA_COLS_MAX,
                    )
                    if want_cols > cfg_grilla.num_columnas:
                        cfg_grilla.num_columnas = want_cols
                        cfg_grilla.save(update_fields=['num_columnas', 'updated_at'])

            msg_ok = (
                f'Se creó la escuela «{escuela.titulo_tarjeta()}» con la primera edición del curso. '
                f'Ya puedes gestionarla en Mis escuelas.'
            )
            if creadas:
                msg_ok += (
                    f' Se generaron {creadas} clases en el calendario '
                    f'(una por cada día de clase en el período).'
                )
            messages.success(request, msg_ok)
            return redirect(f"{reverse('hechos:mis_escuelas_profesor')}?escuela={escuela.id}")
        except IntegrityError:
            messages.error(request, 'No se pudo crear: puede haber un duplicado. Revisa o intenta de nuevo.')
        except ValueError as e:
            messages.error(request, str(e))
        except Exception as e:
            messages.error(request, f'Error al crear la instancia: {e}')

    y0 = max(2016, min(2030, default_anio_escuela()))
    hoy = timezone.localdate()
    return render(
        request,
        'hechos/profesor_crear_instancia_escuela.html',
        {
            'user_sede': user_sede,
            'niveles': niveles,
            'plantillas_data': plantillas_data,
            'salones_libres': salones_libres,
            'anios_academicos': list(range(2016, 2031)),
            'anio_escuela_default': y0,
            'fecha_inicio_default': hoy.strftime('%Y-%m-%d'),
            'fecha_fin_default': (hoy + timedelta(weeks=12)).strftime('%Y-%m-%d'),
            'modalidad_default': Escuela.Modalidad.PRESENCIAL,
        },
    )


@login_required
@require_POST
def profesor_escuela_dar_de_baja(request, escuela_id):
    """
    Baja lógica de una instancia de escuela operativa. Solo el maestro asignado
    (p. ej. quien la creó desde «Crear instancia») en su sede.

    Desactiva escuela, ruta, cursos, ediciones y clases; cancela matrículas activas
    vinculadas a esas ediciones; elimina celdas de grilla de asistencia; libera salón.
    """
    if not hasattr(request.user, 'profesor_profile'):
        messages.error(request, 'Solo los profesores pueden usar esta acción.')
        return redirect('core:dashboard')
    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada.')
        return redirect('core:dashboard')
    prof = request.user.profesor_profile
    escuela = get_object_or_404(
        Escuela,
        id=escuela_id,
        sede=user_sede,
        maestro=prof,
        is_active=True,
    )
    nombre = escuela.titulo_tarjeta()
    try:
        with transaction.atomic():
            now = timezone.now()
            Matricula.objects.filter(
                sede=user_sede,
                is_active=True,
                escuela=escuela,
            ).update(is_active=False, estado='cancelada', updated_at=now)
            AsistenciaGrillaCelda.objects.filter(escuela=escuela, sede=user_sede).delete()
            NotasGrillaCelda.objects.filter(escuela=escuela, sede=user_sede).delete()
            NotasGrillaEstadoFinal.objects.filter(escuela=escuela, sede=user_sede).delete()
            EscuelaGrillaNotasConfig.objects.filter(escuela=escuela, sede=user_sede).delete()
            Salon.objects.filter(sede=user_sede, escuela=escuela).update(escuela=None)

            escuela.is_active = False
            escuela.save(update_fields=['is_active', 'updated_at'])

            RutaEstudio.objects.filter(escuela=escuela, is_active=True).update(is_active=False)
            Curso.objects.filter(ruta_estudio__escuela=escuela, is_active=True).update(is_active=False)
            Clase.objects.filter(
                sede=user_sede,
                escuela=escuela,
                is_active=True,
            ).update(is_active=False)
    except Exception as exc:
        messages.error(request, f'No se pudo dar de baja la escuela: {exc}')
        return redirect(f"{reverse('hechos:mis_escuelas_profesor')}?escuela={escuela.id}")

    messages.success(
        request,
        f'Se dio de baja la escuela «{nombre}». Ya no aparecerá en Mis escuelas; '
        'las personas del sistema no se eliminaron.',
    )
    return redirect('hechos:mis_escuelas_profesor')


@login_required
def profesores_list(request):
    """
    Lista de profesores (director o coordinador pedagógico)
    """
    allowed = (
        request.user.is_super_admin()
        or (
            hasattr(request.user, 'admin_escuela_profile')
            and ca.has_capacidad(request.user, 'pedagogico')
        )
    )
    if not allowed:
        messages.error(request, 'No tienes permisos para ver esta página.')
        return redirect('core:dashboard')
    
    # Obtener la sede del usuario
    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada. Contacta al administrador.')
        return redirect('core:dashboard')
    
    base_qs = Profesor.objects.filter(sede=user_sede, is_active=True)
    experiencia_promedio = base_qs.aggregate(avg=Avg('experiencia_anos'))['avg'] or 0
    escuela_asignada = Escuela.objects.filter(
        maestro_id=OuterRef('pk'),
        is_active=True,
        sede=user_sede,
    )
    profesores_con_asignacion = (
        base_qs.annotate(_tiene_escuela=Exists(escuela_asignada)).filter(_tiene_escuela=True).count()
    )
    total_ediciones_en_sede = Escuela.objects.filter(
        sede=user_sede,
        is_active=True,
        maestro__isnull=False,
    ).count()

    profesores = (
        base_qs.select_related('user')
        .annotate(
            num_ediciones_activas=Count(
                'escuelas_dirigidas',
                filter=Q(
                    escuelas_dirigidas__is_active=True,
                    escuelas_dirigidas__sede=user_sede,
                ),
            )
        )
        .order_by('user__first_name', 'user__last_name')
    )

    es_pedagogico = ca.has_capacidad(request.user, 'pedagogico')

    context = {
        'profesores': profesores,
        'user_sede': user_sede,
        'total_profesores': base_qs.count(),
        'profesores_con_asignacion': profesores_con_asignacion,
        'profesores_sin_asignacion': max(0, base_qs.count() - profesores_con_asignacion),
        'total_ediciones_en_sede': total_ediciones_en_sede,
        'experiencia_promedio': round(float(experiencia_promedio), 1),
        'es_coordinador_pedagogico': bool(es_pedagogico and not request.user.is_super_admin()),
    }
    
    return render(request, 'hechos/profesores_list_modern.html', context)


def _dias_horarios_to_string(rows):
    """
    rows: iterable de (dia_semana:int, hora:datetime.time)
    """
    by_day = {d: h for d, h in rows}
    ordered_days = [d for d in range(0, 7) if d in by_day]
    day_labels = {
        0: "Lunes",
        1: "Martes",
        2: "Miércoles",
        3: "Jueves",
        4: "Viernes",
        5: "Sábado",
        6: "Domingo",
    }
    parts = []
    for d in ordered_days:
        h = by_day[d]
        parts.append(f"{day_labels.get(d, str(d))} {h.strftime('%I:%M %p').lstrip('0')}")
    return " · ".join(parts)


def _calcular_fechas_clase(fecha_inicio, fecha_fin, dias):
    """
    dias: set[int] (0=Lunes .. 6=Domingo) usando weekday() estándar.
    """
    if not fecha_inicio or not fecha_fin or not dias:
        return []
    if fecha_fin < fecha_inicio:
        return []
    fechas = []
    d = fecha_inicio
    while d <= fecha_fin:
        if d.weekday() in dias:
            fechas.append(d)
        d += timedelta(days=1)
    return fechas


def _crear_clases_si_vacias_desde_estructura(escuela, sede, profesor_para_clase):
    """
    Si la escuela tiene días/horarios guardados y aún no hay clases (Clase),
    genera una fila por cada fecha entre fecha_inicio y fecha_fin que coincida
    con esos días (misma lógica que el resumen del asistente).
    """
    if not escuela.fecha_inicio or not escuela.fecha_fin:
        return 0
    if Clase.objects.filter(escuela=escuela, is_active=True).exists():
        return 0
    horarios = list(
        EscuelaHorario.objects.filter(escuela=escuela).order_by("dia_semana", "hora")
    )
    if not horarios:
        return 0
    by_weekday = {h.dia_semana: h.hora for h in horarios}
    dias_set = set(by_weekday.keys())
    fechas = _calcular_fechas_clase(escuela.fecha_inicio, escuela.fecha_fin, dias_set)
    if not fechas:
        return 0
    tz = timezone.get_current_timezone()
    to_create = []
    n = 0
    for d in sorted(fechas):
        hh = by_weekday.get(d.weekday())
        if hh is None:
            continue
        n += 1
        naive = datetime.combine(d, hh)
        aware = timezone.make_aware(naive, tz)
        to_create.append(
            Clase(
                sede=sede,
                escuela=escuela,
                numero_clase=n,
                titulo=f"Sesión {n}",
                descripcion="",
                fecha_clase=aware,
                duracion_minutos=90,
                profesor=profesor_para_clase,
                aula="",
                is_active=True,
            )
        )
    if not to_create:
        return 0
    Clase.objects.bulk_create(to_create)
    return len(to_create)


def _profesor_puede_gestionar_clase(clase, profesor):
    """Titular de la clase, profesor de la edición o maestro de la escuela."""
    if clase.profesor_id and clase.profesor_id == profesor.id:
        return True
    if clase.escuela_id and clase.escuela.maestro_id == profesor.id:
        return True
    esc = None
    try:
        esc = clase.curso.ruta_estudio.escuela
    except Exception:
        return False
    return bool(esc and esc.maestro_id == profesor.id)


def _estructura_wizard_session_key(escuela_id: int) -> str:
    return f"estructura_wiz_{escuela_id}"


def _parse_iso_date(s):
    s = (s or "").strip()
    if not s:
        return None
    try:
        return datetime.strptime(s, "%Y-%m-%d").date()
    except ValueError:
        return None


def _parse_fecha_sesion_texto(s: str) -> date | None:
    """Acepta dd/mm/aa, dd/mm/aaaa, d/m/aa (como en la grilla) o aaaa-mm-dd."""
    s = (s or "").strip()
    if not s:
        return None
    iso = _parse_iso_date(s)
    if iso:
        return iso
    parts = re.split(r"[/.\-]", s)
    if len(parts) != 3:
        return None
    try:
        d, m, y = int(parts[0]), int(parts[1]), int(parts[2])
    except ValueError:
        return None
    if not (1 <= m <= 12 and 1 <= d <= 31):
        return None
    if y < 100:
        y += 2000 if y < 70 else 1900
    try:
        return date(y, m, d)
    except ValueError:
        return None


def _time_from_12h(hour_12: int, minute: int, ampm: str) -> dt_time:
    ampm = (ampm or "").strip().upper()
    if hour_12 < 1 or hour_12 > 12 or minute < 0 or minute > 59:
        raise ValueError("Hora fuera de rango")
    if ampm not in ("AM", "PM"):
        raise ValueError("Indica AM o PM")
    if hour_12 == 12:
        h24 = 0 if ampm == "AM" else 12
    else:
        h24 = hour_12 if ampm == "AM" else hour_12 + 12
    return dt_time(hour=h24, minute=minute)


def _time_to_display_12h(t: dt_time) -> str:
    h = t.hour % 12
    if h == 0:
        h = 12
    suf = "a. m." if t.hour < 12 else "p. m."
    return f"{h}:{t.minute:02d} {suf}"


def _parse_session_time(s: str) -> dt_time:
    parts = (s or "").strip().split(":")
    hh = int(parts[0])
    mm = int(parts[1]) if len(parts) > 1 else 0
    return dt_time(hour=hh, minute=mm)


def _time_to_sel_parts(tt: dt_time):
    h12 = tt.hour % 12
    if h12 == 0:
        h12 = 12
    return h12, tt.minute, "AM" if tt.hour < 12 else "PM"


def _minutos_opts_for(sel_mm: int):
    base = [0, 15, 30, 45]
    if sel_mm not in base:
        return sorted(set(base + [sel_mm]))
    return base


def _build_dias_horario_rows_paso2(dias, day_labels, horas_por_dia, legacy_hora, existing_by_day):
    """
    Una fila por día con selectores 12h. horas_por_dia: dict str(día)->'HH:MM:SS'.
    """
    rows = []
    for d in sorted(dias):
        label = day_labels[d]
        tt = None
        if horas_por_dia and str(d) in horas_por_dia:
            try:
                tt = _parse_session_time(horas_por_dia[str(d)])
            except (ValueError, TypeError, IndexError):
                tt = None
        if tt is None and legacy_hora:
            try:
                tt = _parse_session_time(legacy_hora)
            except (ValueError, TypeError, IndexError):
                tt = None
        if tt is None and d in existing_by_day:
            tt = existing_by_day[d]
        if tt is None:
            tt = dt_time(19, 0)
        h12, mm, ampm = _time_to_sel_parts(tt)
        rows.append(
            {
                "dia_int": d,
                "label": label,
                "sel_h12": h12,
                "sel_mm": mm,
                "sel_ampm": ampm,
                "minutos_opts": _minutos_opts_for(mm),
            }
        )
    return rows


@login_required
def estructura_ediciones(request):
    """
    Coordinador pedagógico: escuelas de la sede; cada una se configura con el asistente (por edición).
    """
    if not request.user.is_super_admin():
        r = ca.require_pedagogico(request)
        if r:
            return r

    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, "No tienes una sede asignada. Contacta al administrador.")
        return redirect("core:dashboard")

    escuelas = Escuela.objects.filter(sede=user_sede, is_active=True).order_by(
        "nombre", "-anio", "ciclo"
    )

    escuelas_feed = []
    for escuela in escuelas:
        escuelas_feed.append(
            {
                "escuela": escuela,
                "num_ediciones": 1,
            }
        )

    return render(
        request,
        "hechos/estructura_escuelas_list.html",
        {"user_sede": user_sede, "escuelas_feed": escuelas_feed},
    )


@login_required
def estructura_escuela_portal(request, escuela_id):
    """Entrada al asistente de estructura (oferta única por escuela)."""
    if not request.user.is_super_admin():
        r = ca.require_pedagogico(request)
        if r:
            return r

    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, "No tienes una sede asignada. Contacta al administrador.")
        return redirect("core:dashboard")

    escuela = get_object_or_404(Escuela, id=escuela_id, sede=user_sede, is_active=True)
    return redirect("hechos:estructura_configurar", escuela_id=escuela.id, step=1)


@login_required
def estructura_configurar_paso(request, escuela_id, step):
    """
    Asistente: 1) fechas + ciclo → 2) días → 3) horario → 4) profesor → 5) resumen y guardar.
    """
    if step not in (1, 2, 3, 4, 5):
        raise Http404

    if not request.user.is_super_admin():
        r = ca.require_pedagogico(request)
        if r:
            return r

    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, "No tienes una sede asignada. Contacta al administrador.")
        return redirect("core:dashboard")

    escuela = get_object_or_404(
        Escuela.objects.select_related("sede", "maestro__user"),
        id=escuela_id,
        sede=user_sede,
        is_active=True,
    )
    edicion = escuela
    sk = _estructura_wizard_session_key(escuela_id)
    data = dict(request.session.get(sk, {}))

    day_labels = {int(k): v for k, v in EscuelaHorario.DiaSemana.choices}
    total_steps = 5

    wizard_ctx = {
        "user_sede": user_sede,
        "edicion": edicion,
        "escuela": escuela,
        "total_steps": total_steps,
    }

    if step == 1:
        if request.method == "POST":
            fi = _parse_iso_date(request.POST.get("fecha_inicio"))
            ff = _parse_iso_date(request.POST.get("fecha_fin"))
            ciclo = (request.POST.get("ciclo") or "").strip().upper()
            errores = []
            if not fi:
                errores.append("Indica la fecha de inicio del ciclo.")
            if not ff:
                errores.append("Indica la fecha de fin del ciclo.")
            if fi and ff and ff < fi:
                errores.append("La fecha de fin debe ser igual o posterior a la de inicio.")
            if ciclo not in ("A", "B"):
                errores.append("Elige ciclo A o B.")
            anio_ref = fi.year if fi else None
            if not errores and anio_ref is not None:
                dup = Escuela.objects.filter(
                    sede=escuela.sede,
                    nombre__iexact=(escuela.nombre or "").strip(),
                    anio=anio_ref,
                    ciclo=ciclo,
                    is_active=True,
                ).exclude(pk=escuela.pk)
                if dup.exists():
                    errores.append(
                        "Ya existe otra escuela en esta sede con el mismo nombre, año y ciclo. "
                        "Ajusta ciclo o el nombre de la escuela desde coordinación."
                    )

            if errores:
                for e in errores:
                    messages.error(request, e)
                return render(
                    request,
                    "hechos/estructura_wizard_paso1.html",
                    {
                        **wizard_ctx,
                        "step": 1,
                        "w_fecha_inicio": request.POST.get("fecha_inicio") or "",
                        "w_fecha_fin": request.POST.get("fecha_fin") or "",
                        "w_ciclo": ciclo if ciclo in ("A", "B") else "A",
                    },
                )

            escuela.anio = anio_ref
            escuela.ciclo = ciclo
            escuela.fecha_inicio = fi
            escuela.fecha_fin = ff
            escuela.save(
                update_fields=["anio", "ciclo", "fecha_inicio", "fecha_fin", "updated_at"],
            )
            messages.success(request, "Período y ciclo guardados. Sigue con los días de clase.")
            return redirect("hechos:estructura_configurar", escuela_id=escuela_id, step=2)

        return render(
            request,
            "hechos/estructura_wizard_paso1.html",
            {
                **wizard_ctx,
                "step": 1,
                "w_fecha_inicio": (
                    escuela.fecha_inicio.strftime("%Y-%m-%d")
                    if escuela.fecha_inicio
                    else timezone.localdate().strftime("%Y-%m-%d")
                ),
                "w_fecha_fin": (
                    escuela.fecha_fin.strftime("%Y-%m-%d")
                    if escuela.fecha_fin
                    else timezone.localdate().strftime("%Y-%m-%d")
                ),
                "w_ciclo": escuela.ciclo,
            },
        )

    if step == 2:
        existing_days = sorted(
            EscuelaHorario.objects.filter(escuela=escuela).values_list(
                "dia_semana", flat=True
            )
        )
        checked_vals = existing_days
        if request.method == "POST":
            dias_sel = []
            for val, _lbl in EscuelaHorario.DiaSemana.choices:
                if request.POST.get(f"dia_{val}") == "on":
                    dias_sel.append(int(val))
            dias_sel = sorted(set(dias_sel))
            if not dias_sel:
                messages.error(request, "Selecciona al menos un día en que se darán clases.")
                checked_vals = []
            else:
                request.session[sk] = {"dias": dias_sel}
                request.session.modified = True
                return redirect("hechos:estructura_configurar", escuela_id=escuela_id, step=3)
        elif "dias" in data:
            checked_vals = data["dias"]
        day_checks = []
        for val, label in EscuelaHorario.DiaSemana.choices:
            day_checks.append(
                {"value": val, "label": label, "checked": int(val) in checked_vals}
            )
        return render(
            request,
            "hechos/estructura_wizard_paso2.html",
            {
                **wizard_ctx,
                "step": 2,
                "day_checks": day_checks,
            },
        )

    if step == 3:
        dias = data.get("dias")
        if not dias:
            messages.warning(request, "Primero elige los días en que se darán clases.")
            return redirect("hechos:estructura_configurar", escuela_id=escuela_id, step=2)

        dias_sorted = sorted(dias)
        horas_por_dia = data.get("horas_por_dia") or {}
        legacy_hora = data.get("hora")
        existing_by_day = {
            h.dia_semana: h.hora
            for h in EscuelaHorario.objects.filter(escuela=escuela)
        }

        if request.method == "POST":
            nuevas = {}
            errores = []
            for d in dias_sorted:
                try:
                    h12 = int(request.POST.get(f"hora_12_{d}") or 0)
                    mm = int(request.POST.get(f"minuto_{d}") or 0)
                    ampm = request.POST.get(f"ampm_{d}") or ""
                    tt = _time_from_12h(h12, mm, ampm)
                    nuevas[str(d)] = tt.strftime("%H:%M:%S")
                except (ValueError, TypeError):
                    errores.append(
                        f"Horario inválido para {day_labels.get(d, d)} (usa 12 h con a. m. / p. m.)."
                    )
            if errores:
                for e in errores:
                    messages.error(request, e)
                dias_rows = []
                for d in dias_sorted:
                    try:
                        h12 = int(request.POST.get(f"hora_12_{d}") or 7)
                    except (TypeError, ValueError):
                        h12 = 7
                    try:
                        mm = int(request.POST.get(f"minuto_{d}") or 0)
                    except (TypeError, ValueError):
                        mm = 0
                    ampm = (request.POST.get(f"ampm_{d}") or "PM").strip().upper()
                    if ampm not in ("AM", "PM"):
                        ampm = "PM"
                    if mm < 0 or mm > 59:
                        mm = 0
                    if h12 < 1 or h12 > 12:
                        h12 = 7
                    dias_rows.append(
                        {
                            "dia_int": d,
                            "label": day_labels[d],
                            "sel_h12": h12,
                            "sel_mm": mm,
                            "sel_ampm": ampm,
                            "minutos_opts": _minutos_opts_for(mm),
                        }
                    )
                return render(
                    request,
                    "hechos/estructura_wizard_paso3.html",
                    {
                        **wizard_ctx,
                        "step": 3,
                        "dias_rows": dias_rows,
                        "horas_12": list(range(1, 13)),
                    },
                )
            data["horas_por_dia"] = nuevas
            data.pop("hora", None)
            request.session[sk] = data
            request.session.modified = True
            return redirect("hechos:estructura_configurar", escuela_id=escuela_id, step=4)

        dias_rows = _build_dias_horario_rows_paso2(
            dias, day_labels, horas_por_dia, legacy_hora, existing_by_day
        )
        return render(
            request,
            "hechos/estructura_wizard_paso3.html",
            {
                **wizard_ctx,
                "step": 3,
                "dias_rows": dias_rows,
                "horas_12": list(range(1, 13)),
            },
        )

    if step in (4, 5):
        dias = data.get("dias")
        horas_por_dia = dict(data.get("horas_por_dia") or {})
        legacy_hora = data.get("hora")
        if not dias:
            messages.warning(request, "Completa los pasos anteriores (días de clase).")
            return redirect("hechos:estructura_configurar", escuela_id=escuela_id, step=2)

        dias_sorted = sorted(dias)
        if legacy_hora and not horas_por_dia:
            for d in dias_sorted:
                horas_por_dia[str(d)] = legacy_hora

        if any(str(d) not in horas_por_dia for d in dias_sorted):
            messages.warning(request, "Define el horario de cada día en el paso anterior.")
            return redirect("hechos:estructura_configurar", escuela_id=escuela_id, step=3)

        horas_parseadas = []
        resumen_filas = []
        try:
            for d in dias_sorted:
                tt = _parse_session_time(horas_por_dia[str(d)])
                horas_parseadas.append((d, tt))
                resumen_filas.append(
                    {"dia": day_labels[d], "hora_txt": _time_to_display_12h(tt)}
                )
        except (ValueError, TypeError, IndexError):
            messages.error(request, "Sesión de configuración inválida. Vuelve al paso 3.")
            return redirect("hechos:estructura_configurar", escuela_id=escuela_id, step=3)

        fi, ff = escuela.fecha_inicio, escuela.fecha_fin
        if not fi or not ff:
            messages.error(request, "Indica primero las fechas de inicio y fin en el paso 1.")
            return redirect("hechos:estructura_configurar", escuela_id=escuela_id, step=1)
        fechas = _calcular_fechas_clase(fi, ff, set(dias_sorted))
        dias_seleccionados = [day_labels[d] for d in dias_sorted]

        profesores = Profesor.objects.filter(sede=user_sede, is_active=True).select_related(
            "user"
        ).order_by("user__first_name", "user__last_name", "id")

        if step == 4:
            if "profesor_wizard_id" in data:
                selected_profesor_id = data["profesor_wizard_id"]
            else:
                selected_profesor_id = escuela.maestro_id

            if request.method == "POST":
                raw = (request.POST.get("profesor_id") or "").strip()
                if not raw:
                    data["profesor_wizard_id"] = None
                else:
                    try:
                        pid = int(raw)
                    except ValueError:
                        messages.error(request, "Selección de profesor inválida.")
                        return render(
                            request,
                            "hechos/estructura_wizard_paso4.html",
                            {
                                **wizard_ctx,
                                "step": 4,
                                "profesores": profesores,
                                "selected_profesor_id": None,
                            },
                        )
                    prof_ok = Profesor.objects.filter(
                        id=pid, sede=user_sede, is_active=True
                    ).first()
                    if not prof_ok:
                        messages.error(
                            request,
                            "El profesor elegido no pertenece a tu sede o no está activo.",
                        )
                        return render(
                            request,
                            "hechos/estructura_wizard_paso4.html",
                            {
                                **wizard_ctx,
                                "step": 4,
                                "profesores": profesores,
                                "selected_profesor_id": pid,
                            },
                        )
                    data["profesor_wizard_id"] = pid
                request.session[sk] = data
                request.session.modified = True
                return redirect("hechos:estructura_configurar", escuela_id=escuela_id, step=5)

            return render(
                request,
                "hechos/estructura_wizard_paso4.html",
                {
                    **wizard_ctx,
                    "step": 4,
                    "profesores": profesores,
                    "selected_profesor_id": selected_profesor_id,
                },
            )

        if "profesor_wizard_id" not in data:
            messages.warning(request, "Antes elige al profesor en el paso anterior.")
            return redirect("hechos:estructura_configurar", escuela_id=escuela_id, step=4)

        pid_final = data["profesor_wizard_id"]
        prof = None
        if pid_final is not None:
            prof = Profesor.objects.filter(
                id=pid_final, sede=user_sede, is_active=True
            ).select_related("user").first()
            if not prof:
                messages.error(request, "El profesor guardado en la clase ya no es válido. Vuelve a elegirlo.")
                return redirect("hechos:estructura_configurar", escuela_id=escuela_id, step=4)

        if prof:
            profesor_resumen = {
                "nombre": prof.user.get_full_name() or prof.user.username,
                "codigo": prof.codigo_profesor or "",
            }
        else:
            profesor_resumen = {"nombre": "Sin asignar", "codigo": ""}

        if request.method == "POST":
            EscuelaHorario.objects.filter(escuela=escuela).delete()
            for d, tt in horas_parseadas:
                EscuelaHorario.objects.create(escuela=escuela, dia_semana=d, hora=tt)
            escuela.horario = _dias_horarios_to_string(horas_parseadas)
            escuela.maestro = prof
            escuela.save(update_fields=["horario", "maestro", "updated_at"])
            creadas = _crear_clases_si_vacias_desde_estructura(escuela, user_sede, prof)
            request.session.pop(sk, None)
            request.session.modified = True
            if prof:
                msg = "Estructura guardada y profesor asignado a esta edición."
            else:
                msg = (
                    "Estructura guardada. Puedes asignar un profesor más adelante si lo necesitas."
                )
            if creadas:
                msg += f" Se generaron {creadas} clases en el calendario."
            messages.success(request, msg)
            return redirect("hechos:estructura_ediciones")

        ciclo_label = dict(Escuela.Ciclo.choices).get(escuela.ciclo, escuela.ciclo)
        return render(
            request,
            "hechos/estructura_wizard_paso5.html",
            {
                **wizard_ctx,
                "step": 5,
                "dias_seleccionados": dias_seleccionados,
                "resumen_filas": resumen_filas,
                "num_clases": len(fechas),
                "fechas_preview": fechas[:10],
                "profesor_resumen": profesor_resumen,
                "ciclo_label": ciclo_label,
            },
        )


@login_required
def estructura_edicion_legacy_redirect(request, edicion_id):
    """Enlaces antiguos con id de edición: ya no aplica; usar estructura por escuela."""
    if not request.user.is_super_admin():
        r = ca.require_pedagogico(request)
        if r:
            return r
    messages.info(
        request,
        'Las rutas antiguas por «edición» ya no se usan. Abre la escuela desde la lista y configura la estructura.',
    )
    return redirect("hechos:estructura_ediciones")


@login_required
def cursos_list(request):
    """
    Lista de cursos
    """
    # Obtener la sede del usuario
    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada. Contacta al administrador.')
        return redirect('core:dashboard')
    
    cursos = Curso.objects.filter(sede=user_sede, is_active=True).select_related('ruta_estudio')
    
    # Filtrar según el perfil del usuario
    if hasattr(request.user, 'estudiante_profile'):
        # Estudiante ve solo sus cursos matriculados
        matriculas_qs = Matricula.objects.filter(
            estudiante=request.user.estudiante_profile,
            sede=user_sede,
            is_active=True
        ).select_related('escuela', 'escuela__maestro__user')
        escuela_ids = matriculas_qs.values_list('escuela_id', flat=True)
        cursos_ids = Curso.objects.filter(
            ruta_estudio__escuela_id__in=escuela_ids,
            sede=user_sede,
            is_active=True,
        ).values_list('id', flat=True)
        cursos = cursos.filter(id__in=cursos_ids)
        context = {
            'matriculas': matriculas_qs,
            'user_sede': user_sede,
        }
        return render(request, 'hechos/cursos_list_estudiante.html', context)
    elif hasattr(request.user, 'profesor_profile'):
        escuelas_docente = Escuela.objects.filter(
            maestro=request.user.profesor_profile,
            sede=user_sede,
            is_active=True,
        ).values_list('id', flat=True)
        cursos = cursos.filter(ruta_estudio__escuela_id__in=escuelas_docente)
    elif hasattr(request.user, 'admin_escuela_profile') and not request.user.is_super_admin():
        resp = ca.require_academico_o_pedagogico(request)
        if resp:
            return resp
    
    context = {
        'cursos': cursos,
        'user_sede': user_sede,
    }
    
    return render(request, 'hechos/cursos_list_modern.html', context)


@login_required
def escuelas_disponibles(request):
    """
    Lista pública autenticada de escuelas disponibles para estudiantes.
    """
    if not hasattr(request.user, 'estudiante_profile'):
        messages.error(request, 'Solo los estudiantes pueden ver las escuelas disponibles.')
        return redirect('core:dashboard')

    estudiante = request.user.estudiante_profile
    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.info(request, 'Selecciona tu sede para ver escuelas disponibles.')
        return redirect('hechos:seleccionar_sede_estudiante')

    solicitudes = {
        solicitud.escuela_id: solicitud
        for solicitud in SolicitudMatricula.objects.filter(estudiante=estudiante).select_related(
            'escuela', 'escuela__sede', 'escuela__maestro__user'
        )
    }
    matriculas_activas = set(
        Matricula.objects.filter(estudiante=estudiante, is_active=True).values_list('escuela_id', flat=True)
    )

    escuelas_operativas = (
        Escuela.objects.filter(sede=user_sede, is_active=True)
        .select_related('sede', 'maestro__user')
        .order_by('nombre', 'id')
    )

    escuelas_by_id = {}
    for escuela in escuelas_operativas:
        curso = (
            Curso.objects.filter(ruta_estudio__escuela=escuela, is_active=True)
            .order_by('orden', 'id')
            .first()
        )
        escuela_data = escuelas_by_id.setdefault(
            escuela.id,
            {
                'escuela': escuela,
                'sede': escuela.sede,
                'niveles': [],
            },
        )
        solicitud = solicitudes.get(escuela.id)
        escuela_data['niveles'].append({
            'edicion': escuela,
            'curso': curso,
            'ya_matriculado': escuela.id in matriculas_activas,
            'solicitud': solicitud,
            'puede_solicitar': escuela.id not in matriculas_activas and solicitud is None,
        })

    sede_ids = set()
    for escuela_data in escuelas_by_id.values():
        s = escuela_data.get('sede')
        if s is not None and getattr(s, 'id', None):
            sede_ids.add(s.id)
        elif escuela_data['escuela'].sede_id:
            sede_ids.add(escuela_data['escuela'].sede_id)

    ep_lookup = _build_escuela_programa_lookup_por_nombre(sede_ids)

    escuelas_list = []
    for _eid, escuela_data in escuelas_by_id.items():
        escuela = escuela_data['escuela']
        sede_row = escuela_data.get('sede')
        sid = getattr(sede_row, 'id', None) or getattr(escuela, 'sede_id', None)
        ep_escuela = _escuela_programa_para_escuela_operativa(escuela, sid, ep_lookup)
        nivel_programa_linea = (
            ep_escuela.nivel_programa.titulo_acordeon() if ep_escuela else ''
        )

        niveles = escuela_data['niveles']
        ediciones_list = [n['edicion'] for n in niveles]
        ed_pri = edicion_prioritaria_para_matricula(ediciones_list)
        ya_mat = any(n['ya_matriculado'] for n in niveles)
        solicitud_any = None
        for n in niveles:
            s = n['solicitud']
            if not s:
                continue
            if s.estado == 'pendiente':
                solicitud_any = s
                break
            if solicitud_any is None:
                solicitud_any = s
        puede = (
            not ya_mat
            and solicitud_any is None
            and any(e.cupos_disponibles > 0 for e in ediciones_list)
        )

        partes = [
            escuela_data['escuela'].nombre or '',
            getattr(escuela_data.get('sede'), 'nombre', None) or '',
            escuela_data['escuela'].descripcion or '',
        ]
        if nivel_programa_linea:
            partes.append(nivel_programa_linea)
        for n in niveles:
            if n['curso']:
                partes.append(n['curso'].nombre or '')
            if n['edicion'].maestro and n['edicion'].maestro.user:
                partes.append(n['edicion'].maestro.user.get_full_name() or '')

        escuela_data['search_text'] = ' '.join(p for p in partes if p)
        escuela_data['oferta'] = {
            'edicion': ed_pri,
            'curso': next((n['curso'] for n in niveles if n.get('curso')), None),
            'nivel_programa_linea': nivel_programa_linea,
            'cupos_disponibles': sum(e.cupos_disponibles for e in ediciones_list),
            'ya_matriculado': ya_mat,
            'solicitud': solicitud_any,
            'puede_solicitar': puede,
        }
        del escuela_data['niveles']
        escuelas_list.append(escuela_data)

    context = {
        'escuelas': escuelas_list,
        'estudiante': estudiante,
        'user_sede': user_sede,
    }
    return render(request, 'hechos/escuelas_disponibles.html', context)


@login_required
def quejas_reclamos(request):
    if not hasattr(request.user, "estudiante_profile"):
        messages.error(request, "Solo los estudiantes pueden enviar quejas y reclamos.")
        return redirect("core:dashboard")

    estudiante = request.user.estudiante_profile
    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.info(request, "Selecciona tu sede para contactar al coordinador académico.")
        return redirect("hechos:seleccionar_sede_estudiante")

    base_coordinadores = AdminEscuela.objects.filter(
        sede=user_sede,
        is_active=True,
        user__is_active=True,
    )
    coordinador_academico = (
        base_coordinadores.filter(capacidad_asignaciones__capacidad__codigo='academico')
        .select_related('user')
        .distinct()
        .first()
    )
    if not coordinador_academico:
        coordinador_academico = (
            base_coordinadores.filter(tipo_coordinador=AdminEscuela.TipoCoordinador.ACADEMICO)
            .select_related('user')
            .first()
        )

    wa_prefill = (
        f"Hola, soy {request.user.get_full_name() or request.user.email} "
        f"(estudiante, sede {user_sede.nombre}). "
        f"Escribo por el canal de quejas y reclamos en HechosHub."
    )
    coordinador_whatsapp_url = None
    if coordinador_academico and coordinador_academico.user.phone:
        base_wa = _whatsapp_me_url(coordinador_academico.user.phone)
        if base_wa:
            coordinador_whatsapp_url = f"{base_wa}?text={quote(wa_prefill)}"

    if request.method == "POST":
        form = QuejaReclamoForm(request.POST)
        if form.is_valid():
            tipo = form.cleaned_data["asunto"]
            etiqueta_asunto = QuejaReclamo.Tipo(tipo).label
            mensaje = form.cleaned_data["mensaje"]

            if tipo == QuejaReclamo.Tipo.SOLICITUD:
                sol = SolicitudEspecialEstudiante.objects.create(
                    sede=user_sede,
                    estudiante=estudiante,
                    tipo=SolicitudEspecialEstudiante.Tipo.OTRA,
                    asunto=etiqueta_asunto,
                    descripcion=mensaje,
                )
                if coordinador_academico and coordinador_academico.user.email:
                    preview = (mensaje[:70] + "…") if len(mensaje) > 70 else mensaje
                    subject = f"[HechosHub] Solicitud - {user_sede.nombre}: {preview}"
                    body = (
                        f"Sede: {user_sede.nombre}\n"
                        f"Estudiante: {request.user.get_full_name()} ({request.user.email})\n"
                        f"Asunto: {etiqueta_asunto}\n\n"
                        f"Mensaje:\n{mensaje}\n\n"
                        f"ID bandeja: {sol.id}\n"
                    )
                    from_email = getattr(settings, "DEFAULT_FROM_EMAIL", None) or "no-reply@hechoshub.local"
                    email = EmailMessage(
                        subject=subject,
                        body=body,
                        from_email=from_email,
                        to=[coordinador_academico.user.email],
                        reply_to=[request.user.email] if request.user.email else None,
                    )
                    try:
                        email.send(fail_silently=False)
                    except Exception:
                        messages.warning(
                            request,
                            "Se registró tu solicitud, pero no se pudo enviar el correo al coordinador.",
                        )
                    else:
                        messages.success(
                            request,
                            "Tu solicitud fue registrada y se notificó al coordinador académico.",
                        )
                else:
                    messages.success(
                        request,
                        "Tu solicitud fue registrada. El coordinador académico la verá en su bandeja.",
                    )
                return redirect("hechos:quejas_reclamos")

            if not coordinador_academico or not coordinador_academico.user.email:
                messages.error(
                    request,
                    "No hay coordinador académico configurado para tu sede, o no tiene correo. Contacta al administrador.",
                )
                return redirect("hechos:quejas_reclamos")

            qr = QuejaReclamo(
                tipo=tipo,
                asunto=etiqueta_asunto,
                mensaje=mensaje,
                estudiante=estudiante,
                sede=user_sede,
                asignado_a=coordinador_academico,
            )
            qr.save()

            subject = f"[HechosHub] {etiqueta_asunto} · {user_sede.nombre}"
            body = (
                f"Sede: {user_sede.nombre}\n"
                f"Estudiante: {request.user.get_full_name()} ({request.user.email})\n"
                f"Asunto: {etiqueta_asunto}\n\n"
                f"Mensaje:\n{qr.mensaje}\n\n"
                f"ID interno: {qr.id}\n"
            )

            from_email = getattr(settings, "DEFAULT_FROM_EMAIL", None) or "no-reply@hechoshub.local"
            email = EmailMessage(
                subject=subject,
                body=body,
                from_email=from_email,
                to=[coordinador_academico.user.email],
                reply_to=[request.user.email] if request.user.email else None,
            )
            try:
                email.send(fail_silently=False)
            except Exception:
                messages.warning(
                    request,
                    "Se guardó tu mensaje, pero no se pudo enviar el correo. Revisa la configuración de email.",
                )
            else:
                messages.success(request, "Tu mensaje fue enviado al coordinador académico.")

            return redirect("hechos:quejas_reclamos")
    else:
        form = QuejaReclamoForm()

    return render(
        request,
        "hechos/quejas_reclamos_form.html",
        {
            "form": form,
            "user_sede": user_sede,
            "coordinador_academico": coordinador_academico,
            "coordinador_whatsapp_url": coordinador_whatsapp_url,
        },
    )


@login_required
def mis_certificados(request):
    if not hasattr(request.user, "estudiante_profile"):
        messages.error(request, "Solo los estudiantes pueden acceder a esta sección.")
        return redirect("core:dashboard")
    return render(request, "hechos/mis_certificados.html")


def _procesar_solicitud_matricula_estudiante(request, estudiante, escuela: Escuela):
    """
    Crea solicitud de matrícula por escuela operativa (sin duplicar solicitud/matrícula).
    """
    if not escuela or not escuela.sede_id:
        messages.error(request, 'Oferta no válida.')
        return redirect('hechos:escuelas_disponibles')

    if Matricula.objects.filter(
        estudiante=estudiante,
        is_active=True,
        escuela=escuela,
    ).exists():
        messages.info(request, 'Ya tienes una matrícula activa en esta escuela.')
        return redirect('hechos:escuelas_disponibles')

    if SolicitudMatricula.objects.filter(
        estudiante=estudiante,
        escuela=escuela,
    ).exists():
        messages.info(request, 'Ya enviaste una solicitud para esta escuela.')
        return redirect('hechos:escuelas_disponibles')

    if escuela.cupos_disponibles <= 0:
        messages.warning(request, 'Esta escuela no tiene cupos disponibles por ahora.')
        return redirect('hechos:escuelas_disponibles')

    SolicitudMatricula.objects.create(
        sede=escuela.sede,
        estudiante=estudiante,
        escuela=escuela,
    )
    messages.success(request, f'Solicitud enviada para {escuela.nombre}.')
    return redirect('hechos:mis_escuelas')


@login_required
def solicitar_matricula_escuela(request, escuela_id):
    """Matrícula solicitada por escuela (sin exponer id de edición al estudiante)."""
    if request.method != 'POST':
        return redirect('hechos:escuelas_disponibles')

    if not hasattr(request.user, 'estudiante_profile'):
        messages.error(request, 'Solo los estudiantes pueden solicitar matrícula.')
        return redirect('core:dashboard')

    estudiante = request.user.estudiante_profile
    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.info(request, 'Selecciona tu sede para solicitar matrícula.')
        return redirect('hechos:seleccionar_sede_estudiante')

    escuela = get_object_or_404(Escuela, id=escuela_id, sede=user_sede, is_active=True)
    return _procesar_solicitud_matricula_estudiante(request, estudiante, escuela)


@login_required
def solicitar_matricula(request, edicion_id):
    """
    Compatibilidad de URL: el segmento numérico es el id de escuela operativa.
    Preferir ``solicitar_matricula_escuela``.
    """
    if request.method != 'POST':
        return redirect('hechos:escuelas_disponibles')

    if not hasattr(request.user, 'estudiante_profile'):
        messages.error(request, 'Solo los estudiantes pueden solicitar matrícula.')
        return redirect('core:dashboard')

    estudiante = request.user.estudiante_profile
    user_sede = get_user_sede(request.user)
    escuela = get_object_or_404(
        Escuela.objects.select_related('sede', 'maestro__user'),
        id=edicion_id,
        is_active=True,
        sede=user_sede,
    )

    return _procesar_solicitud_matricula_estudiante(request, estudiante, escuela)


@login_required
def mis_escuelas(request):
    """
    Lista de solicitudes de matrícula del estudiante (Mis escuelas).
    """
    if not hasattr(request.user, 'estudiante_profile'):
        messages.error(request, 'Solo los estudiantes pueden ver sus escuelas.')
        return redirect('core:dashboard')

    estudiante = request.user.estudiante_profile
    solicitudes = list(
        SolicitudMatricula.objects.filter(estudiante=estudiante).select_related(
            'escuela',
            'escuela__sede',
            'escuela__maestro__user',
        )
    )
    sede_ids = {s.escuela.sede_id for s in solicitudes if s.escuela_id and s.escuela.sede_id}
    ep_lookup = _build_escuela_programa_lookup_por_nombre(sede_ids)
    for sol in solicitudes:
        esc = sol.escuela
        sid = esc.sede_id if esc else None
        ep = _escuela_programa_para_escuela_operativa(esc, sid, ep_lookup) if esc else None
        sol.nivel_programa_catalogo = ep.nivel_programa.titulo_acordeon() if ep else ''
        sol.curso_portal_id = None
        if esc:
            sol.curso_portal_id = (
                Curso.objects.filter(ruta_estudio__escuela=esc, sede=sol.sede, is_active=True)
                .order_by('orden', 'id')
                .values_list('id', flat=True)
                .first()
            )

    context = {
        'solicitudes': solicitudes,
        'estudiante': estudiante,
    }
    return render(request, 'hechos/mis_escuelas.html', context)


@login_required
def solicitudes_matricula_admin(request):
    if not request.user.is_super_admin():
        r = ca.require_academico(request)
        if r:
            return r

    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada. Contacta al administrador.')
        return redirect('core:dashboard')

    pendientes = SolicitudMatricula.objects.filter(
        sede=user_sede,
        estado='pendiente',
    ).select_related(
        'estudiante__user',
        'escuela',
        'escuela__maestro__user',
    ).order_by('-fecha_solicitud')

    revisadas = SolicitudMatricula.objects.filter(
        sede=user_sede,
    ).exclude(
        estado='pendiente'
    ).select_related(
        'estudiante__user',
        'escuela',
        'revisado_por',
    ).order_by('-fecha_revision', '-updated_at')[:30]

    return render(
        request,
        'hechos/solicitudes_matricula_admin.html',
        {
            'pendientes': pendientes,
            'revisadas': revisadas,
            'user_sede': user_sede,
        },
    )


@login_required
def profesor_solicitudes_matricula(request):
    if not hasattr(request.user, 'profesor_profile'):
        messages.error(request, 'Solo los profesores pueden acceder a esta sección.')
        return redirect('core:dashboard')

    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada. Contacta al administrador.')
        return redirect('core:dashboard')

    prof = request.user.profesor_profile
    base_qs = SolicitudMatricula.objects.filter(
        sede=user_sede,
        escuela__maestro=prof,
    )

    pendientes = (
        base_qs.filter(estado='pendiente')
        .select_related(
            'estudiante__user',
            'escuela',
            'escuela__maestro__user',
        )
        .order_by('-fecha_solicitud')
    )
    revisadas = (
        base_qs.exclude(estado='pendiente')
        .select_related(
            'estudiante__user',
            'escuela',
            'revisado_por',
        )
        .order_by('-fecha_revision', '-updated_at')[:30]
    )

    return render(
        request,
        'hechos/solicitudes_matricula_profesor.html',
        {
            'pendientes': pendientes,
            'revisadas': revisadas,
            'user_sede': user_sede,
        },
    )


@login_required
def revisar_solicitud_matricula(request, solicitud_id):
    if request.method != 'POST':
        if hasattr(request.user, 'profesor_profile'):
            return redirect('hechos:profesor_solicitudes_matricula')
        return redirect('hechos:solicitudes_matricula_admin')

    es_profesor = hasattr(request.user, 'profesor_profile')
    if not request.user.is_super_admin() and not es_profesor:
        r = ca.require_academico(request)
        if r:
            return r

    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada. Contacta al administrador.')
        return redirect('core:dashboard')

    solicitud = get_object_or_404(
        SolicitudMatricula.objects.select_related('estudiante__user', 'escuela__sede'),
        id=solicitud_id,
        sede=user_sede,
    )
    if es_profesor:
        prof = request.user.profesor_profile
        if solicitud.escuela.maestro_id != prof.id:
            messages.error(request, 'No tienes permisos para revisar esta solicitud.')
            return redirect('hechos:profesor_solicitudes_matricula')

    redirect_target = (
        'hechos:profesor_solicitudes_matricula' if es_profesor else 'hechos:solicitudes_matricula_admin'
    )
    if solicitud.estado != 'pendiente':
        messages.info(request, 'Esta solicitud ya fue revisada.')
        return redirect(redirect_target)

    accion = (request.POST.get('accion') or '').strip().lower()
    observaciones = (request.POST.get('observaciones') or '').strip()
    ahora = timezone.now()

    if accion == 'aprobar':
        solicitud.estado = 'aprobada'
        solicitud.observaciones = observaciones
        solicitud.revisado_por = request.user
        solicitud.fecha_revision = ahora
        solicitud.save(update_fields=['estado', 'observaciones', 'revisado_por', 'fecha_revision', 'updated_at'])

        estudiante = solicitud.estudiante
        if estudiante.sede_id != solicitud.sede_id:
            estudiante.sede = solicitud.sede
            estudiante.save(update_fields=['sede'])
        if estudiante.user.sede_id != solicitud.sede_id:
            estudiante.user.sede = solicitud.sede
            estudiante.user.save(update_fields=['sede'])

        Matricula.objects.get_or_create(
            estudiante=estudiante,
            escuela=solicitud.escuela,
            defaults={
                'sede': solicitud.sede,
                'periodo': get_current_period(),
                'estado': 'activa',
                'is_active': True,
            },
        )

        curso_ref = (
            Curso.objects.filter(
                ruta_estudio__escuela_id=solicitud.escuela_id,
                is_active=True,
            )
            .order_by('orden', 'id')
            .first()
        )
        curso_txt = (curso_ref.nombre or '').strip() if curso_ref else ''
        esc_txt = (solicitud.escuela.nombre or '').strip() if solicitud.escuela else ''
        linea_sol = f'{esc_txt} · {curso_txt}' if (esc_txt and curso_txt) else (esc_txt or curso_txt or 'la escuela')
        NotificacionEstudiante.objects.create(
            estudiante=estudiante,
            tipo=NotificacionEstudiante.Tipo.APROBACION,
            titulo='Solicitud aprobada',
            mensaje=(
                f'Tu solicitud para {linea_sol} fue aprobada.'
                + (f" Observación: {observaciones}" if observaciones else "")
            ),
        )
        messages.success(request, 'Solicitud aprobada y matrícula creada.')
    elif accion == 'rechazar':
        solicitud.estado = 'rechazada'
        solicitud.observaciones = observaciones
        solicitud.revisado_por = request.user
        solicitud.fecha_revision = ahora
        solicitud.save(update_fields=['estado', 'observaciones', 'revisado_por', 'fecha_revision', 'updated_at'])

        curso_ref = (
            Curso.objects.filter(
                ruta_estudio__escuela_id=solicitud.escuela_id,
                is_active=True,
            )
            .order_by('orden', 'id')
            .first()
        )
        curso_txt = (curso_ref.nombre or '').strip() if curso_ref else ''
        esc_txt = (solicitud.escuela.nombre or '').strip() if solicitud.escuela else ''
        linea_sol = f'{esc_txt} · {curso_txt}' if (esc_txt and curso_txt) else (esc_txt or curso_txt or 'la escuela')
        NotificacionEstudiante.objects.create(
            estudiante=solicitud.estudiante,
            tipo=NotificacionEstudiante.Tipo.RECHAZO,
            titulo='Solicitud rechazada',
            mensaje=(
                f'Tu solicitud para {linea_sol} fue rechazada.'
                + (f" Motivo: {observaciones}" if observaciones else "")
            ),
        )
        messages.warning(request, 'Solicitud rechazada.')
    else:
        messages.error(request, 'Acción no válida.')

    return redirect(redirect_target)


@login_required
def solicitudes_especiales_estudiante(request):
    """Compatibilidad: la gestión vive en la página unificada de quejas y reclamos."""
    return redirect("hechos:quejas_reclamos")


@login_required
def solicitudes_especiales_admin(request):
    if not request.user.is_super_admin():
        r = ca.require_academico(request)
        if r:
            return r

    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada. Contacta al administrador.')
        return redirect('core:dashboard')

    pendientes = SolicitudEspecialEstudiante.objects.filter(
        sede=user_sede,
        estado='pendiente',
    ).select_related('estudiante__user').order_by('-fecha_solicitud')
    revisadas = SolicitudEspecialEstudiante.objects.filter(
        sede=user_sede,
    ).exclude(estado='pendiente').select_related('estudiante__user', 'revisado_por').order_by('-fecha_revision', '-updated_at')[:30]
    total_bandeja = SolicitudEspecialEstudiante.objects.filter(sede=user_sede).count()

    return render(
        request,
        'hechos/solicitudes_especiales_admin.html',
        {'pendientes': pendientes, 'revisadas': revisadas, 'user_sede': user_sede, 'total_bandeja': total_bandeja},
    )


@login_required
def revisar_solicitud_especial_admin(request, solicitud_id):
    if request.method != 'POST':
        return redirect('hechos:solicitudes_especiales_admin')

    if not request.user.is_super_admin():
        r = ca.require_academico(request)
        if r:
            return r

    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada. Contacta al administrador.')
        return redirect('core:dashboard')

    solicitud = get_object_or_404(
        SolicitudEspecialEstudiante.objects.select_related('estudiante__user'),
        id=solicitud_id,
        sede=user_sede,
    )
    if solicitud.estado != SolicitudEspecialEstudiante.Estado.PENDIENTE:
        messages.info(request, 'Esta solicitud ya fue revisada.')
        return redirect('hechos:solicitudes_especiales_admin')

    accion = (request.POST.get('accion') or '').strip().lower()
    observaciones = (request.POST.get('observaciones') or '').strip()
    if accion == 'atender':
        solicitud.estado = SolicitudEspecialEstudiante.Estado.ATENDIDA
    elif accion == 'cerrar':
        solicitud.estado = SolicitudEspecialEstudiante.Estado.CERRADA
    else:
        messages.error(request, 'Acción no válida.')
        return redirect('hechos:solicitudes_especiales_admin')

    solicitud.observaciones = observaciones
    solicitud.revisado_por = request.user
    solicitud.fecha_revision = timezone.now()
    solicitud.save(update_fields=['estado', 'observaciones', 'revisado_por', 'fecha_revision', 'updated_at'])

    NotificacionEstudiante.objects.create(
        estudiante=solicitud.estudiante,
        tipo=NotificacionEstudiante.Tipo.INFO,
        titulo='Actualización de tu solicitud',
        mensaje=(
            f"Tu solicitud «{solicitud.asunto}» fue marcada como {solicitud.get_estado_display().lower()}."
            + (f" Comentario: {observaciones}" if observaciones else "")
        ),
    )
    messages.success(request, 'Solicitud actualizada.')
    return redirect('hechos:solicitudes_especiales_admin')


@login_required
def rutas_estudio_list(request):
    """
    Lista de rutas de estudio (solo para admins)
    """
    if not request.user.is_super_admin():
        r = ca.require_academico_o_pedagogico(request)
        if r:
            return r
    
    # Obtener la sede del usuario
    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada. Contacta al administrador.')
        return redirect('core:dashboard')
    
    rutas = RutaEstudio.objects.filter(sede=user_sede, is_active=True).prefetch_related('cursos')
    
    # Agregar información de cursos para cada ruta
    rutas_con_cursos = []
    for ruta in rutas:
        cursos = ruta.cursos.filter(is_active=True).order_by('orden', 'nombre')
        rutas_con_cursos.append({
            'ruta': ruta,
            'cursos': cursos,
            'total_cursos': cursos.count(),
            'duracion_total': sum(curso.duracion_semanas for curso in cursos)
        })
    
    context = {
        'rutas_con_cursos': rutas_con_cursos,
        'user_sede': user_sede,
    }
    
    return render(request, 'hechos/rutas_estudio_list.html', context)


@login_required
def ruta_estudio_detail(request, ruta_id):
    """
    Detalle de una ruta de estudio específica
    """
    if not request.user.is_super_admin():
        r = ca.require_academico_o_pedagogico(request)
        if r:
            return r
    
    # Obtener la sede del usuario
    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada. Contacta al administrador.')
        return redirect('core:dashboard')
    
    ruta = get_object_or_404(RutaEstudio, id=ruta_id, sede=user_sede, is_active=True)
    cursos = ruta.cursos.filter(is_active=True).order_by('orden', 'nombre')
    
    # Estadísticas de la ruta
    total_estudiantes = 0
    for curso in cursos:
        total_estudiantes += curso.total_estudiantes_inscritos
    
    # Obtener profesores disponibles para asignar a nuevos cursos
    profesores = Profesor.objects.filter(sede=user_sede, is_active=True).select_related('user')
    
    context = {
        'ruta': ruta,
        'cursos': cursos,
        'total_cursos': cursos.count(),
        'total_estudiantes': total_estudiantes,
        'duracion_total': sum(curso.duracion_semanas for curso in cursos),
        'profesores': profesores,
        'user_sede': user_sede,
    }
    
    return render(request, 'hechos/ruta_estudio_detail.html', context)


@login_required
def crear_curso_ruta(request, ruta_id):
    """
    Crear nuevo curso para una ruta específica
    """
    if not request.user.is_super_admin():
        r = ca.require_academico_o_pedagogico(request)
        if r:
            return r
    
    # Obtener la sede del usuario
    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada. Contacta al administrador.')
        return redirect('core:dashboard')
    
    ruta = get_object_or_404(RutaEstudio, id=ruta_id, sede=user_sede, is_active=True)
    profesores = Profesor.objects.filter(sede=user_sede, is_active=True).select_related('user')
    
    if request.method == 'POST':
        try:
            profesor_id = request.POST.get('profesor')
            profesor = Profesor.objects.get(id=profesor_id, sede=user_sede)
            
            curso = Curso.objects.create(
                nombre=request.POST.get('nombre'),
                descripcion=request.POST.get('descripcion', ''),
                ruta_estudio=ruta,
                orden=request.POST.get('orden', 1) or 1,
                duracion_semanas=request.POST.get('duracion_semanas', 4) or 4,
                sede=user_sede
            )
            
            messages.success(request, f'Curso "{request.POST.get("nombre")}" creado exitosamente en la ruta {ruta.nombre}.')
            return redirect('hechos:ruta_estudio_detail', ruta_id=ruta_id)
            
        except Exception as e:
            messages.error(request, f'Error al crear curso: {str(e)}')
    
    # Obtener el siguiente orden disponible
    ultimo_orden = ruta.cursos.filter(is_active=True).aggregate(
        max_orden=models.Max('orden')
    )['max_orden'] or 0
    siguiente_orden = ultimo_orden + 1
    
    context = {
        'ruta': ruta,
        'profesores': profesores,
        'siguiente_orden': siguiente_orden,
        'user_sede': user_sede,
    }
    
    return render(request, 'hechos/crear_curso_ruta.html', context)


@login_required
def profesor_bandeja_entregas(request):
    """
    Entregas de actividades con evidencia pendientes de calificar (vista tipo bandeja).
    """
    if not hasattr(request.user, 'profesor_profile'):
        messages.error(request, 'Solo los profesores pueden acceder.')
        return redirect('core:dashboard')
    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada.')
        return redirect('core:dashboard')
    entregas = list(entregas_pendientes_calificacion_qs(request.user, user_sede))
    return render(
        request,
        'hechos/profesor_bandeja_entregas.html',
        {
            'entregas': entregas,
            'user_sede': user_sede,
        },
    )


@login_required
def clases_list(request):
    """
    Lista de clases
    """
    # Obtener la sede del usuario
    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada. Contacta al administrador.')
        return redirect('core:dashboard')

    if (
        request.resolver_match.url_name == 'clases_list'
        and hasattr(request.user, 'profesor_profile')
        and not hasattr(request.user, 'estudiante_profile')
        and not request.user.is_super_admin()
    ):
        target = reverse('hechos:mis_escuelas_profesor')
        q = request.GET.urlencode()
        if q:
            target = f'{target}?{q}'
        return redirect(target)

    if hasattr(request.user, 'admin_escuela_profile') and not request.user.is_super_admin():
        if not ca.admin_may_access_clases_list(request.user):
            messages.error(request, 'No tienes permisos para ver las clases.')
            return redirect('core:dashboard')

    es_mis_escuelas_docente = request.resolver_match.url_name == 'mis_escuelas_profesor'
    profesor = getattr(request.user, 'profesor_profile', None)

    if es_mis_escuelas_docente:
        if not profesor:
            return redirect('hechos:clases_list')
        escuelas_asignadas = list(
            Escuela.objects.filter(
                sede=user_sede,
                maestro=profesor,
                is_active=True,
            )
            .select_related('maestro__user')
            .prefetch_related('estructura_horarios')
            .order_by('-anio', 'ciclo', 'nombre')
        )
        escuela_pk = request.GET.get('escuela')
        if escuela_pk:
            escuela_filtro = get_object_or_404(
                Escuela,
                id=escuela_pk,
                sede=user_sede,
                maestro=profesor,
                is_active=True,
            )
            num_estudiantes_escuela = (
                Matricula.objects.filter(
                    sede=user_sede,
                    is_active=True,
                    escuela=escuela_filtro,
                )
                .values('estudiante_id')
                .distinct()
                .count()
            )
            return render(
                request,
                'hechos/clases_list_modern.html',
                {
                    'user_sede': user_sede,
                    'docente_mis_escuelas_modo': 'clases_por_escuela',
                    'escuelas_asignadas': escuelas_asignadas,
                    'escuela_filtro': escuela_filtro,
                    'clases': [],
                    'num_estudiantes_escuela': num_estudiantes_escuela,
                    'clases_completadas': [],
                    'clases_en_progreso': [],
                    'promedio_asistencia': None,
                },
            )

        escuelas_feed = []
        for escuela in escuelas_asignadas:
            n_est = (
                Matricula.objects.filter(
                    sede=user_sede,
                    is_active=True,
                    escuela=escuela,
                )
                .values('estudiante_id')
                .distinct()
                .count()
            )
            horario_ui = _horario_edicion_a_texto_ui(escuela)
            aula_ui = (escuela.aula or "").strip()
            if not aula_ui:
                salon_asignado = (
                    Salon.objects.filter(
                        sede=user_sede,
                        escuela=escuela,
                        is_active=True,
                    )
                    .order_by('nombre')
                    .first()
                )
                if salon_asignado:
                    aula_ui = salon_asignado.nombre
            escuelas_feed.append(
                {
                    'escuela': escuela,
                    'num_estudiantes': n_est,
                    'horario_ui': horario_ui,
                    'aula_ui': aula_ui,
                }
            )
        return render(
            request,
            'hechos/clases_list_modern.html',
            {
                'user_sede': user_sede,
                'docente_mis_escuelas_modo': 'escuelas',
                'escuelas_asignadas': escuelas_asignadas,
                'escuelas_feed': escuelas_feed,
                'clases': [],
                'clases_completadas': [],
                'clases_en_progreso': [],
                'promedio_asistencia': None,
            },
        )

    clases = Clase.objects.filter(sede=user_sede, is_active=True).select_related(
        'escuela',
        'profesor__user',
    ).prefetch_related('asistencias')

    # Filtrar según el perfil del usuario
    if hasattr(request.user, 'profesor_profile'):
        # Profesor ve solo sus clases
        clases = clases.filter(profesor=request.user.profesor_profile)
    elif hasattr(request.user, 'estudiante_profile'):
        # Estudiante ve clases de las ediciones en las que está matriculado
        edicion_ids = Matricula.objects.filter(
            estudiante=request.user.estudiante_profile,
            sede=user_sede,
            is_active=True,
        ).values_list('escuela_id', flat=True)
        clases = clases.filter(escuela_id__in=edicion_ids)

    clases = list(clases.order_by('fecha_clase', 'escuela__nombre', 'numero_clase'))
    for c in clases:
        regs = list(c.asistencias.all())
        total = len(regs)
        presentes = sum(1 for a in regs if a.estado == 'presente')
        c.porcentaje_asistencia_ui = round(100 * presentes / total, 1) if total else None
        c.n_registros_asistencia = total

    context = {
        'clases': clases,
        'user_sede': user_sede,
        'docente_mis_escuelas_modo': '',
        'clases_completadas': [],
        'clases_en_progreso': [],
        'promedio_asistencia': None,
    }

    return render(request, 'hechos/clases_list_modern.html', context)


def _profesor_escuela_docente(request, escuela_id):
    """Escuela activa del maestro en esta sede, o (None, redirect response)."""
    if not hasattr(request.user, 'profesor_profile'):
        messages.error(request, 'Solo los profesores pueden acceder.')
        return None, redirect('core:dashboard')
    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada.')
        return None, redirect('core:dashboard')
    prof = request.user.profesor_profile
    escuela = get_object_or_404(
        Escuela,
        id=escuela_id,
        sede=user_sede,
        maestro=prof,
        is_active=True,
    )
    return (user_sede, escuela, prof), None


def _estudiante_matriculado_en_escuela(estudiante, escuela, user_sede):
    if not escuela or not user_sede:
        return False
    return Matricula.objects.filter(
        estudiante=estudiante,
        sede=user_sede,
        is_active=True,
        escuela=escuela,
    ).exists()


def _porcentajes_evaluacion_por_estudiante_escuela(estudiante_ids, user_sede, escuela):
    """
    Porcentajes por estudiante: notas de curso (modelo Nota) y entregas de actividad
    calificadas (EntregaActividad), siempre en el contexto de esta escuela y sede.
    """
    pct_by_est = defaultdict(list)
    if not estudiante_ids:
        return pct_by_est

    for row in (
        Nota.objects.filter(
            sede=user_sede,
            curso__ruta_estudio__escuela=escuela,
            puntaje_maximo__gt=0,
            estudiante_id__in=estudiante_ids,
        ).values('estudiante_id', 'puntaje_obtenido', 'puntaje_maximo')
    ):
        po = float(row['puntaje_obtenido'])
        pm = float(row['puntaje_maximo'])
        pct_by_est[row['estudiante_id']].append(po * 100.0 / pm)

    entregas = EntregaActividad.objects.filter(
        sede=user_sede,
        actividad__escuela=escuela,
        actividad__is_active=True,
        estudiante_id__in=estudiante_ids,
        puntaje_asignado__isnull=False,
    ).select_related('actividad')
    for ent in entregas:
        max_pts = ent.actividad.puntos_posibles
        if max_pts is None:
            continue
        fm = float(max_pts)
        if fm <= 0 or ent.puntaje_asignado is None:
            continue
        pct_by_est[ent.estudiante_id].append(float(ent.puntaje_asignado) * 100.0 / fm)

    return pct_by_est


@login_required
def profesor_escuela_estudiantes(request, escuela_id):
    """Listado de estudiantes matriculados en cursos de la escuela (docente asignado)."""
    pack, redir = _profesor_escuela_docente(request, escuela_id)
    if redir:
        return redir
    user_sede, escuela, _prof = pack
    estudiante_ids = list(
        Matricula.objects.filter(
            sede=user_sede,
            is_active=True,
            escuela=escuela,
        )
        .values_list('estudiante_id', flat=True)
        .distinct()
    )
    estudiantes = (
        Estudiante.objects.filter(id__in=estudiante_ids)
        .select_related('user')
        .order_by('-id')
    )
    pct_by_est = _porcentajes_evaluacion_por_estudiante_escuela(
        estudiante_ids, user_sede, escuela
    )
    filas = []
    for e in estudiantes:
        pcts = pct_by_est.get(e.id) or []
        filas.append(
            {
                'estudiante': e,
                'promedio_pct': (sum(pcts) / len(pcts)) if pcts else None,
                'n_calificaciones': len(pcts),
            }
        )
    volver_url = f"{reverse('hechos:mis_escuelas_profesor')}?escuela={escuela.id}"
    return render(
        request,
        'hechos/profesor_escuela_estudiantes.html',
        {
            'escuela': escuela,
            'user_sede': user_sede,
            'filas': filas,
            'volver_url': volver_url,
        },
    )


@login_required
@require_POST
def profesor_escuela_estudiante_quitar_matricula(request, escuela_id, estudiante_id):
    """
    Anula la matrícula del estudiante en todos los cursos de esta escuela (misma sede).
    Útil si hubo error al digitar documento o nombre: libera cupo y permite volver a matricular bien.
    No borra el usuario ni el registro de estudiante en el sistema.
    """
    pack, redir = _profesor_escuela_docente(request, escuela_id)
    if redir:
        return redir
    user_sede, escuela, _prof = pack

    redirect_to = (request.POST.get('redirect_to') or 'estudiantes').strip()
    if redirect_to not in ('estudiantes', 'asistencia', 'notas'):
        redirect_to = 'estudiantes'

    mats = list(
        Matricula.objects.filter(
            estudiante_id=estudiante_id,
            sede=user_sede,
            is_active=True,
            escuela=escuela,
        ).select_related('estudiante__user', 'escuela')
    )
    if not mats:
        messages.warning(
            request,
            'Ese estudiante no tiene una matrícula activa en esta escuela.',
        )
        if redirect_to == 'asistencia':
            return redirect('hechos:profesor_escuela_asistencia_panel', escuela_id=escuela.id)
        if redirect_to == 'notas':
            return redirect('hechos:profesor_escuela_notas_panel', escuela_id=escuela.id)
        return redirect('hechos:profesor_escuela_estudiantes', escuela_id=escuela.id)

    est = mats[0].estudiante
    nombre = est.user.get_full_name() or est.user.email or str(est.user_id)

    try:
        with transaction.atomic():
            for m in mats:
                m.is_active = False
                m.estado = 'cancelada'
                m.save(update_fields=['is_active', 'estado', 'updated_at'])
            AsistenciaGrillaCelda.objects.filter(
                escuela=escuela,
                sede=user_sede,
                estudiante_id=estudiante_id,
            ).delete()
            NotasGrillaCelda.objects.filter(
                escuela=escuela,
                sede=user_sede,
                estudiante_id=estudiante_id,
            ).delete()
            NotasGrillaEstadoFinal.objects.filter(
                escuela=escuela,
                estudiante_id=estudiante_id,
            ).delete()
    except Exception as exc:
        messages.error(request, f'No se pudo quitar la matrícula: {exc}')
        if redirect_to == 'asistencia':
            return redirect('hechos:profesor_escuela_asistencia_panel', escuela_id=escuela.id)
        if redirect_to == 'notas':
            return redirect('hechos:profesor_escuela_notas_panel', escuela_id=escuela.id)
        return redirect('hechos:profesor_escuela_estudiantes', escuela_id=escuela.id)

    n = len(mats)
    if n == 1:
        messages.success(
            request,
            f'Se anuló la matrícula de {nombre} en esta escuela. Puede volver a matricular al estudiante con los datos correctos.',
        )
    else:
        messages.success(
            request,
            f'Se anularon {n} matrículas activas de {nombre} en esta escuela. Puede volver a matricular al estudiante con los datos correctos.',
        )

    if redirect_to == 'asistencia':
        return redirect('hechos:profesor_escuela_asistencia_panel', escuela_id=escuela.id)
    if redirect_to == 'notas':
        return redirect('hechos:profesor_escuela_notas_panel', escuela_id=escuela.id)
    return redirect('hechos:profesor_escuela_estudiantes', escuela_id=escuela.id)


@login_required
def profesor_escuela_estudiante_notas(request, escuela_id, estudiante_id):
    """Notas del estudiante en cursos que pertenecen a esta escuela."""
    pack, redir = _profesor_escuela_docente(request, escuela_id)
    if redir:
        return redir
    user_sede, escuela, _prof = pack
    inscrito = Matricula.objects.filter(
        estudiante_id=estudiante_id,
        sede=user_sede,
        is_active=True,
        escuela=escuela,
    ).exists()
    if not inscrito:
        raise Http404
    estudiante = get_object_or_404(Estudiante, id=estudiante_id)
    notas = (
        Nota.objects.filter(
            estudiante=estudiante,
            sede=user_sede,
            curso__ruta_estudio__escuela=escuela,
        )
        .select_related('curso', 'profesor__user')
        .order_by('-fecha_evaluacion', 'curso__nombre')
    )
    entregas_calif = (
        EntregaActividad.objects.filter(
            estudiante=estudiante,
            sede=user_sede,
            actividad__escuela=escuela,
            actividad__is_active=True,
            puntaje_asignado__isnull=False,
        )
        .select_related('actividad')
        .order_by('-evaluado_en', '-actualizado_en')
    )
    evaluaciones = []
    for n in notas:
        pct = float(n.porcentaje) if n.puntaje_maximo and n.puntaje_maximo > 0 else None
        evaluaciones.append(
            {
                'kind': 'nota',
                'fecha': n.fecha_evaluacion,
                'curso_o_ctx': n.curso.nombre,
                'titulo': n.titulo,
                'tipo_label': n.get_tipo_display(),
                'puntaje_celda': f'{n.puntaje_obtenido} / {n.puntaje_maximo}',
                'porcentaje': pct,
                'nota': n,
            }
        )
    for ent in entregas_calif:
        max_pts = ent.actividad.puntos_posibles
        pct = None
        if max_pts is not None and float(max_pts) > 0 and ent.puntaje_asignado is not None:
            pct = float(ent.puntaje_asignado) * 100.0 / float(max_pts)
        fecha_ev = ent.evaluado_en.date() if ent.evaluado_en else None
        if max_pts is not None:
            puntaje_celda = f'{ent.puntaje_asignado} / {max_pts}'
        else:
            puntaje_celda = str(ent.puntaje_asignado)
        evaluaciones.append(
            {
                'kind': 'entrega',
                'fecha': fecha_ev,
                'curso_o_ctx': 'Actividad de escuela',
                'titulo': ent.actividad.titulo,
                'tipo_label': 'Entrega calificada',
                'puntaje_celda': puntaje_celda,
                'porcentaje': pct,
                'entrega': ent,
            }
        )

    def _sort_key(ev):
        fd = ev['fecha']
        if isinstance(fd, datetime):
            fd = fd.date()
        return fd or date.min

    evaluaciones.sort(key=_sort_key, reverse=True)

    pcts_prom = [e['porcentaje'] for e in evaluaciones if e['porcentaje'] is not None]
    promedio_pct = (sum(pcts_prom) / len(pcts_prom)) if pcts_prom else None
    volver_lista = reverse('hechos:profesor_escuela_estudiantes', kwargs={'escuela_id': escuela.id})
    volver_escuela = f"{reverse('hechos:mis_escuelas_profesor')}?escuela={escuela.id}"
    return render(
        request,
        'hechos/profesor_escuela_estudiante_notas.html',
        {
            'escuela': escuela,
            'estudiante': estudiante,
            'evaluaciones': evaluaciones,
            'promedio_pct': promedio_pct,
            'user_sede': user_sede,
            'volver_lista_url': volver_lista,
            'volver_escuela_url': volver_escuela,
        },
    )


GRILLA_ASISTENCIA_COLS_MAX = 60
GRILLA_ASISTENCIA_COLS_DEFECTO = 0

_GRILLA_VALORES_PERMITIDOS = frozenset(
    {
        AsistenciaGrillaCelda.Valor.ASISTIO,
        AsistenciaGrillaCelda.Valor.ESTADO_P,
        AsistenciaGrillaCelda.Valor.INASISTENCIA,
    }
)


def _clases_ordenadas_para_grilla_asistencia(user_sede, escuela):
    return list(
        Clase.objects.filter(
            sede=user_sede,
            is_active=True,
            escuela__is_active=True,
            escuela=escuela,
        )
        .order_by("fecha_clase", "id")[:GRILLA_ASISTENCIA_COLS_MAX]
    )


def _edicion_principal_escuela(user_sede, escuela):
    """Compat nombre: devuelve la escuela operativa activa (ya no existe edición aparte)."""
    if not escuela or escuela.sede_id != getattr(user_sede, 'id', None):
        return None
    if not escuela.is_active:
        return None
    return escuela


def _grilla_asistencia_config_asegurada(escuela, user_sede) -> EscuelaGrillaAsistenciaConfig:
    cfg, _ = EscuelaGrillaAsistenciaConfig.objects.get_or_create(
        escuela=escuela,
        sede=user_sede,
        defaults={'num_columnas': GRILLA_ASISTENCIA_COLS_DEFECTO},
    )
    n_clases = Clase.objects.filter(
        sede=user_sede,
        is_active=True,
        escuela__is_active=True,
        escuela=escuela,
    ).count()
    target = min(n_clases, GRILLA_ASISTENCIA_COLS_MAX)
    if target != cfg.num_columnas:
        cfg.num_columnas = target
        cfg.save(update_fields=['num_columnas', 'updated_at'])
    return cfg


def _grilla_num_clases_calendario_escuela(user_sede, escuela) -> int:
    return Clase.objects.filter(
        sede=user_sede,
        is_active=True,
        escuela__is_active=True,
        escuela=escuela,
    ).count()


def _fecha_sugerida_nueva_sesion_grilla(escuela: Escuela | None, clases_cols: list) -> date:
    """
    Próxima fecha de clase según EscuelaHorario dentro del período; si ya no cabe
    otra fecha programada, la sugerencia es última sesión + 7 días (ritmo semanal típico),
    aunque pase fecha_fin — no se sustituye por fecha_fin si ese día no es de clase.
    """
    hoy = timezone.localdate()
    if not escuela:
        return hoy
    ini = escuela.fecha_inicio
    fin = escuela.fecha_fin
    dias_set = set(
        EscuelaHorario.objects.filter(escuela=escuela).values_list(
            "dia_semana", flat=True
        )
    )
    if dias_set:
        candidatas = _calcular_fechas_clase(ini, fin, dias_set)
        if clases_cols:
            ultima = _grilla_fecha_local_clase(clases_cols[-1].fecha_clase)
            for d in sorted(candidatas):
                if d > ultima:
                    return d
            return ultima + timedelta(days=7)
        start = max(ini, hoy)
        for d in sorted(candidatas):
            if d >= start:
                return d
        candidata = max(ini, hoy)
    else:
        if clases_cols:
            return _grilla_fecha_local_clase(clases_cols[-1].fecha_clase) + timedelta(
                days=7
            )
        candidata = max(ini, hoy)
    if candidata < ini:
        candidata = ini
    if candidata > fin:
        candidata = fin
    return candidata


def _grilla_nueva_clase_en_fecha(escuela: Escuela, sede, profesor, d: date) -> Clase:
    """Crea una Clase en la fecha indicada (hora del primer día configurado o 19:00)."""
    tz = timezone.get_current_timezone()
    horarios = list(
        EscuelaHorario.objects.filter(escuela=escuela).order_by('dia_semana', 'hora')
    )
    hh = horarios[0].hora if horarios else dt_time(19, 0)
    mx = (
        Clase.objects.filter(escuela=escuela).aggregate(m=Max('numero_clase'))['m']
        or 0
    )
    naive = datetime.combine(d, hh)
    aware = timezone.make_aware(naive, tz)
    prof_asignado = profesor or escuela.maestro
    return Clase.objects.create(
        sede=sede,
        escuela=escuela,
        numero_clase=mx + 1,
        titulo=f'Sesión {mx + 1}',
        descripcion='',
        fecha_clase=aware,
        duracion_minutos=90,
        profesor=prof_asignado,
        aula='',
        is_active=True,
    )


def _grilla_doc_normalizado(raw) -> str:
    if raw is None:
        return ''
    if isinstance(raw, float) and not isinstance(raw, bool):
        if raw == int(raw):
            raw = str(int(raw))
        else:
            raw = str(raw)
    elif isinstance(raw, int):
        raw = str(raw)
    else:
        raw = str(raw).strip()
    return ''.join(c for c in raw if c.isdigit())


def _grilla_fecha_local_clase(dt) -> date:
    if dt is None:
        return None
    if isinstance(dt, date) and not isinstance(dt, datetime):
        return dt
    if timezone.is_naive(dt):
        dt = timezone.make_aware(dt, timezone.get_current_timezone())
    return timezone.localtime(dt).date()


def _grilla_excel_celda_a_fecha(cell_data, cell_formula, prev_date: date | None):
    """Interpreta encabezado de sesión: datetime, serial Excel o fórmula +7 como en plantilla exportada."""
    v = cell_data.value if cell_data is not None else None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        try:
            from openpyxl.utils.datetime import from_excel

            dtx = from_excel(v)
            if isinstance(dtx, datetime):
                return dtx.date()
            return dtx
        except Exception:
            pass
    fv = cell_formula.value if cell_formula is not None else None
    if isinstance(fv, str) and fv.startswith('=') and prev_date is not None:
        if re.search(r'\+[\s]*7\s*\)?\s*$', fv.replace('\n', '')):
            return prev_date + timedelta(days=7)
    return None


def _grilla_excel_es_inicio_cola_totales(ws, header_row: int, col: int) -> bool:
    a = ws.cell(row=header_row, column=col).value
    b = ws.cell(row=header_row, column=col + 1).value
    c2 = ws.cell(row=header_row, column=col + 2).value
    sa = (str(a).strip().upper() if a is not None else '')
    sb = (str(b).strip().upper() if b is not None else '')
    sc = (str(c2).strip().upper() if c2 is not None else '')
    return sa == 'A' and sb == 'N' and sc == 'P'


def _grilla_excel_columnas_sesion(ws_data, ws_formula, header_row: int = 6):
    """Lista (columna_excel 1-based, fecha local) hasta la cola A/N/P."""
    prev = None
    out = []
    max_col = min(
        ws_data.max_column or 0,
        ws_formula.max_column or 0,
        4 + GRILLA_ASISTENCIA_COLS_MAX + 8,
    )
    c = 5
    while c <= max_col:
        if _grilla_excel_es_inicio_cola_totales(ws_data, header_row, c):
            break
        cell_d = ws_data.cell(row=header_row, column=c)
        cell_f = ws_formula.cell(row=header_row, column=c)
        d = _grilla_excel_celda_a_fecha(cell_d, cell_f, prev)
        if d is None:
            break
        out.append((c, d))
        prev = d
        c += 1
    return out


def _grilla_reindex_celdas_tras_cambio_clases(escuela, sede, old_clases: list, new_clases: list):
    """Reasigna columna 1..n según orden actual de Clase (por id estable)."""
    if old_clases == new_clases:
        return
    new_id_to_col = {c.id: i + 1 for i, c in enumerate(new_clases)}
    temp_base = 1_000_000
    final_by_pk = {}
    for cel in AsistenciaGrillaCelda.objects.filter(escuela=escuela, sede=sede):
        oc = cel.columna
        if oc < 1 or oc > len(old_clases):
            cel.delete()
            continue
        oid = old_clases[oc - 1].id
        nc = new_id_to_col.get(oid)
        if nc is None:
            cel.delete()
            continue
        if nc != oc:
            final_by_pk[cel.pk] = nc
            cel.columna = temp_base + cel.pk
            cel.save(update_fields=['columna', 'updated_at'])
    for pk, nc in final_by_pk.items():
        AsistenciaGrillaCelda.objects.filter(pk=pk).update(columna=nc)


def _grilla_ensure_clases_para_fechas_excel(
    escuela: Escuela,
    sede,
    profesor,
    fechas_en_orden: list,
) -> int:
    """
    Crea tantas Clase como haga falta para que cada fecha del Excel (incl. repetidas)
    tenga una clase distinta en la escuela. Devuelve cantidad creada.
    """
    if not fechas_en_orden or not escuela:
        return 0
    counts = Counter(fechas_en_orden)
    clases = list(
        Clase.objects.filter(
            sede=sede,
            escuela=escuela,
            is_active=True,
        ).order_by('fecha_clase', 'id')
    )
    by_date = defaultdict(deque)
    for cl in clases:
        by_date[_grilla_fecha_local_clase(cl.fecha_clase)].append(cl)
    creadas = 0
    for d, need in counts.items():
        have = len(by_date.get(d, ()))
        while have < need:
            if _grilla_num_clases_calendario_escuela(sede, escuela) >= GRILLA_ASISTENCIA_COLS_MAX:
                break
            nueva = _grilla_nueva_clase_en_fecha(escuela, sede, profesor, d)
            by_date[d].append(nueva)
            have += 1
            creadas += 1
    return creadas


def _grilla_map_excel_col_a_columna_grilla(new_clases: list, session_cols: list):
    """
    session_cols: [(col_excel, fecha), ...] en orden de izquierda a derecha.
    Asigna cada columna Excel a columna grilla 1..n usando orden de clases con misma fecha.
    """
    by_date = defaultdict(deque)
    for i, c in enumerate(new_clases):
        by_date[_grilla_fecha_local_clase(c.fecha_clase)].append(i + 1)
    out = {}
    for _cx, d in session_cols:
        q = by_date.get(d)
        if not q:
            continue
        out[_cx] = q.popleft()
    return out


def _grilla_valor_excel_a_choice(raw) -> str | None:
    if raw is None or raw == '':
        return None
    if isinstance(raw, float) and not isinstance(raw, bool):
        if raw == int(raw):
            raw = str(int(raw))
        else:
            raw = str(raw)
    s = str(raw).strip().upper()[:1]
    if s in _GRILLA_VALORES_PERMITIDOS:
        return s
    return None


@login_required
def profesor_escuela_asistencia_import_excel(request, escuela_id):
    """Importa A/N/P desde Excel: celdas con valor actualizan; vacío no cambia; fechas nuevas crean Clase."""
    if request.method != 'POST':
        return redirect('hechos:profesor_escuela_asistencia_panel', escuela_id=escuela_id)
    pack, redir = _profesor_escuela_docente(request, escuela_id)
    if redir:
        return redir
    user_sede, escuela, prof = pack
    f = request.FILES.get('archivo_excel')
    if not f:
        messages.error(request, 'Selecciona un archivo Excel (.xlsx).')
        return redirect('hechos:profesor_escuela_asistencia_panel', escuela_id=escuela.id)
    if f.size > 6 * 1024 * 1024:
        messages.error(request, 'El archivo es demasiado grande (máx. 6 MB).')
        return redirect('hechos:profesor_escuela_asistencia_panel', escuela_id=escuela.id)
    head = f.read(4)
    f.seek(0)
    if head != b'PK\x03\x04' and head != b'PK\x05\x06':
        messages.error(request, 'El archivo no parece un .xlsx válido.')
        return redirect('hechos:profesor_escuela_asistencia_panel', escuela_id=escuela.id)

    try:
        from openpyxl import load_workbook
    except Exception:
        messages.error(request, 'Falta openpyxl para importar Excel.')
        return redirect('hechos:profesor_escuela_asistencia_panel', escuela_id=escuela.id)

    estudiantes = list(
        Estudiante.objects.filter(
            sede=user_sede,
            matriculas__sede=user_sede,
            matriculas__is_active=True,
            matriculas__escuela=escuela,
        )
        .select_related('user')
        .distinct()
        .order_by('user__last_name', 'user__first_name', 'user__email')
    )
    doc_a_est = {}
    for est in estudiantes:
        k = _grilla_doc_normalizado(est.numero_documento)
        if k:
            doc_a_est[k] = est
    edicion = _edicion_principal_escuela(user_sede, escuela)

    try:
        data = f.read()
        from io import BytesIO

        wb_data = load_workbook(BytesIO(data), data_only=True)
        wb_formula = load_workbook(BytesIO(data), data_only=False)
        ws_d = wb_data.active
        ws_f = wb_formula.active
    except Exception as exc:
        messages.error(request, f'No se pudo leer el Excel: {exc}')
        return redirect('hechos:profesor_escuela_asistencia_panel', escuela_id=escuela.id)

    header_row = 6
    h1 = ws_d.cell(row=header_row, column=1).value
    if str(h1 or '').strip().upper() not in ('NO', 'Nº', 'N°', '#'):
        header_row = None
        for r in range(1, 12):
            v = ws_d.cell(row=r, column=1).value
            if str(v or '').strip().upper() == 'NO':
                header_row = r
                break
        if header_row is None:
            messages.error(
                request,
                'No se encontró la fila de encabezados (esperada columna «NO» en la primera columna).',
            )
            return redirect('hechos:profesor_escuela_asistencia_panel', escuela_id=escuela.id)

    session_cols = _grilla_excel_columnas_sesion(ws_d, ws_f, header_row=header_row)
    if not session_cols:
        messages.error(
            request,
            'No se encontraron columnas de fechas en el Excel (después de «Celular», antes de A/N/P).',
        )
        return redirect('hechos:profesor_escuela_asistencia_panel', escuela_id=escuela.id)
    if len(session_cols) > GRILLA_ASISTENCIA_COLS_MAX:
        messages.error(
            request,
            f'El Excel tiene demasiadas sesiones (máx. {GRILLA_ASISTENCIA_COLS_MAX}).',
        )
        return redirect('hechos:profesor_escuela_asistencia_panel', escuela_id=escuela.id)

    if not edicion:
        messages.error(
            request,
            'La escuela no está disponible; no se pueden crear sesiones nuevas desde el Excel.',
        )
        return redirect('hechos:profesor_escuela_asistencia_panel', escuela_id=escuela.id)

    fechas_en_orden = [d for _c, d in session_cols]
    try:
        with transaction.atomic():
            old_clases = list(_clases_ordenadas_para_grilla_asistencia(user_sede, escuela))
            creadas = _grilla_ensure_clases_para_fechas_excel(
                edicion, user_sede, prof, fechas_en_orden
            )
            new_clases = list(_clases_ordenadas_para_grilla_asistencia(user_sede, escuela))
            if len(new_clases) > GRILLA_ASISTENCIA_COLS_MAX:
                raise ValueError(
                    f'Se superaría el máximo de {GRILLA_ASISTENCIA_COLS_MAX} sesiones en la grilla.'
                )
            _grilla_reindex_celdas_tras_cambio_clases(escuela, user_sede, old_clases, new_clases)
            col_excel_a_grilla = _grilla_map_excel_col_a_columna_grilla(new_clases, session_cols)
            if len(col_excel_a_grilla) < len(session_cols):
                raise ValueError(
                    'No se pudieron enlazar todas las columnas de fechas del Excel con sesiones del sistema.'
                )
            _grilla_asistencia_config_asegurada(escuela, user_sede)

            actualizados = 0
            sin_doc = 0
            max_r = ws_d.max_row or 0
            for r in range(header_row + 1, max_r + 1):
                raw_doc = ws_d.cell(row=r, column=2).value
                doc_key = _grilla_doc_normalizado(raw_doc)
                if not doc_key:
                    continue
                est = doc_a_est.get(doc_key)
                if not est:
                    sin_doc += 1
                    continue
                for cx, _fecha in session_cols:
                    grilla_col = col_excel_a_grilla.get(cx)
                    if grilla_col is None:
                        continue
                    val = _grilla_valor_excel_a_choice(ws_d.cell(row=r, column=cx).value)
                    if val is None:
                        continue
                    obj, created = AsistenciaGrillaCelda.objects.get_or_create(
                        escuela=escuela,
                        sede=user_sede,
                        estudiante_id=est.id,
                        columna=grilla_col,
                        defaults={
                            'valor': val,
                            'registrado_por': request.user,
                        },
                    )
                    if created:
                        actualizados += 1
                    else:
                        if obj.valor != val:
                            obj.valor = val
                            obj.registrado_por = request.user
                            obj.save(update_fields=['valor', 'registrado_por', 'updated_at'])
                            actualizados += 1
    except ValueError as e:
        messages.error(request, str(e))
        return redirect('hechos:profesor_escuela_asistencia_panel', escuela_id=escuela.id)
    except Exception as exc:
        messages.error(request, f'Error al importar: {exc}')
        return redirect('hechos:profesor_escuela_asistencia_panel', escuela_id=escuela.id)

    parts = [
        f'Se importaron datos desde el Excel ({actualizados} celdas actualizadas o creadas).',
    ]
    if creadas:
        parts.append(f'Se añadieron {creadas} sesión(es) nueva(s) al calendario según las fechas del archivo.')
    if sin_doc:
        parts.append(
            f'{sin_doc} fila(s) con documento no reconocido entre los matriculados se omitieron.'
        )
    messages.success(request, ' '.join(parts))
    return redirect('hechos:profesor_escuela_asistencia_panel', escuela_id=escuela.id)


@login_required
def profesor_escuela_asistencia_agregar_sesion(request, escuela_id):
    """POST: crea una Clase en la fecha elegida y reordena columnas de la grilla."""
    if request.method != 'POST':
        return redirect('hechos:profesor_escuela_asistencia_panel', escuela_id=escuela_id)
    pack, redir = _profesor_escuela_docente(request, escuela_id)
    if redir:
        return redir
    user_sede, escuela, prof = pack
    edicion = _edicion_principal_escuela(user_sede, escuela)
    if not edicion:
        messages.error(
            request,
            'La escuela no está disponible; no se puede crear una sesión.',
        )
        return redirect('hechos:profesor_escuela_asistencia_panel', escuela_id=escuela.id)
    if _grilla_num_clases_calendario_escuela(user_sede, escuela) >= GRILLA_ASISTENCIA_COLS_MAX:
        messages.error(
            request,
            f'Ya se alcanzó el máximo de {GRILLA_ASISTENCIA_COLS_MAX} sesiones en la grilla.',
        )
        return redirect('hechos:profesor_escuela_asistencia_panel', escuela_id=escuela.id)
    raw = (request.POST.get('fecha_sesion') or '').strip()
    d = _parse_fecha_sesion_texto(raw)
    if not d:
        messages.error(request, 'Indica una fecha válida (día/mes/año, como en la tabla).')
        return redirect('hechos:profesor_escuela_asistencia_panel', escuela_id=escuela.id)
    try:
        with transaction.atomic():
            old_clases = list(_clases_ordenadas_para_grilla_asistencia(user_sede, escuela))
            _grilla_nueva_clase_en_fecha(edicion, user_sede, prof, d)
            new_clases = list(_clases_ordenadas_para_grilla_asistencia(user_sede, escuela))
            _grilla_reindex_celdas_tras_cambio_clases(escuela, user_sede, old_clases, new_clases)
            _grilla_asistencia_config_asegurada(escuela, user_sede)
    except Exception as exc:
        messages.error(request, f'No se pudo crear la sesión: {exc}')
        return redirect('hechos:profesor_escuela_asistencia_panel', escuela_id=escuela.id)
    messages.success(
        request,
        f'Sesión añadida para el {d.strftime("%d/%m/%Y")}. La columna aparece en la grilla y en el Excel al descargarlo.',
    )
    return redirect('hechos:profesor_escuela_asistencia_panel', escuela_id=escuela.id)


@login_required
def profesor_escuela_asistencia_quitar_sesion(request, escuela_id):
    """POST: desactiva la Clase, borra celdas de esa columna y reindexa la grilla."""
    if request.method != 'POST':
        return redirect('hechos:profesor_escuela_asistencia_panel', escuela_id=escuela_id)
    pack, redir = _profesor_escuela_docente(request, escuela_id)
    if redir:
        return redir
    user_sede, escuela, _prof = pack
    raw_id = (request.POST.get('clase_id') or '').strip()
    try:
        clase_id = int(raw_id)
    except ValueError:
        messages.error(request, 'Solicitud no válida.')
        return redirect('hechos:profesor_escuela_asistencia_panel', escuela_id=escuela.id)
    base_qs = Clase.objects.filter(
        sede=user_sede,
        is_active=True,
        escuela__is_active=True,
        escuela=escuela,
    )
    clase = get_object_or_404(base_qs, pk=clase_id)
    fecha_txt = timezone.localtime(clase.fecha_clase).strftime('%d/%m/%Y')
    try:
        with transaction.atomic():
            old_clases = list(_clases_ordenadas_para_grilla_asistencia(user_sede, escuela))
            old_col = None
            for i, c in enumerate(old_clases):
                if c.id == clase.id:
                    old_col = i + 1
                    break
            if old_col is None:
                raise ValueError('La sesión no está en la grilla actual.')
            AsistenciaGrillaCelda.objects.filter(
                escuela=escuela,
                sede=user_sede,
                columna=old_col,
            ).delete()
            clase.is_active = False
            clase.save(update_fields=['is_active', 'updated_at'])
            new_clases = list(_clases_ordenadas_para_grilla_asistencia(user_sede, escuela))
            _grilla_reindex_celdas_tras_cambio_clases(escuela, user_sede, old_clases, new_clases)
            _grilla_asistencia_config_asegurada(escuela, user_sede)
    except ValueError as e:
        messages.error(request, str(e))
        return redirect('hechos:profesor_escuela_asistencia_panel', escuela_id=escuela.id)
    except Exception as exc:
        messages.error(request, f'No se pudo quitar la sesión: {exc}')
        return redirect('hechos:profesor_escuela_asistencia_panel', escuela_id=escuela.id)
    messages.success(
        request,
        f'Se eliminó la sesión del {fecha_txt} y la asistencia guardada en esa columna.',
    )
    return redirect('hechos:profesor_escuela_asistencia_panel', escuela_id=escuela.id)


@login_required
def profesor_escuela_asistencia_panel(request, escuela_id):
    """
    Una columna por cada clase (solo fecha en encabezado). Letras como el Excel: A verde, P amarillo, N rojo;
    % inasistencia = N / número de sesiones.
    """
    pack, redir = _profesor_escuela_docente(request, escuela_id)
    if redir:
        return redir
    user_sede, escuela, prof = pack

    estudiantes = list(
        Estudiante.objects.filter(
            sede=user_sede,
            matriculas__sede=user_sede,
            matriculas__is_active=True,
            matriculas__escuela=escuela,
        )
        .select_related('user')
        .distinct()
        .order_by('user__last_name', 'user__first_name', 'user__email')
    )
    est_ids = [e.id for e in estudiantes]
    cfg = _grilla_asistencia_config_asegurada(escuela, user_sede)
    num_columnas = cfg.num_columnas

    valores: dict[tuple[int, int], str] = {}
    if est_ids:
        for cel in AsistenciaGrillaCelda.objects.filter(
            escuela=escuela,
            sede=user_sede,
            estudiante_id__in=est_ids,
            columna__lte=num_columnas,
        ).only('estudiante_id', 'columna', 'valor'):
            valores[(cel.estudiante_id, cel.columna)] = cel.valor

    if request.method == 'POST':
        est_set = frozenset(est_ids)
        try:
            with transaction.atomic():
                for key, raw_val in request.POST.items():
                    if not key.startswith('g_'):
                        continue
                    parts = key.split('_')
                    if len(parts) != 3:
                        continue
                    try:
                        est_id = int(parts[1])
                        col = int(parts[2])
                    except ValueError:
                        continue
                    if (
                        est_id not in est_set
                        or num_columnas < 1
                        or col < 1
                        or col > num_columnas
                    ):
                        continue
                    val = (raw_val or '').strip().upper()[:1]
                    if not val:
                        AsistenciaGrillaCelda.objects.filter(
                            escuela=escuela,
                            sede=user_sede,
                            estudiante_id=est_id,
                            columna=col,
                        ).delete()
                        continue
                    if val not in _GRILLA_VALORES_PERMITIDOS:
                        continue
                    obj, created = AsistenciaGrillaCelda.objects.get_or_create(
                        escuela=escuela,
                        sede=user_sede,
                        estudiante_id=est_id,
                        columna=col,
                        defaults={
                            'valor': val,
                            'registrado_por': request.user,
                        },
                    )
                    if not created:
                        obj.valor = val
                        obj.registrado_por = request.user
                        obj.save(update_fields=['valor', 'registrado_por', 'updated_at'])
        except Exception as exc:
            messages.error(request, f'No se pudo guardar: {exc}')
            return redirect('hechos:profesor_escuela_asistencia_panel', escuela_id=escuela.id)

        messages.success(request, 'Asistencia guardada.')
        return redirect('hechos:profesor_escuela_asistencia_panel', escuela_id=escuela.id)

    clases_cols = _clases_ordenadas_para_grilla_asistencia(user_sede, escuela)
    edicion = _edicion_principal_escuela(user_sede, escuela)
    docente_user = None
    if edicion and edicion.maestro_id:
        docente_user = edicion.maestro.user
    elif prof:
        docente_user = prof.user
    nombre_docente = (
        (docente_user.get_full_name() or docente_user.email or docente_user.username) if docente_user else '—'
    )
    horario_txt = _horario_edicion_a_texto_ui(edicion) if edicion else '—'
    modalidad_txt = _modalidad_edicion_label(edicion)

    encabezados_columnas = []
    for i in range(1, num_columnas + 1):
        fecha_txt = '—'
        fecha_etiqueta = ''
        clase_id = None
        if i <= len(clases_cols):
            cl = clases_cols[i - 1]
            clase_id = cl.id
            fc = timezone.localtime(cl.fecha_clase)
            fecha_txt = fc.strftime('%d/%m/%y')
            fecha_etiqueta = fc.strftime('%d/%m/%Y')
        encabezados_columnas.append(
            {
                'fecha': fecha_txt,
                'clase_id': clase_id,
                'fecha_etiqueta': fecha_etiqueta,
            }
        )

    filas = []
    for est in estudiantes:
        celdas = []
        cnt_a = cnt_p = cnt_n = 0
        for col in range(1, num_columnas + 1):
            v = valores.get((est.id, col), '')
            if v == AsistenciaGrillaCelda.Valor.ASISTIO:
                cnt_a += 1
            elif v == AsistenciaGrillaCelda.Valor.ESTADO_P:
                cnt_p += 1
            elif v == AsistenciaGrillaCelda.Valor.INASISTENCIA:
                cnt_n += 1
            celdas.append({'col': col, 'value': v})
        denom = num_columnas if num_columnas else 0
        pct_inas = (cnt_n / denom) if denom else 0.0
        filas.append(
            {
                'estudiante': est,
                'celdas': celdas,
                'cnt_a': cnt_a,
                'cnt_n': cnt_n,
                'cnt_p': cnt_p,
                'pct_inasistencia': pct_inas,
            }
        )

    volver_url = f"{reverse('hechos:mis_escuelas_profesor')}?escuela={escuela.id}"
    n_cal = _grilla_num_clases_calendario_escuela(user_sede, escuela)
    if not edicion:
        agregar_sesion_deshabilitado_motivo = 'La escuela no está disponible para añadir sesiones.'
    elif n_cal >= GRILLA_ASISTENCIA_COLS_MAX:
        agregar_sesion_deshabilitado_motivo = (
            f'Ya hay {GRILLA_ASISTENCIA_COLS_MAX} sesiones (límite de la grilla).'
        )
    else:
        agregar_sesion_deshabilitado_motivo = ''
    _sug_fecha = (
        _fecha_sugerida_nueva_sesion_grilla(edicion, clases_cols)
        if edicion
        else timezone.localdate()
    )
    fecha_sesion_input_default_display = _sug_fecha.strftime('%d/%m/%y')
    ret_asistencia = reverse('hechos:profesor_escuela_asistencia_panel', args=[escuela.id])
    matricular_modal_activo = bool(
        prof
        and _profesor_puede_gestionar_escuela(prof, escuela)
        and escuela.cupos_disponibles > 0
    )
    return render(
        request,
        'hechos/profesor_escuela_asistencia_panel.html',
        {
            'escuela': escuela,
            'user_sede': user_sede,
            'volver_url': volver_url,
            'encabezados_columnas': encabezados_columnas,
            'filas': filas,
            'num_columnas': num_columnas,
            'tiene_estudiantes': bool(estudiantes),
            'tiene_columnas_clase': num_columnas > 0,
            'nombre_docente': nombre_docente,
            'horario_txt': horario_txt,
            'modalidad_txt': modalidad_txt,
            'agregar_sesion_url': reverse(
                'hechos:profesor_escuela_asistencia_agregar_sesion',
                args=[escuela.id],
            ),
            'agregar_sesion_deshabilitado_motivo': agregar_sesion_deshabilitado_motivo,
            'fecha_sesion_input_default_display': fecha_sesion_input_default_display,
            'quitar_sesion_url': reverse(
                'hechos:profesor_escuela_asistencia_quitar_sesion',
                args=[escuela.id],
            ),
            'matricular_modal_escuela': escuela,
            'matricular_modal_activo': matricular_modal_activo,
            'matricular_modal_return_url': ret_asistencia,
            'matricular_modal_post_url': reverse('hechos:matricular_estudiante'),
            'matricular_modal_sugerencias_url': reverse(
                'hechos:sugerencias_estudiante_doc_matricula'
            ),
        },
    )


def _profesor_escuela_asistencia_excel_response(request, escuela_id, *, plantilla_vacia: bool):
    """
    Excel de grilla de asistencia. Si plantilla_vacia=True, mismas columnas y estudiantes
    pero celdas de sesión y totales (A/N/P/%) vacías.
    """
    pack, redir = _profesor_escuela_docente(request, escuela_id)
    if redir:
        return redir
    user_sede, escuela, prof = pack
    estudiantes = list(
        Estudiante.objects.filter(
            sede=user_sede,
            matriculas__sede=user_sede,
            matriculas__is_active=True,
            matriculas__escuela=escuela,
        )
        .select_related('user')
        .distinct()
        .order_by('user__last_name', 'user__first_name', 'user__email')
    )
    cfg = _grilla_asistencia_config_asegurada(escuela, user_sede)
    num_columnas = cfg.num_columnas
    clases_cols = _clases_ordenadas_para_grilla_asistencia(user_sede, escuela)
    edicion = _edicion_principal_escuela(user_sede, escuela)
    docente_user = None
    if edicion and edicion.maestro_id:
        docente_user = edicion.maestro.user
    elif prof:
        docente_user = prof.user
    nombre_docente = (
        (docente_user.get_full_name() or docente_user.email or docente_user.username)
        if docente_user
        else '—'
    )
    texto_escuela = f'{escuela.nombre} — {escuela.anio} · {escuela.get_ciclo_display()}'
    horario_txt = _horario_edicion_a_texto_ui(edicion) if edicion else '—'
    modalidad_txt = _modalidad_edicion_label(edicion)

    est_ids = [e.id for e in estudiantes]
    valores_grid: dict[tuple[int, int], str] = {}
    if not plantilla_vacia and est_ids:
        for cel in AsistenciaGrillaCelda.objects.filter(
            escuela=escuela,
            sede=user_sede,
            estudiante_id__in=est_ids,
            columna__lte=num_columnas,
        ):
            valores_grid[(cel.estudiante_id, cel.columna)] = cel.valor

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side, Protection
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation
        from openpyxl.formatting.rule import FormulaRule
        from openpyxl.worksheet.datavalidation import DataValidation
    except Exception:
        messages.error(request, "Falta la dependencia para generar Excel. Instala openpyxl.")
        return redirect('hechos:profesor_escuela_asistencia_panel', escuela_id=escuela.id)

    wb = Workbook()
    ws = wb.active
    ws.title = "Formato asistencia" if plantilla_vacia else "Asistencia"

    fixed_headers = ["NO", "Documento", "Nombres y Apellidos", "Celular"]
    attendance_cols = num_columnas
    tail_headers = ["A", "N", "P", "% Inasistencia"]
    total_cols = len(fixed_headers) + attendance_cols + len(tail_headers)
    base_cols = len(fixed_headers)
    sess_start_col = base_cols + 1
    sess_end_col = base_cols + attendance_cols
    first_tail = base_cols + attendance_cols + 1

    header_fill_dates = PatternFill("solid", fgColor="D9E1F2")
    header_fill_fixed_no = PatternFill("solid", fgColor="D9D9D9")
    header_fill_fixed_ids = PatternFill("solid", fgColor="FFFFFF")
    header_font_fixed_no = Font(color="000000", bold=True)
    header_font_fixed_ids = Font(color="434343", bold=True)
    header_font_dates = Font(color="000000", bold=True)
    header_fill_tail = [
        PatternFill("solid", fgColor="B6D7A8"),
        PatternFill("solid", fgColor="DA8585"),
        PatternFill("solid", fgColor="FBBC04"),
        PatternFill("solid", fgColor="073763"),
    ]
    header_font_tail = Font(color="FFFFFF", bold=True)
    thin_side = Side(style="thin", color="B4B4B4")
    border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")
    meta_label_fill = PatternFill("solid", fgColor="0B5394")
    meta_label_font = Font(color="FFFFFF", bold=True)
    meta_val_fill = PatternFill("solid", fgColor="FFFFFF")
    data_font = Font(color="000000")

    fill_asistio = PatternFill("solid", fgColor="C6EFCE")
    fill_p = PatternFill("solid", fgColor="FFEB9C")
    fill_inas = PatternFill("solid", fgColor="FFC7CE")

    def _meta_celda(r, c, val, es_etiqueta=False):
        cell = ws.cell(row=r, column=c, value=val)
        cell.border = border
        if es_etiqueta:
            cell.fill = meta_label_fill
            cell.font = meta_label_font
        else:
            cell.fill = meta_val_fill
            cell.font = data_font
        cell.alignment = left

    _meta_celda(1, 2, "Docente", es_etiqueta=True)
    _meta_celda(1, 3, nombre_docente)
    _meta_celda(2, 2, "Curso", es_etiqueta=True)
    _meta_celda(2, 3, texto_escuela)
    _meta_celda(3, 2, "Horario", es_etiqueta=True)
    _meta_celda(3, 3, horario_txt)
    _meta_celda(4, 2, "Modalidad", es_etiqueta=True)
    _meta_celda(4, 3, modalidad_txt)

    header_row = 6
    for col in range(1, total_cols + 1):
        c = ws.cell(row=header_row, column=col)
        if col <= base_cols:
            c.value = fixed_headers[col - 1]
            if col == 1:
                c.fill = header_fill_fixed_no
                c.font = header_font_fixed_no
            else:
                c.fill = header_fill_fixed_ids
                c.font = header_font_fixed_ids
        elif col <= base_cols + attendance_cols:
            ci = col - base_cols - 1
            if ci < len(clases_cols):
                fc = timezone.localtime(clases_cols[ci].fecha_clase)
                c.value = fc.replace(tzinfo=None)
                c.number_format = "dd/mm/yyyy"
            else:
                c.value = ""
            c.fill = header_fill_dates
            c.font = header_font_dates
        else:
            ti = col - base_cols - attendance_cols - 1
            c.value = tail_headers[ti]
            c.fill = header_fill_tail[ti]
            c.font = header_font_tail
        c.border = border
        c.alignment = center

    row = header_row + 1
    ls = get_column_letter(sess_start_col) if attendance_cols else None
    le = get_column_letter(sess_end_col) if attendance_cols else None
    la = get_column_letter(first_tail) if attendance_cols else None
    ln = get_column_letter(first_tail + 1) if attendance_cols else None
    lp = get_column_letter(first_tail + 2) if attendance_cols else None

    for idx, est in enumerate(estudiantes, start=1):
        ws.cell(row=row, column=1, value=idx)
        doc = (est.numero_documento or "").strip()
        if doc.isdigit():
            try:
                ws.cell(row=row, column=2, value=int(doc))
            except ValueError:
                ws.cell(row=row, column=2, value=doc)
        else:
            ws.cell(row=row, column=2, value=doc or "")
        ws.cell(row=row, column=3, value=est.user.get_full_name() or est.user.email or "—")
        cel_txt = (est.user.phone or "").strip() or (est.telefono_emergencia or "").strip()
        ws.cell(row=row, column=4, value=cel_txt)
        if plantilla_vacia:
            for ci in range(attendance_cols):
                col_excel = base_cols + 1 + ci
                ce = ws.cell(row=row, column=col_excel, value=None)
                ce.font = data_font
                ce.border = border
                ce.alignment = center
        else:
            for ci in range(attendance_cols):
                col_excel = base_cols + 1 + ci
                v = valores_grid.get((est.id, ci + 1), "")
                if v == AsistenciaGrillaCelda.Valor.ASISTIO:
                    cell_val = "A"
                elif v == AsistenciaGrillaCelda.Valor.ESTADO_P:
                    cell_val = "P"
                elif v == AsistenciaGrillaCelda.Valor.INASISTENCIA:
                    cell_val = "N"
                else:
                    cell_val = ""
                ce = ws.cell(row=row, column=col_excel, value=cell_val)
                ce.font = data_font
                ce.border = border
                ce.alignment = center
                if cell_val == "A":
                    ce.fill = fill_asistio
                elif cell_val == "P":
                    ce.fill = fill_p
                elif cell_val == "N":
                    ce.fill = fill_inas

        if attendance_cols and ls and le and la and ln and lp:
            ws.cell(row=row, column=first_tail, value=f"=COUNTIF({ls}{row}:{le}{row},{la}$6)")
            ws.cell(row=row, column=first_tail + 1, value=f"=COUNTIF({ls}{row}:{le}{row},{ln}$6)")
            ws.cell(row=row, column=first_tail + 2, value=f"=COUNTIF({ls}{row}:{le}{row},{lp}$6)")
            pct_cell = ws.cell(
                row=row,
                column=first_tail + 3,
                value=f"=IF(COUNTA({ls}$6:{le}$6)=0,0,{ln}{row}/COUNTA({ls}$6:{le}$6))",
            )
            pct_cell.number_format = "0.00%"
        else:
            ws.cell(row=row, column=first_tail, value=0)
            ws.cell(row=row, column=first_tail + 1, value=0)
            ws.cell(row=row, column=first_tail + 2, value=0)
            zpct = ws.cell(row=row, column=first_tail + 3, value=0)
            zpct.number_format = "0.00%"

        for j in range(4):
            tc = ws.cell(row=row, column=first_tail + j)
            tc.font = data_font
            tc.border = border
            tc.alignment = center
        row += 1

    first_data_row = header_row + 1
    data_last_row = row - 1
    if attendance_cols and ls and le:
        dv_end_row = max(data_last_row, first_data_row) + 25
        dv = DataValidation(
            type="list",
            formula1='"A,N,P"',
            allow_blank=True,
            showErrorMessage=True,
        )
        dv.error = "Use A, N o P."
        dv.errorTitle = "Valor no permitido"
        ws.add_data_validation(dv)
        dv.add(f"{ls}{first_data_row}:{le}{dv_end_row}")
        ws.freeze_panes = f"{ls}{first_data_row}"
    else:
        ws.freeze_panes = f"{get_column_letter(first_tail)}{first_data_row}"

    for r in range(header_row, row):
        for c in range(1, total_cols + 1):
            cell = ws.cell(row=r, column=c)
            if r > header_row and base_cols < c <= base_cols + attendance_cols:
                continue
            cell.border = border
            if r == header_row:
                continue
            cell.alignment = left if c in (2, 3) else center
            if r > header_row and c <= 4:
                cell.font = data_font

    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 34
    ws.column_dimensions["D"].width = 14
    for col_idx in range(base_cols + 1, total_cols + 1):
        letter = ws.cell(row=header_row, column=col_idx).column_letter
        if col_idx <= base_cols + attendance_cols:
            ws.column_dimensions[letter].width = 11
        elif col_idx < total_cols:
            ws.column_dimensions[letter].width = 6
        else:
            ws.column_dimensions[letter].width = 14

    nombre_corto = (slugify(escuela.nombre) or "escuela")[:24].strip("-")
    fecha = timezone.now().strftime("%Y%m%d")
    if plantilla_vacia:
        filename = f"formato-asis-e{escuela.id}-{nombre_corto}-{fecha}.xlsx"
    else:
        filename = f"asis-e{escuela.id}-{nombre_corto}-{fecha}.xlsx"
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


@login_required
def profesor_escuela_asistencia_excel(request, escuela_id):
    """Exporta Excel con asistencia guardada (valores A/P/N y totales)."""
    return _profesor_escuela_asistencia_excel_response(request, escuela_id, plantilla_vacia=False)


@login_required
def profesor_escuela_asistencia_excel_formato(request, escuela_id):
    """Plantilla Excel: mismos estudiantes y fechas; celdas de asistencia y totales vacías."""
    return _profesor_escuela_asistencia_excel_response(request, escuela_id, plantilla_vacia=True)


GRILLA_NOTAS_COLS_DEFECTO = 5
GRILLA_NOTAS_COLS_MAX = 40


def _grilla_notas_valor_guardar(raw) -> str:
    if raw is None:
        return ''
    if isinstance(raw, float) and not isinstance(raw, bool):
        if math.isnan(raw):
            return ''
        s = str(raw).strip()
    elif isinstance(raw, int):
        s = str(raw)
    else:
        s = str(raw).strip()
    return s[:64]


def _grilla_notas_parse_num_promedio(s: str):
    if not s or not str(s).strip():
        return None
    t = str(s).strip().replace(',', '.')
    try:
        return float(t)
    except ValueError:
        return None


def _grilla_notas_promedio_display(valores_por_col: dict[int, str], num_columnas: int) -> str:
    nums = []
    for c in range(1, num_columnas + 1):
        v = valores_por_col.get(c) or ''
        n = _grilla_notas_parse_num_promedio(v)
        if n is not None:
            nums.append(n)
    if not nums:
        return ''
    avg = sum(nums) / len(nums)
    return f'{avg:.4f}'.replace('.', ',')


def _grilla_notas_config_asegurada(escuela, user_sede) -> EscuelaGrillaNotasConfig:
    cfg, _ = EscuelaGrillaNotasConfig.objects.get_or_create(
        escuela=escuela,
        sede=user_sede,
        defaults={'num_columnas': GRILLA_NOTAS_COLS_DEFECTO},
    )
    if cfg.num_columnas < 1:
        cfg.num_columnas = 1
        cfg.save(update_fields=['num_columnas', 'updated_at'])
    return cfg


def _grilla_notas_estudiantes_qs(user_sede, escuela):
    return (
        Estudiante.objects.filter(
            sede=user_sede,
            matriculas__sede=user_sede,
            matriculas__is_active=True,
            matriculas__escuela=escuela,
        )
        .select_related('user')
        .distinct()
        .order_by('user__last_name', 'user__first_name', 'user__email')
    )


def _grilla_notas_excel_find_header_row(ws, max_scan=15):
    for r in range(1, max_scan + 1):
        v = ws.cell(row=r, column=1).value
        if str(v or '').strip().upper() in ('NO', 'Nº', 'N°', '#'):
            return r
    return 6


def _grilla_notas_import_mapeo_columnas(ws, header_row: int, cfg: EscuelaGrillaNotasConfig):
    """Devuelve nota_cols: {excel_col: nota_index 1..n}, estado_col, promedio_col."""
    max_c = min(ws.max_column or 0, 4 + GRILLA_NOTAS_COLS_MAX + 4)
    nota_cols = {}
    estado_col = None
    promedio_col = None
    rx = re.compile(r'^nota\s*(\d+)\s*$', re.I)
    for c in range(5, max_c + 1):
        h = ws.cell(row=header_row, column=c).value
        hs = str(h or '').strip()
        if not hs:
            continue
        hum = hs.lower()
        if 'promed' in hum:
            promedio_col = c
            continue
        if 'estado' in hum:
            estado_col = c
            continue
        m = rx.match(hs.strip())
        if m:
            nota_cols[c] = int(m.group(1))
    if not nota_cols:
        for c in range(5, 5 + cfg.num_columnas):
            nota_cols[c] = c - 4
    max_idx = max(nota_cols.values()) if nota_cols else cfg.num_columnas
    return nota_cols, estado_col, promedio_col, max_idx


def _grilla_notas_estado_desde_excel(raw) -> str:
    if raw is None or raw == '':
        return ''
    s = str(raw).strip().lower()
    s = ' '.join(s.split())
    if not s:
        return ''
    if 'no aprob' in s or s.startswith('no ') and 'aprob' in s:
        return NotasGrillaEstadoFinal.Estado.NO_APROBADO
    if 'aprob' in s:
        return NotasGrillaEstadoFinal.Estado.APROBADO
    return ''


@login_required
@require_POST
def profesor_escuela_notas_agregar_columna(request, escuela_id):
    pack, redir = _profesor_escuela_docente(request, escuela_id)
    if redir:
        return redir
    user_sede, escuela, _prof = pack
    cfg = _grilla_notas_config_asegurada(escuela, user_sede)
    if cfg.num_columnas >= GRILLA_NOTAS_COLS_MAX:
        messages.error(request, f'Ya hay {GRILLA_NOTAS_COLS_MAX} columnas de nota (límite).')
        return redirect('hechos:profesor_escuela_notas_panel', escuela_id=escuela.id)
    cfg.num_columnas += 1
    cfg.save(update_fields=['num_columnas', 'updated_at'])
    messages.success(request, f'Se añadió «Nota {cfg.num_columnas}».')
    return redirect('hechos:profesor_escuela_notas_panel', escuela_id=escuela.id)


@login_required
@require_POST
def profesor_escuela_notas_quitar_columna(request, escuela_id):
    pack, redir = _profesor_escuela_docente(request, escuela_id)
    if redir:
        return redir
    user_sede, escuela, _prof = pack
    cfg = _grilla_notas_config_asegurada(escuela, user_sede)
    if cfg.num_columnas <= 1:
        messages.warning(request, 'Debe quedar al menos una columna de nota.')
        return redirect('hechos:profesor_escuela_notas_panel', escuela_id=escuela.id)
    NotasGrillaCelda.objects.filter(
        escuela=escuela, sede=user_sede, columna=cfg.num_columnas
    ).delete()
    cfg.num_columnas -= 1
    cfg.save(update_fields=['num_columnas', 'updated_at'])
    messages.success(request, 'Se eliminó la última columna de nota y sus valores.')
    return redirect('hechos:profesor_escuela_notas_panel', escuela_id=escuela.id)


@login_required
def profesor_escuela_notas_import_excel(request, escuela_id):
    if request.method != 'POST':
        return redirect('hechos:profesor_escuela_notas_panel', escuela_id=escuela_id)
    pack, redir = _profesor_escuela_docente(request, escuela_id)
    if redir:
        return redir
    user_sede, escuela, _prof = pack
    f = request.FILES.get('archivo_excel')
    if not f:
        messages.error(request, 'Selecciona un archivo Excel (.xlsx).')
        return redirect('hechos:profesor_escuela_notas_panel', escuela_id=escuela.id)
    if f.size > 6 * 1024 * 1024:
        messages.error(request, 'El archivo es demasiado grande (máx. 6 MB).')
        return redirect('hechos:profesor_escuela_notas_panel', escuela_id=escuela.id)
    head = f.read(4)
    f.seek(0)
    if head != b'PK\x03\x04' and head != b'PK\x05\x06':
        messages.error(request, 'El archivo no parece un .xlsx válido.')
        return redirect('hechos:profesor_escuela_notas_panel', escuela_id=escuela.id)
    try:
        from openpyxl import load_workbook
    except Exception:
        messages.error(request, 'Falta openpyxl para importar Excel.')
        return redirect('hechos:profesor_escuela_notas_panel', escuela_id=escuela.id)

    estudiantes = list(_grilla_notas_estudiantes_qs(user_sede, escuela))
    doc_a_est = {}
    for est in estudiantes:
        k = _grilla_doc_normalizado(est.numero_documento)
        if k:
            doc_a_est[k] = est

    cfg = _grilla_notas_config_asegurada(escuela, user_sede)
    try:
        from io import BytesIO

        data = f.read()
        wb = load_workbook(BytesIO(data), data_only=True)
        ws = wb.active
    except Exception as exc:
        messages.error(request, f'No se pudo leer el Excel: {exc}')
        return redirect('hechos:profesor_escuela_notas_panel', escuela_id=escuela.id)

    header_row = _grilla_notas_excel_find_header_row(ws)
    nota_cols, estado_col, _prom_col, max_idx = _grilla_notas_import_mapeo_columnas(
        ws, header_row, cfg
    )
    target_n = min(max(max_idx, cfg.num_columnas), GRILLA_NOTAS_COLS_MAX)
    if target_n != cfg.num_columnas:
        cfg.num_columnas = target_n
        cfg.save(update_fields=['num_columnas', 'updated_at'])

    actualizados = 0
    sin_doc = 0
    try:
        with transaction.atomic():
            max_r = ws.max_row or 0
            for r in range(header_row + 1, max_r + 1):
                raw_doc = ws.cell(row=r, column=2).value
                doc_key = _grilla_doc_normalizado(raw_doc)
                if not doc_key:
                    continue
                est = doc_a_est.get(doc_key)
                if not est:
                    sin_doc += 1
                    continue
                for excel_c, nota_i in nota_cols.items():
                    if nota_i < 1 or nota_i > cfg.num_columnas:
                        continue
                    val = _grilla_notas_valor_guardar(ws.cell(row=r, column=excel_c).value)
                    if val == '':
                        NotasGrillaCelda.objects.filter(
                            escuela=escuela,
                            sede=user_sede,
                            estudiante_id=est.id,
                            columna=nota_i,
                        ).delete()
                        continue
                    obj, created = NotasGrillaCelda.objects.get_or_create(
                        escuela=escuela,
                        sede=user_sede,
                        estudiante_id=est.id,
                        columna=nota_i,
                        defaults={
                            'valor': val,
                            'registrado_por': request.user,
                        },
                    )
                    if created:
                        actualizados += 1
                    elif obj.valor != val:
                        obj.valor = val
                        obj.registrado_por = request.user
                        obj.save(update_fields=['valor', 'registrado_por', 'updated_at'])
                        actualizados += 1
                if estado_col:
                    ev = _grilla_notas_estado_desde_excel(ws.cell(row=r, column=estado_col).value)
                    if ev != '':
                        NotasGrillaEstadoFinal.objects.update_or_create(
                            escuela=escuela,
                            estudiante_id=est.id,
                            defaults={
                                'sede': user_sede,
                                'estado': ev,
                            },
                        )
                        actualizados += 1
    except Exception as exc:
        messages.error(request, f'Error al importar: {exc}')
        return redirect('hechos:profesor_escuela_notas_panel', escuela_id=escuela.id)

    parts = [f'Importación aplicada ({actualizados} cambios registrados).']
    if sin_doc:
        parts.append(
            f'{sin_doc} fila(s) con documento no reconocido entre los matriculados se omitieron.'
        )
    messages.success(request, ' '.join(parts))
    return redirect('hechos:profesor_escuela_notas_panel', escuela_id=escuela.id)


def _profesor_escuela_notas_excel_response(request, escuela_id, *, plantilla_vacia: bool):
    pack, redir = _profesor_escuela_docente(request, escuela_id)
    if redir:
        return redir
    user_sede, escuela, prof = pack
    estudiantes = list(_grilla_notas_estudiantes_qs(user_sede, escuela))
    cfg = _grilla_notas_config_asegurada(escuela, user_sede)
    n = cfg.num_columnas
    edicion = _edicion_principal_escuela(user_sede, escuela)
    docente_user = None
    if edicion and edicion.maestro_id:
        docente_user = edicion.maestro.user
    elif prof:
        docente_user = prof.user
    nombre_docente = (
        (docente_user.get_full_name() or docente_user.email or docente_user.username)
        if docente_user
        else '—'
    )
    texto_escuela = f'{escuela.nombre} — {escuela.anio} · {escuela.get_ciclo_display()}'
    horario_txt = _horario_edicion_a_texto_ui(edicion) if edicion else '—'

    est_ids = [e.id for e in estudiantes]
    valores_grid: dict[tuple[int, int], str] = {}
    estados: dict[int, str] = {}
    if not plantilla_vacia and est_ids:
        for cel in NotasGrillaCelda.objects.filter(
            escuela=escuela,
            sede=user_sede,
            estudiante_id__in=est_ids,
            columna__lte=n,
        ):
            valores_grid[(cel.estudiante_id, cel.columna)] = cel.valor
        for ef in NotasGrillaEstadoFinal.objects.filter(
            escuela=escuela, estudiante_id__in=est_ids
        ):
            estados[ef.estudiante_id] = ef.estado

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side, Protection
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation
        from openpyxl.formatting.rule import FormulaRule
    except Exception:
        messages.error(request, 'Falta la dependencia para generar Excel. Instala openpyxl.')
        return redirect('hechos:profesor_escuela_notas_panel', escuela_id=escuela.id)

    wb = Workbook()
    ws = wb.active
    ws.title = 'Formato notas' if plantilla_vacia else 'Notas'

    fixed_headers = ['NO', 'Documento', 'Nombres y Apellidos', 'Celular']
    nota_headers = [f'Nota {i}' for i in range(1, n + 1)]
    tail_headers = ['Promedio', 'Estado']
    base_cols = len(fixed_headers)
    total_cols = base_cols + n + len(tail_headers)

    header_fill = PatternFill('solid', fgColor='D9E1F2')
    header_font = Font(color='000000', bold=True)
    thin_side = Side(style='thin', color='B4B4B4')
    border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    center = Alignment(horizontal='center', vertical='center')
    left = Alignment(horizontal='left', vertical='center')
    meta_label_fill = PatternFill('solid', fgColor='0B5394')
    meta_label_font = Font(color='FFFFFF', bold=True)
    meta_val_fill = PatternFill('solid', fgColor='FFFFFF')
    data_font = Font(color='000000')
    estado_fill_aprobado = PatternFill('solid', fgColor='C6EFCE')
    estado_fill_no_aprobado = PatternFill('solid', fgColor='FFC7CE')
    estado_font_aprobado = Font(color='166534', bold=True)
    estado_font_no_aprobado = Font(color='991B1B', bold=True)

    def _meta_celda(r, c, val, es_etiqueta=False):
        cell = ws.cell(row=r, column=c, value=val)
        cell.border = border
        if es_etiqueta:
            cell.fill = meta_label_fill
            cell.font = meta_label_font
        else:
            cell.fill = meta_val_fill
            cell.font = data_font
        cell.alignment = left

    _meta_celda(1, 2, 'Docente', es_etiqueta=True)
    _meta_celda(1, 3, nombre_docente)
    _meta_celda(2, 2, 'Curso', es_etiqueta=True)
    _meta_celda(2, 3, texto_escuela)
    _meta_celda(3, 2, 'Horario', es_etiqueta=True)
    _meta_celda(3, 3, horario_txt)

    header_row = 6
    col_idx = 1
    for h in fixed_headers + nota_headers + tail_headers:
        c = ws.cell(row=header_row, column=col_idx, value=h)
        c.fill = header_fill
        c.font = header_font
        c.border = border
        c.alignment = center
        col_idx += 1

    first_nota_col = base_cols + 1
    last_nota_col = base_cols + n
    promedio_col = base_cols + n + 1
    estado_col = base_cols + n + 2
    ls = get_column_letter(first_nota_col) if n else None
    le = get_column_letter(last_nota_col) if n else None

    row = header_row + 1
    for idx, est in enumerate(estudiantes, start=1):
        ws.cell(row=row, column=1, value=idx)
        doc = (est.numero_documento or '').strip()
        if doc.isdigit():
            try:
                ws.cell(row=row, column=2, value=int(doc))
            except ValueError:
                ws.cell(row=row, column=2, value=doc)
        else:
            ws.cell(row=row, column=2, value=doc or '')
        ws.cell(row=row, column=3, value=est.user.get_full_name() or est.user.email or '—')
        cel_txt = (est.user.phone or '').strip() or (est.telefono_emergencia or '').strip()
        ws.cell(row=row, column=4, value=cel_txt)
        vals_row = {}
        for ci in range(1, n + 1):
            v = '' if plantilla_vacia else valores_grid.get((est.id, ci), '')
            vals_row[ci] = v
            ce = ws.cell(row=row, column=base_cols + ci, value=v or None)
            ce.font = data_font
            ce.border = border
            ce.alignment = center
        # Promedio siempre automático desde columnas de nota.
        if n and ls and le:
            prom_cell = ws.cell(
                row=row,
                column=promedio_col,
                value=f'=IFERROR(AVERAGE({ls}{row}:{le}{row}),"")',
            )
            prom_cell.number_format = '0,####'
        else:
            ws.cell(row=row, column=promedio_col, value='')
        est_lbl = ''
        if not plantilla_vacia:
            ev = estados.get(est.id, '')
            if ev == NotasGrillaEstadoFinal.Estado.APROBADO:
                est_lbl = 'Aprobado'
            elif ev == NotasGrillaEstadoFinal.Estado.NO_APROBADO:
                est_lbl = 'No aprobado'
        ws.cell(row=row, column=estado_col, value=est_lbl or None)
        for tc in (promedio_col, estado_col):
            tcell = ws.cell(row=row, column=tc)
            tcell.font = data_font
            tcell.border = border
            tcell.alignment = center
        if est_lbl == 'Aprobado':
            ecell = ws.cell(row=row, column=estado_col)
            ecell.fill = estado_fill_aprobado
            ecell.font = estado_font_aprobado
        elif est_lbl == 'No aprobado':
            ecell = ws.cell(row=row, column=estado_col)
            ecell.fill = estado_fill_no_aprobado
            ecell.font = estado_font_no_aprobado
        for c in range(1, 5):
            cell = ws.cell(row=row, column=c)
            cell.font = data_font
            cell.border = border
            cell.alignment = left if c in (2, 3) else center
        row += 1

    ws.column_dimensions['A'].width = 5
    ws.column_dimensions['B'].width = 14
    ws.column_dimensions['C'].width = 34
    ws.column_dimensions['D'].width = 14
    for ci in range(first_nota_col, total_cols + 1):
        letter = get_column_letter(ci)
        ws.column_dimensions[letter].width = 12 if ci <= last_nota_col else 14

    first_data_row = header_row + 1
    data_last_row = row - 1
    if data_last_row >= first_data_row:
        estado_letter = get_column_letter(estado_col)
        dv_end_row = max(data_last_row, first_data_row) + 25
        dv_estado = DataValidation(
            type='list',
            formula1='"Aprobado,No aprobado"',
            allow_blank=True,
            showErrorMessage=True,
        )
        dv_estado.error = 'Use solo Aprobado o No aprobado.'
        dv_estado.errorTitle = 'Valor no permitido'
        ws.add_data_validation(dv_estado)
        dv_estado.add(f'{estado_letter}{first_data_row}:{estado_letter}{dv_end_row}')

        rango_estado = f'${estado_letter}${first_data_row}:${estado_letter}${data_last_row}'
        ws.conditional_formatting.add(
            rango_estado,
            FormulaRule(
                formula=[f'${estado_letter}{first_data_row}="Aprobado"'],
                stopIfTrue=False,
                fill=estado_fill_aprobado,
                font=estado_font_aprobado,
            ),
        )
        ws.conditional_formatting.add(
            rango_estado,
            FormulaRule(
                formula=[f'${estado_letter}{first_data_row}="No aprobado"'],
                stopIfTrue=False,
                fill=estado_fill_no_aprobado,
                font=estado_font_no_aprobado,
            ),
        )

        # Bloquear hoja para evitar cambios accidentales en columnas fijas.
        # Solo quedan editables: columnas de Nota y Estado.
        for rr in range(first_data_row, dv_end_row + 1):
            for cc in range(1, 5):
                ws.cell(row=rr, column=cc).protection = Protection(locked=True)
            for cc in range(first_nota_col, last_nota_col + 1):
                ws.cell(row=rr, column=cc).protection = Protection(locked=False)
            ws.cell(row=rr, column=estado_col).protection = Protection(locked=False)
            ws.cell(row=rr, column=promedio_col).protection = Protection(locked=True)
        ws.protection.sheet = True
        ws.protection.selectLockedCells = True
        ws.protection.selectUnlockedCells = True

    if n:
        ws.freeze_panes = f'{get_column_letter(first_nota_col)}{header_row + 1}'

    nombre_corto = (slugify(escuela.nombre) or 'escuela')[:24].strip('-')
    fecha = timezone.now().strftime('%Y%m%d')
    filename = (
        f'formato-notas-e{escuela.id}-{nombre_corto}-{fecha}.xlsx'
        if plantilla_vacia
        else f'notas-e{escuela.id}-{nombre_corto}-{fecha}.xlsx'
    )
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


@login_required
def profesor_escuela_notas_excel(request, escuela_id):
    return _profesor_escuela_notas_excel_response(request, escuela_id, plantilla_vacia=False)


@login_required
def profesor_escuela_notas_excel_formato(request, escuela_id):
    return _profesor_escuela_notas_excel_response(request, escuela_id, plantilla_vacia=True)


@login_required
def profesor_escuela_notas_panel(request, escuela_id):
    pack, redir = _profesor_escuela_docente(request, escuela_id)
    if redir:
        return redir
    user_sede, escuela, prof = pack

    estudiantes = list(_grilla_notas_estudiantes_qs(user_sede, escuela))
    est_ids = [e.id for e in estudiantes]
    cfg = _grilla_notas_config_asegurada(escuela, user_sede)
    num_columnas = cfg.num_columnas

    valores: dict[tuple[int, int], str] = {}
    if est_ids:
        for cel in NotasGrillaCelda.objects.filter(
            escuela=escuela,
            sede=user_sede,
            estudiante_id__in=est_ids,
            columna__lte=num_columnas,
        ).only('estudiante_id', 'columna', 'valor'):
            valores[(cel.estudiante_id, cel.columna)] = cel.valor

    estados_por_est: dict[int, str] = {}
    if est_ids:
        for ef in NotasGrillaEstadoFinal.objects.filter(
            escuela=escuela, estudiante_id__in=est_ids
        ):
            estados_por_est[ef.estudiante_id] = ef.estado

    if request.method == 'POST':
        est_set = frozenset(est_ids)
        try:
            with transaction.atomic():
                for key, raw_val in request.POST.items():
                    if key.startswith('n_'):
                        parts = key.split('_')
                        if len(parts) != 3:
                            continue
                        try:
                            est_id = int(parts[1])
                            col = int(parts[2])
                        except ValueError:
                            continue
                        if (
                            est_id not in est_set
                            or col < 1
                            or col > num_columnas
                        ):
                            continue
                        val = _grilla_notas_valor_guardar(raw_val)
                        if val == '':
                            NotasGrillaCelda.objects.filter(
                                escuela=escuela,
                                sede=user_sede,
                                estudiante_id=est_id,
                                columna=col,
                            ).delete()
                            continue
                        obj, created = NotasGrillaCelda.objects.get_or_create(
                            escuela=escuela,
                            sede=user_sede,
                            estudiante_id=est_id,
                            columna=col,
                            defaults={
                                'valor': val,
                                'registrado_por': request.user,
                            },
                        )
                        if not created:
                            obj.valor = val
                            obj.registrado_por = request.user
                            obj.save(
                                update_fields=['valor', 'registrado_por', 'updated_at']
                            )
                    elif key.startswith('estado_'):
                        try:
                            est_id = int(key.split('_', 1)[1])
                        except (ValueError, IndexError):
                            continue
                        if est_id not in est_set:
                            continue
                        ev = (raw_val or '').strip()
                        if ev not in (
                            '',
                            NotasGrillaEstadoFinal.Estado.APROBADO,
                            NotasGrillaEstadoFinal.Estado.NO_APROBADO,
                        ):
                            continue
                        if ev == '':
                            NotasGrillaEstadoFinal.objects.filter(
                                escuela=escuela, estudiante_id=est_id
                            ).delete()
                        else:
                            NotasGrillaEstadoFinal.objects.update_or_create(
                                escuela=escuela,
                                estudiante_id=est_id,
                                defaults={'sede': user_sede, 'estado': ev},
                            )
        except Exception as exc:
            messages.error(request, f'No se pudo guardar: {exc}')
            return redirect('hechos:profesor_escuela_notas_panel', escuela_id=escuela.id)

        messages.success(request, 'Notas guardadas.')
        return redirect('hechos:profesor_escuela_notas_panel', escuela_id=escuela.id)

    encabezados_columnas = [{'idx': i, 'label': f'Nota {i}'} for i in range(1, num_columnas + 1)]

    filas = []
    for est in estudiantes:
        por_col = {c: valores.get((est.id, c), '') for c in range(1, num_columnas + 1)}
        celdas = [{'col': c, 'value': por_col.get(c, '')} for c in range(1, num_columnas + 1)]
        promedio_ui = _grilla_notas_promedio_display(por_col, num_columnas)
        filas.append(
            {
                'estudiante': est,
                'celdas': celdas,
                'promedio_ui': promedio_ui,
                'estado_final': estados_por_est.get(est.id, ''),
            }
        )

    volver_url = f"{reverse('hechos:mis_escuelas_profesor')}?escuela={escuela.id}"
    edicion = _edicion_principal_escuela(user_sede, escuela)
    docente_user = None
    if edicion and edicion.maestro_id:
        docente_user = edicion.maestro.user
    elif prof:
        docente_user = prof.user
    nombre_docente = (
        (docente_user.get_full_name() or docente_user.email or docente_user.username)
        if docente_user
        else '—'
    )
    horario_txt = _horario_edicion_a_texto_ui(edicion) if edicion else '—'
    modalidad_txt = _modalidad_edicion_label(edicion)
    ret_notas = reverse('hechos:profesor_escuela_notas_panel', args=[escuela.id])
    matricular_modal_activo = bool(
        prof
        and _profesor_puede_gestionar_escuela(prof, escuela)
        and escuela.cupos_disponibles > 0
    )

    puede_quitar_col = num_columnas > 1
    puede_agregar_col = num_columnas < GRILLA_NOTAS_COLS_MAX

    return render(
        request,
        'hechos/profesor_escuela_notas_panel.html',
        {
            'escuela': escuela,
            'user_sede': user_sede,
            'volver_url': volver_url,
            'encabezados_columnas': encabezados_columnas,
            'filas': filas,
            'num_columnas': num_columnas,
            'tiene_estudiantes': bool(estudiantes),
            'nombre_docente': nombre_docente,
            'horario_txt': horario_txt,
            'modalidad_txt': modalidad_txt,
            'agregar_nota_url': reverse(
                'hechos:profesor_escuela_notas_agregar_columna', args=[escuela.id]
            ),
            'quitar_nota_url': reverse(
                'hechos:profesor_escuela_notas_quitar_columna', args=[escuela.id]
            ),
            'puede_quitar_col': puede_quitar_col,
            'puede_agregar_col': puede_agregar_col,
            'matricular_modal_escuela': escuela,
            'matricular_modal_activo': matricular_modal_activo,
            'matricular_modal_return_url': ret_notas,
            'matricular_modal_post_url': reverse('hechos:matricular_estudiante'),
            'matricular_modal_sugerencias_url': reverse(
                'hechos:sugerencias_estudiante_doc_matricula'
            ),
        },
    )


@login_required
def asistencia_clase(request, clase_id):
    """
    Tomar asistencia de una clase (solo para profesores)
    """
    if not hasattr(request.user, 'profesor_profile'):
        messages.error(request, 'Solo los profesores pueden tomar asistencia.')
        return redirect('core:dashboard')
    
    # Obtener la sede del usuario
    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada. Contacta al administrador.')
        return redirect('core:dashboard')
    
    prof = request.user.profesor_profile
    clase = get_object_or_404(
        Clase.objects.select_related(
            "profesor",
            "escuela__maestro",
            "escuela",
            "escuela",
            "escuela",
        ),
        id=clase_id,
        sede=user_sede,
    )
    if not _profesor_puede_gestionar_clase(clase, prof):
        messages.error(request, "No tienes permisos para tomar asistencia en esta clase.")
        return redirect("core:dashboard")

    estudiantes_qs = Estudiante.objects.filter(
        matriculas__escuela=clase.escuela,
        matriculas__sede=user_sede,
        matriculas__is_active=True,
        sede=user_sede,
    ).distinct().select_related("user")

    if request.method == 'POST':
        valid_states = {'presente', 'ausente', 'tarde', 'justificado'}
        for estudiante in estudiantes_qs:
            raw = (request.POST.get(f'estudiante_{estudiante.id}') or 'ausente').strip()
            value = raw if raw in valid_states else 'ausente'
            asistencia, created = Asistencia.objects.get_or_create(
                clase=clase,
                estudiante=estudiante,
                sede=user_sede,
                defaults={
                    'estado': value,
                    'registrado_por': request.user,
                },
            )
            if not created:
                asistencia.estado = value
                asistencia.registrado_por = request.user
                asistencia.save(update_fields=['estado', 'registrado_por'])

        messages.success(request, 'Asistencia registrada correctamente.')
        return redirect('hechos:asistencia_clase', clase_id=clase_id)

    estudiantes = estudiantes_qs
    
    asistencias = list(
        Asistencia.objects.filter(clase=clase, sede=user_sede).select_related('estudiante__user')
    )
    asistencia_dict = {a.estudiante_id: a.estado for a in asistencias}
    filas_asistencia = [
        {'estudiante': e, 'estado': asistencia_dict.get(e.id)} for e in estudiantes
    ]
    presentes_count = sum(1 for a in asistencias if a.estado == 'presente')
    ausentes_count = sum(1 for a in asistencias if a.estado == 'ausente')
    tarde_count = sum(1 for a in asistencias if a.estado == 'tarde')
    just_count = sum(1 for a in asistencias if a.estado == 'justificado')

    escuela_pk = None
    rt = getattr(clase.curso, 'ruta_estudio', None) if clase.escuela_id else None
    if rt is not None:
        escuela_pk = getattr(rt, 'escuela_id', None)
    volver_url = reverse('hechos:mis_escuelas_profesor')
    if escuela_pk:
        volver_url = f'{volver_url}?escuela={escuela_pk}'

    context = {
        'clase': clase,
        'estudiantes': estudiantes,
        'filas_asistencia': filas_asistencia,
        'presentes_count': presentes_count,
        'ausentes_count': ausentes_count,
        'tarde_count': tarde_count,
        'justificado_count': just_count,
        'user_sede': user_sede,
        'volver_mis_escuelas_url': volver_url,
    }

    return render(request, 'hechos/asistencia_clase_modern.html', context)


def _sesion_herramienta_placeholder(request, clase_id, titulo, texto_ayuda):
    if not hasattr(request.user, 'profesor_profile'):
        messages.error(request, 'Solo los profesores pueden acceder.')
        return redirect('core:dashboard')
    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada. Contacta al administrador.')
        return redirect('core:dashboard')
    prof = request.user.profesor_profile
    clase = get_object_or_404(
        Clase.objects.select_related(
            'escuela',
            'escuela',
        ),
        id=clase_id,
        sede=user_sede,
    )
    if not _profesor_puede_gestionar_clase(clase, prof):
        messages.error(request, 'No tienes permisos para gestionar esta clase.')
        return redirect('core:dashboard')
    rt = getattr(clase.curso, 'ruta_estudio', None) if clase.escuela_id else None
    escuela_pk = getattr(rt, 'escuela_id', None) if rt is not None else None
    volver_url = reverse('hechos:mis_escuelas_profesor')
    if escuela_pk:
        volver_url = f'{volver_url}?escuela={escuela_pk}'
    return render(
        request,
        'hechos/sesion_herramienta_placeholder.html',
        {
            'clase': clase,
            'titulo_herramienta': titulo,
            'texto_ayuda': texto_ayuda,
            'volver_url': volver_url,
            'user_sede': user_sede,
        },
    )


@login_required
def escuela_crear_actividad(request, escuela_id):
    """
    Crear actividad para la escuela (flujo tipo assignment de LMS: tipo, título,
    instrucciones, fecha límite, puntos, archivo opcional).
    """
    if not hasattr(request.user, 'profesor_profile'):
        messages.error(request, 'Solo los profesores pueden acceder.')
        return redirect('core:dashboard')
    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada. Contacta al administrador.')
        return redirect('core:dashboard')
    prof = request.user.profesor_profile
    escuela = get_object_or_404(
        Escuela,
        id=escuela_id,
        sede=user_sede,
        maestro=prof,
        is_active=True,
    )
    volver_url = f"{reverse('hechos:mis_escuelas_profesor')}?escuela={escuela.id}"

    if request.method == 'POST':
        form = ActividadEscuelaForm(request.POST, request.FILES)
        if form.is_valid():
            obj = form.save(commit=False)
            obj.sede = user_sede
            obj.escuela = escuela
            obj.creada_por = prof
            obj.save()
            messages.success(
                request,
                f'Actividad «{obj.titulo}» publicada para esta escuela.',
            )
            return redirect('hechos:escuela_crear_actividad', escuela_id=escuela.id)
        messages.error(request, 'Revisa los datos del formulario.')
    else:
        form = ActividadEscuelaForm()

    actividades = (
        ActividadEscuela.objects.filter(
            escuela=escuela,
            sede=user_sede,
            is_active=True,
        )
        .annotate(n_entregas=Count('entregas'))
        .select_related('creada_por__user')
        .order_by('-created_at')[:100]
    )

    return render(
        request,
        'hechos/escuela_actividad_form.html',
        {
            'escuela': escuela,
            'form': form,
            'volver_url': volver_url,
            'user_sede': user_sede,
            'actividades': actividades,
        },
    )


def _estudiante_matriculado_en_curso(curso, estudiante, sede):
    return Matricula.objects.filter(
        estudiante=estudiante,
        escuela=curso,
        sede=sede,
        is_active=True,
    ).exists()


@login_required
def entregar_actividad_estudiante(request, curso_id, actividad_id):
    """
    El estudiante sube evidencia (texto y/o archivo) para una actividad de su escuela.
    """
    if not hasattr(request.user, 'estudiante_profile'):
        messages.error(request, 'Solo los estudiantes pueden entregar actividades.')
        return redirect('core:dashboard')
    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada.')
        return redirect('core:dashboard')
    estudiante = request.user.estudiante_profile
    curso = get_object_or_404(Curso, id=curso_id, sede=user_sede, is_active=True)
    if not _estudiante_matriculado_en_curso(curso, estudiante, user_sede):
        messages.error(request, 'No tienes matrícula activa en este curso.')
        return redirect('hechos:mis_escuelas')
    escuela = curso.escuela
    if not escuela:
        messages.error(request, 'Este curso no está asociado a una escuela.')
        return redirect('hechos:mis_escuelas')
    actividad = get_object_or_404(
        ActividadEscuela,
        id=actividad_id,
        escuela=escuela,
        sede=user_sede,
        is_active=True,
    )
    volver = reverse('hechos:detalle_mis_escuela', kwargs={'curso_id': curso.id})
    if actividad.fecha_limite and timezone.now() > actividad.fecha_limite:
        messages.warning(
            request,
            'La fecha límite de esta actividad ya pasó; ya no puedes enviar ni modificar tu entrega.',
        )
        return redirect(f'{volver}#actividades')

    entrega, _created = EntregaActividad.objects.get_or_create(
        actividad=actividad,
        estudiante=estudiante,
        defaults={'sede': user_sede},
    )
    if entrega.evaluado_en:
        messages.info(request, 'Tu entrega ya fue evaluada por el docente y no se puede cambiar.')
        return redirect(f'{volver}#actividades')

    if request.method == 'POST':
        # Texto con ModelForm; archivo con request.FILES + input nativo visible (mismo patrón que el avatar).
        form = EntregaActividadTextoForm(request.POST, instance=entrega)
        uf = request.FILES.get("archivo")
        if uf is not None and getattr(uf, "size", 0) == 0:
            messages.error(
                request,
                "El archivo que elegiste pesa 0 bytes: está vacío o no se guardó nada dentro. "
                "Ábrelo (Bloc de notas, Word, etc.), escribe o pega tu tarea, guarda el archivo "
                "y vuelve a subirlo.",
            )
            uf = None
        nuevo_archivo = uf is not None and getattr(uf, "size", 0) > 0
        archivo_rechazado = False
        if nuevo_archivo:
            uf = normalizar_archivo_entrega(uf)
            try:
                entrega_evidencia_archivo_validator(uf)
            except DjangoValidationError as exc:
                archivo_rechazado = True
                err_msg = exc.messages[0] if exc.messages else str(exc)
                messages.error(request, err_msg)

        if form.is_valid() and not archivo_rechazado:
            texto = (form.cleaned_data.get("texto") or "").strip()
            tenia_archivo = bool(entrega.archivo and entrega.archivo.name)
            if not texto and not tenia_archivo and not nuevo_archivo:
                messages.error(
                    request,
                    "Debes escribir una respuesta o subir un archivo con contenido "
                    "(los archivos vacíos no cuentan).",
                )
            else:
                try:
                    mr = str(settings.MEDIA_ROOT)
                    os.makedirs(mr, exist_ok=True)
                    os.makedirs(os.path.join(mr, "actividades"), exist_ok=True)
                    obj = form.save(commit=False)
                    obj.sede = user_sede
                    if nuevo_archivo:
                        obj.archivo = uf
                    obj.save()
                except OSError as exc:
                    messages.error(
                        request,
                        "No se pudo escribir el archivo en la carpeta media del proyecto "
                        "(permisos, carpeta inexistente o carpeta sincronizada en la nube bloqueando). "
                        f"Detalle: {exc}",
                    )
                    return redirect(request.path)
                entrega = EntregaActividad.objects.get(pk=obj.pk)
                form = EntregaActividadTextoForm(instance=entrega)
                archivo_basename = (
                    os.path.basename(entrega.archivo.name)
                    if (entrega.archivo and entrega.archivo.name)
                    else ''
                )
                if nuevo_archivo:
                    messages.success(request, "Archivo guardado en media/actividades/.")
                else:
                    messages.success(request, "Tu entrega se guardó correctamente.")
                return render(
                    request,
                    'hechos/entregar_actividad_estudiante.html',
                    {
                        'form': form,
                        'curso': curso,
                        'actividad': actividad,
                        'entrega': entrega,
                        'volver_url': f'{volver}#actividades',
                        'user_sede': user_sede,
                        'entrega_archivo_basename': archivo_basename,
                    },
                )
        elif not form.is_valid():
            messages.error(request, 'Revisa los datos del formulario.')
    else:
        form = EntregaActividadTextoForm(instance=entrega)

    entrega.refresh_from_db()
    archivo_basename = (
        os.path.basename(entrega.archivo.name)
        if (entrega.archivo and entrega.archivo.name)
        else ''
    )

    return render(
        request,
        'hechos/entregar_actividad_estudiante.html',
        {
            'form': form,
            'curso': curso,
            'actividad': actividad,
            'entrega': entrega,
            'volver_url': f'{volver}#actividades',
            'user_sede': user_sede,
            'entrega_archivo_basename': archivo_basename,
        },
    )


@login_required
def descargar_archivo_entrega_actividad(request, curso_id, actividad_id):
    """Descarga el archivo de la entrega del estudiante (forzar adjunto, misma autorización que entregar)."""
    if not hasattr(request.user, 'estudiante_profile'):
        raise Http404
    user_sede = get_user_sede(request.user)
    if not user_sede:
        raise Http404
    estudiante = request.user.estudiante_profile
    curso = get_object_or_404(Curso, id=curso_id, sede=user_sede, is_active=True)
    if not _estudiante_matriculado_en_curso(curso, estudiante, user_sede):
        raise Http404
    escuela = curso.escuela
    if not escuela:
        raise Http404
    actividad = get_object_or_404(
        ActividadEscuela,
        id=actividad_id,
        escuela=escuela,
        sede=user_sede,
        is_active=True,
    )
    entrega = get_object_or_404(
        EntregaActividad,
        actividad=actividad,
        estudiante=estudiante,
        sede=user_sede,
    )
    if not entrega.archivo or not entrega.archivo.name:
        raise Http404
    fname = os.path.basename(entrega.archivo.name)
    try:
        fh = entrega.archivo.open("rb")
    except FileNotFoundError:
        raise Http404 from None
    return FileResponse(fh, as_attachment=True, filename=fname)


@login_required
def profesor_actividad_entregas(request, escuela_id, actividad_id):
    """Lista entregas de una actividad y permite evaluar (puntaje + comentario)."""
    if not hasattr(request.user, 'profesor_profile'):
        messages.error(request, 'Solo los profesores pueden revisar entregas.')
        return redirect('core:dashboard')
    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada.')
        return redirect('core:dashboard')
    prof = request.user.profesor_profile
    actividad = get_object_or_404(
        ActividadEscuela.objects.select_related('escuela'),
        id=actividad_id,
        escuela_id=escuela_id,
        sede=user_sede,
        escuela__maestro=prof,
        is_active=True,
    )
    volver = reverse('hechos:escuela_crear_actividad', kwargs={'escuela_id': escuela_id})

    failed_entrega_id = None
    failed_form = None
    if request.method == 'POST' and request.POST.get('action') == 'evaluar':
        entrega = get_object_or_404(
            EntregaActividad.objects.select_related('estudiante__user'),
            id=request.POST.get('entrega_id'),
            actividad=actividad,
            sede=user_sede,
        )
        prefix = f'e{entrega.id}'
        form = ProfesorEvaluaEntregaForm(
            request.POST, instance=entrega, actividad=actividad, prefix=prefix,
        )
        if form.is_valid():
            obj = form.save(commit=False)
            obj.evaluado_en = timezone.now()
            obj.evaluado_por = prof
            obj.save()
            messages.success(
                request,
                f'Evaluación guardada para {entrega.estudiante.user.get_full_name()}.',
            )
            return redirect('hechos:profesor_actividad_entregas', escuela_id=escuela_id, actividad_id=actividad_id)
        messages.error(request, 'Revisa el puntaje o el comentario.')
        failed_entrega_id = entrega.id
        failed_form = form

    entregas = (
        EntregaActividad.objects.filter(actividad=actividad, sede=user_sede)
        .select_related('estudiante__user')
        .order_by('estudiante__user__last_name', 'estudiante__user__first_name')
    )
    entregas_rows = []
    for e in entregas:
        if failed_form is not None and e.id == failed_entrega_id:
            frm = failed_form
        else:
            frm = ProfesorEvaluaEntregaForm(
                instance=e, actividad=actividad, prefix=f'e{e.id}',
            )
        entregas_rows.append({'entrega': e, 'form': frm})

    return render(
        request,
        'hechos/profesor_actividad_entregas.html',
        {
            'actividad': actividad,
            'escuela': actividad.escuela,
            'entregas_rows': entregas_rows,
            'volver_url': volver,
            'user_sede': user_sede,
        },
    )


@login_required
def profesor_descargar_entrega_archivo(request, escuela_id, actividad_id, entrega_id):
    """
    Sirve el archivo de evidencia de una entrega (solo el maestro de la escuela).
    Fuerza descarga con nombre de archivo legible.
    """
    if not hasattr(request.user, 'profesor_profile'):
        messages.error(request, 'Solo los profesores pueden descargar entregas.')
        return redirect('core:dashboard')
    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada.')
        return redirect('core:dashboard')
    prof = request.user.profesor_profile
    entrega = get_object_or_404(
        EntregaActividad.objects.select_related('actividad__escuela'),
        id=entrega_id,
        actividad_id=actividad_id,
        actividad__escuela_id=escuela_id,
        sede=user_sede,
        actividad__escuela__maestro=prof,
    )
    if not entrega.archivo or not entrega.archivo.name:
        raise Http404('Esta entrega no tiene archivo.')
    try:
        file_handle = entrega.archivo.open('rb')
    except FileNotFoundError:
        raise Http404('El archivo ya no está en el servidor.')
    filename = os.path.basename(entrega.archivo.name)
    return FileResponse(
        file_handle,
        as_attachment=True,
        filename=filename or 'entrega',
    )


@login_required
def escuela_archivos(request, escuela_id):
    """Subir y administrar materiales de la escuela (docente asignado)."""
    if not hasattr(request.user, 'profesor_profile'):
        messages.error(request, 'Solo los profesores pueden acceder.')
        return redirect('core:dashboard')
    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada. Contacta al administrador.')
        return redirect('core:dashboard')
    prof = request.user.profesor_profile
    escuela = get_object_or_404(
        Escuela,
        id=escuela_id,
        sede=user_sede,
        maestro=prof,
        is_active=True,
    )
    volver_url = f"{reverse('hechos:mis_escuelas_profesor')}?escuela={escuela.id}"

    if request.method == 'POST':
        action = (request.POST.get('action') or '').strip()
        if action == 'eliminar':
            ar = get_object_or_404(
                ArchivoRecursoEscuela,
                id=request.POST.get('archivo_id'),
                escuela=escuela,
                sede=user_sede,
            )
            ar.is_active = False
            ar.save()
            messages.success(request, f'Se quitó «{ar.titulo}» de los recursos visibles.')
            return redirect('hechos:escuela_archivos', escuela_id=escuela.id)

        form = ArchivoRecursoEscuelaForm(request.POST, request.FILES)
        if form.is_valid():
            obj = form.save(commit=False)
            obj.sede = user_sede
            obj.escuela = escuela
            obj.subido_por = prof
            try:
                mr = str(settings.MEDIA_ROOT)
                os.makedirs(mr, exist_ok=True)
                os.makedirs(os.path.join(mr, 'recursos_escuela'), exist_ok=True)
                obj.save()
            except OSError as exc:
                messages.error(
                    request,
                    f'No se pudo guardar el archivo en el servidor. Detalle: {exc}',
                )
            else:
                messages.success(request, f'Recurso «{obj.titulo}» publicado para los estudiantes.')
            return redirect('hechos:escuela_archivos', escuela_id=escuela.id)
        messages.error(request, 'Revisa el formulario: título y archivo son obligatorios.')
    else:
        form = ArchivoRecursoEscuelaForm()

    archivos = (
        ArchivoRecursoEscuela.objects.filter(escuela=escuela, sede=user_sede, is_active=True)
        .select_related('subido_por__user')
        .order_by('-created_at')
    )
    return render(
        request,
        'hechos/escuela_archivos_modern.html',
        {
            'escuela': escuela,
            'form': form,
            'archivos': archivos,
            'volver_url': volver_url,
            'user_sede': user_sede,
        },
    )


@login_required
def estudiante_escuela_recursos(request, escuela_id):
    """Lista de archivos de la escuela para estudiantes matriculados en algún curso de esa escuela."""
    if not hasattr(request.user, 'estudiante_profile'):
        messages.error(request, 'Solo los estudiantes pueden ver recursos de escuela.')
        return redirect('core:dashboard')
    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada.')
        return redirect('core:dashboard')
    estudiante = request.user.estudiante_profile
    escuela = get_object_or_404(
        Escuela,
        id=escuela_id,
        sede=user_sede,
        is_active=True,
    )
    if not _estudiante_matriculado_en_escuela(estudiante, escuela, user_sede):
        messages.error(request, 'No tienes acceso a los recursos de esta escuela.')
        return redirect('hechos:mis_escuelas')

    archivos = (
        ArchivoRecursoEscuela.objects.filter(escuela=escuela, sede=user_sede, is_active=True)
        .select_related('subido_por__user')
        .order_by('-created_at')
    )
    return render(
        request,
        'hechos/estudiante_escuela_recursos.html',
        {
            'escuela': escuela,
            'archivos': archivos,
            'user_sede': user_sede,
        },
    )


@login_required
def descargar_archivo_recurso_escuela(request, escuela_id, archivo_id):
    """Descarga segura: maestro de la escuela o estudiante matriculado en la escuela."""
    user_sede = get_user_sede(request.user)
    if not user_sede:
        raise Http404
    ar = get_object_or_404(
        ArchivoRecursoEscuela.objects.select_related('escuela'),
        id=archivo_id,
        escuela_id=escuela_id,
        sede=user_sede,
        is_active=True,
    )
    esc = ar.escuela
    permitido = False
    if hasattr(request.user, 'profesor_profile'):
        prof = request.user.profesor_profile
        if esc.maestro_id == prof.id and esc.sede_id == user_sede.id:
            permitido = True
    elif hasattr(request.user, 'estudiante_profile'):
        if _estudiante_matriculado_en_escuela(request.user.estudiante_profile, esc, user_sede):
            permitido = True
    if not permitido:
        raise Http404
    if not ar.archivo or not ar.archivo.name:
        raise Http404
    fname = os.path.basename(ar.archivo.name)
    try:
        fh = ar.archivo.open('rb')
    except FileNotFoundError:
        raise Http404 from None
    return FileResponse(
        fh,
        as_attachment=True,
        filename=fname or 'recurso',
    )


@login_required
def sesion_actividad(request, clase_id):
    """Reserva de pantalla: crear actividad ligada a la clase (pendiente de implementar)."""
    return _sesion_herramienta_placeholder(
        request,
        clase_id,
        titulo='Crear actividad',
        texto_ayuda=(
            'Aquí podrás crear actividades para esta clase (por ejemplo tareas o ejercicios). '
            'La función está en preparación.'
        ),
    )


@login_required
def sesion_archivos(request, clase_id):
    """Reserva de pantalla: materiales de la clase (pendiente de implementar)."""
    return _sesion_herramienta_placeholder(
        request,
        clase_id,
        titulo='Subir archivos',
        texto_ayuda=(
            'Aquí podrás subir archivos y materiales para los estudiantes de esta clase. '
            'La función está en preparación.'
        ),
    )


@login_required
def notas_estudiante(request):
    """
    Notas del estudiante
    """
    if not hasattr(request.user, 'estudiante_profile'):
        messages.error(request, 'Solo los estudiantes pueden ver sus notas.')
        return redirect('core:dashboard')
    
    # Obtener la sede del usuario
    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada. Contacta al administrador.')
        return redirect('core:dashboard')
    
    estudiante = request.user.estudiante_profile
    
    # Obtener matrículas activas
    matriculas_activas = Matricula.objects.filter(
        estudiante=estudiante,
        sede=user_sede,
        is_active=True
    ).select_related('escuela', 'escuela__maestro__user')
    
    # Obtener todas las notas del estudiante
    notas = (
        Nota.objects.filter(estudiante=estudiante, sede=user_sede)
        .select_related('escuela', 'profesor__user')
        .order_by('-fecha_evaluacion')
    )
    
    # Calcular promedio general
    promedio_general = 0
    if notas.exists():
        promedio_general = sum(nota.puntaje_obtenido for nota in notas) / notas.count()
    
    # Procesar cursos matriculados con sus datos
    cursos_matriculados = []
    cursos_aprobados = []
    cursos_en_progreso = []
    
    for matricula in matriculas_activas:
        curso = matricula.curso
        escuela_m = matricula.escuela

        notas_curso = notas.filter(escuela=escuela_m)
        
        # Calcular promedio del curso
        promedio_curso = 0
        if notas_curso.exists():
            promedio_curso = sum(nota.puntaje_obtenido for nota in notas_curso) / notas_curso.count()
        
        # Determinar estado del curso (aprobado con 3.5/5.0 o más)
        estado = 'en_progreso'
        if promedio_curso >= 3.5:  # 3.5/5.0 = 70%
            estado = 'aprobado'
            if curso:
                cursos_aprobados.append(curso)
        else:
            if curso:
                cursos_en_progreso.append(curso)
        
        # Obtener asistencia del curso
        from hechos.models import Asistencia
        total_asistencias = Asistencia.objects.filter(
            estudiante=estudiante,
            clase__escuela=escuela_m,
            sede=user_sede,
        ).count()

        asistencias_presentes = Asistencia.objects.filter(
            estudiante=estudiante,
            clase__escuela=escuela_m,
            sede=user_sede,
            estado='presente',
        ).count()
        
        porcentaje_asistencia = (asistencias_presentes / total_asistencias * 100) if total_asistencias > 0 else 0
        
        # Obtener clases del curso
        from hechos.models import Clase
        total_clases = Clase.objects.filter(escuela=escuela_m, sede=user_sede).count()
        
        curso_data = {
            'curso': curso,
            'matricula': matricula,
            'notas_curso': notas_curso,
            'promedio_curso': round(promedio_curso, 2),
            'estado': estado,
            'porcentaje_asistencia': round(porcentaje_asistencia, 1),
            'total_asistencias': total_asistencias,
            'asistencias_presentes': asistencias_presentes,
            'total_clases': total_clases,
            'progreso_curso': min(100, (asistencias_presentes / total_clases * 100)) if total_clases > 0 else 0
        }
        
        cursos_matriculados.append(curso_data)
    
    context = {
        'estudiante': estudiante,
        'notas': notas,
        'cursos_matriculados': cursos_matriculados,
        'cursos_aprobados': cursos_aprobados,
        'cursos_en_progreso': cursos_en_progreso,
        'promedio_general': round(promedio_general, 2),
        'user_sede': user_sede,
    }
    
    return render(request, 'hechos/notas_estudiante_modern.html', context)


@login_required
def notas_estudiante_ruta(request, ruta_id):
    """
    Notas del estudiante por ruta de estudio específica
    """
    if not hasattr(request.user, 'estudiante_profile'):
        messages.error(request, 'Solo los estudiantes pueden ver sus notas.')
        return redirect('core:dashboard')
    
    # Obtener la sede del usuario
    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada. Contacta al administrador.')
        return redirect('core:dashboard')
    
    estudiante = request.user.estudiante_profile
    
    # Obtener la ruta específica
    try:
        ruta = RutaEstudio.objects.get(id=ruta_id, sede=user_sede, is_active=True)
    except RutaEstudio.DoesNotExist:
        messages.error(request, 'La ruta de estudio no existe o no tienes acceso a ella.')
        return redirect('hechos:notas_estudiante')
    
    # Obtener cursos de la ruta
    cursos_ruta = Curso.objects.filter(
        ruta_estudio=ruta,
        sede=user_sede,
        is_active=True
    ).order_by('orden', 'nombre')
    
    # Obtener matrículas del estudiante en esta ruta
    matriculas_ruta = Matricula.objects.filter(
        estudiante=estudiante,
        escuela=ruta,
        sede=user_sede,
        is_active=True
    ).select_related('escuela', 'escuela__maestro__user')
    
    # Obtener todas las notas del estudiante en esta ruta
    notas_ruta = Nota.objects.filter(
        estudiante=estudiante,
        curso__ruta_estudio=ruta,
        sede=user_sede
    ).select_related('curso', 'profesor__user').order_by('-fecha_evaluacion')
    
    # Calcular promedio de la ruta
    promedio_ruta = 0
    if notas_ruta.exists():
        promedio_ruta = sum(nota.puntaje_obtenido for nota in notas_ruta) / notas_ruta.count()
    
    # Procesar cada curso de la ruta
    cursos_con_datos = []
    cursos_completados = 0
    cursos_en_progreso = 0
    cursos_pendientes = 0
    
    for curso in cursos_ruta:
        # Verificar si está matriculado
        matricula_curso = matriculas_ruta.filter(escuela=curso).first()
        
        if matricula_curso:
            # Obtener notas del curso
            notas_curso = notas_ruta.filter(curso=curso)
            
            # Calcular promedio del curso
            promedio_curso = 0
            if notas_curso.exists():
                promedio_curso = sum(nota.puntaje_obtenido for nota in notas_curso) / notas_curso.count()
            
            # Determinar estado del curso (completado con 3.5/5.0 o más)
            estado = 'en_progreso'
            if promedio_curso >= 3.5:  # 3.5/5.0 = 70%
                estado = 'completado'
                cursos_completados += 1
            else:
                cursos_en_progreso += 1
            
            # Obtener asistencia del curso
            total_asistencias = Asistencia.objects.filter(
                estudiante=estudiante,
                clase__escuela=curso,
                sede=user_sede
            ).count()
            
            asistencias_presentes = Asistencia.objects.filter(
                estudiante=estudiante,
                clase__escuela=curso,
                sede=user_sede,
                estado='presente'
            ).count()
            
            porcentaje_asistencia = (asistencias_presentes / total_asistencias * 100) if total_asistencias > 0 else 0
            
            # Obtener clases del curso
            total_clases = Clase.objects.filter(
                escuela=curso,
                sede=user_sede
            ).count()
            
            curso_data = {
                'curso': curso,
                'matricula': matricula_curso,
                'notas_curso': notas_curso,
                'promedio_curso': round(promedio_curso, 2),
                'estado': estado,
                'porcentaje_asistencia': round(porcentaje_asistencia, 1),
                'total_asistencias': total_asistencias,
                'asistencias_presentes': asistencias_presentes,
                'total_clases': total_clases,
                'progreso_curso': min(100, (asistencias_presentes / total_clases * 100)) if total_clases > 0 else 0
            }
            
            cursos_con_datos.append(curso_data)
        else:
            # Curso no matriculado
            curso_data = {
                'curso': curso,
                'matricula': None,
                'notas_curso': [],
                'promedio_curso': 0,
                'estado': 'pendiente',
                'porcentaje_asistencia': 0,
                'total_asistencias': 0,
                'asistencias_presentes': 0,
                'total_clases': 0,
                'progreso_curso': 0
            }
            
            cursos_con_datos.append(curso_data)
            cursos_pendientes += 1
    
    # Calcular progreso general de la ruta
    total_cursos = cursos_ruta.count()
    progreso_ruta = (cursos_completados / total_cursos * 100) if total_cursos > 0 else 0
    
    context = {
        'estudiante': estudiante,
        'ruta': ruta,
        'cursos_con_datos': cursos_con_datos,
        'notas_ruta': notas_ruta,
        'promedio_ruta': round(promedio_ruta, 2),
        'cursos_completados': cursos_completados,
        'cursos_en_progreso': cursos_en_progreso,
        'cursos_pendientes': cursos_pendientes,
        'progreso_ruta': round(progreso_ruta, 1),
        'user_sede': user_sede,
    }
    
    return render(request, 'hechos/notas_estudiante_ruta_modern.html', context)


# ===== VISTAS DE CREACIÓN PARA ADMIN DE SEDE =====

@login_required
def crear_estudiante(request):
    """
    Crear nuevo estudiante (solo para admins)
    """
    if not request.user.is_super_admin():
        r = ca.require_academico(request)
        if r:
            return r
    
    # Obtener la sede del usuario
    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada. Contacta al administrador.')
        return redirect('core:dashboard')
    
    if request.method == 'POST':
        try:
            avatar_file = request.FILES.get('avatar')
            # Crear usuario
            user = User.objects.create_user(
                username=request.POST.get('username'),
                email=request.POST.get('email'),
                password=request.POST.get('password'),
                first_name=request.POST.get('first_name'),
                last_name=request.POST.get('last_name'),
                sede=user_sede,
                avatar=avatar_file,
            )
            
            # Crear perfil de estudiante
            Estudiante.objects.create(
                user=user,
                sede=user_sede,
                telefono_emergencia=request.POST.get('telefono_emergencia', ''),
                contacto_emergencia=request.POST.get('contacto_emergencia', ''),
                direccion=request.POST.get('direccion', ''),
                iglesia=request.POST.get('iglesia', ''),
                fecha_nacimiento=request.POST.get('fecha_nacimiento') or None,
                notas_medicas=request.POST.get('notas_medicas', '')
            )
            
            messages.success(request, f'Estudiante {user.get_full_name()} creado exitosamente.')
            return redirect('hechos:estudiantes_list')
            
        except Exception as e:
            messages.error(request, f'Error al crear estudiante: {str(e)}')
    
    context = {
        'user_sede': user_sede,
    }
    
    return render(request, 'hechos/crear_estudiante.html', context)


@login_required
def crear_profesor(request):
    """
    Crear nuevo profesor (solo para admins)
    """
    if not request.user.is_super_admin():
        r = ca.require_pedagogico(request)
        if r:
            return r
    
    # Obtener la sede del usuario
    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada. Contacta al administrador.')
        return redirect('core:dashboard')
    
    if request.method == 'POST':
        try:
            email = (request.POST.get('email') or '').strip()
            username = generate_unique_username(User, email)
            # Crear usuario
            user = User.objects.create_user(
                username=username,
                email=email,
                password=request.POST.get('password'),
                first_name=request.POST.get('first_name'),
                last_name=request.POST.get('last_name'),
                sede=user_sede,
                avatar=request.FILES.get('avatar'),
            )
            tel = (request.POST.get('telefono') or '').strip()
            if tel:
                user.phone = tel
                user.save(update_fields=['phone'])
            
            # Crear perfil de profesor
            Profesor.objects.create(
                user=user,
                sede=user_sede,
            )
            
            messages.success(request, f'Profesor {user.get_full_name()} creado exitosamente.')
            return redirect('hechos:profesores_list')
            
        except Exception as e:
            messages.error(request, f'Error al crear profesor: {str(e)}')
    
    context = {
        'user_sede': user_sede,
    }
    
    return render(request, 'hechos/crear_profesor.html', context)


@login_required
def crear_ruta_estudio(request):
    """
    Crear nueva ruta de estudio (solo para admins)
    """
    if not request.user.is_super_admin():
        r = ca.require_academico_o_pedagogico(request)
        if r:
            return r
    
    # Obtener la sede del usuario
    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada. Contacta al administrador.')
        return redirect('core:dashboard')
    
    if request.method == 'POST':
        try:
            RutaEstudio.objects.create(
                nombre=request.POST.get('nombre'),
                descripcion=request.POST.get('descripcion', ''),
                duracion_semanas=request.POST.get('duracion_semanas', 0) or 0,
                nivel=request.POST.get('nivel'),
                sede=user_sede
            )
            
            messages.success(request, f'Ruta de estudio "{request.POST.get("nombre")}" creada exitosamente.')
            return redirect('hechos:rutas_estudio_list')
            
        except Exception as e:
            messages.error(request, f'Error al crear ruta de estudio: {str(e)}')
    
    context = {
        'user_sede': user_sede,
        'nivel_choices': [
            ('basico', 'Básico'),
            ('intermedio', 'Intermedio'),
            ('avanzado', 'Avanzado'),
        ],
    }
    
    return render(request, 'hechos/crear_ruta_estudio.html', context)


@login_required
def crear_curso(request):
    """
    Crear nuevo curso (solo para admins)
    """
    if not request.user.is_super_admin():
        r = ca.require_academico_o_pedagogico(request)
        if r:
            return r
    
    # Obtener la sede del usuario
    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada. Contacta al administrador.')
        return redirect('core:dashboard')
    
    # Obtener rutas de estudio y profesores disponibles
    rutas = RutaEstudio.objects.filter(sede=user_sede, is_active=True)
    profesores = Profesor.objects.filter(sede=user_sede, is_active=True)
    
    if request.method == 'POST':
        try:
            ruta_id = request.POST.get('ruta_estudio')
            profesor_id = request.POST.get('profesor')
            
            ruta = RutaEstudio.objects.get(id=ruta_id, sede=user_sede)
            profesor = Profesor.objects.get(id=profesor_id, sede=user_sede)
            
            Curso.objects.create(
                nombre=request.POST.get('nombre'),
                descripcion=request.POST.get('descripcion', ''),
                ruta_estudio=ruta,
                profesor=profesor,
                orden=request.POST.get('orden', 1) or 1,
                fecha_inicio=request.POST.get('fecha_inicio'),
                fecha_fin=request.POST.get('fecha_fin'),
                sede=user_sede
            )
            
            messages.success(request, f'Curso "{request.POST.get("nombre")}" creado exitosamente.')
            return redirect('hechos:cursos_list')
            
        except Exception as e:
            messages.error(request, f'Error al crear curso: {str(e)}')
    
    context = {
        'user_sede': user_sede,
        'rutas': rutas,
        'profesores': profesores,
    }
    
    return render(request, 'hechos/crear_curso.html', context)


def _profesor_puede_gestionar_escuela(profesor, escuela: Escuela | None) -> bool:
    """Docente titular de la escuela operativa."""
    return bool(escuela and escuela.maestro_id == profesor.id)


def _normalize_numero_documento_matricula(raw: str) -> str | None:
    """
    Número/código de identificación (CC, TI, CE, pasaporte, etc.).
    Quita espacios y separadores; conserva letras y dígitos (Unicode) para alinear con lo que suele guardarse en BD.
    Mínimo 3, máximo 15 caracteres (cubre cédula, TI, CE y pasaportes habituales).
    """
    s = (raw or '').strip()
    if not s:
        return None
    s = re.sub(r'\s+', '', s)
    # Quitar prefijos típicos al pegar desde otros sistemas
    s = re.sub(
        r'(?i)^(?:doc(?:umento)?|c\.?c\.?|t\.?i\.?|c\.?e\.?|pas(?:aporte)?|no\.?|n[°º#]?)\s*:?\s*',
        '',
        s,
    )
    s = ''.join(c for c in s if c.isalnum())
    if not s:
        return None
    if len(s) < 3 or len(s) > 15:
        return None
    return s


def _username_unico_desde_documento(doc_norm: str) -> str:
    """Username válido y único; el inicio de sesión puede no coincidir 1:1 con el documento si hay caracteres raros."""
    base = ''.join(c for c in doc_norm if c.isalnum() or c in '._@+-')
    if len(base) < 3:
        base = re.sub(r'[^a-zA-Z0-9]', '', doc_norm)[:30] or 'est'
    base = base[:150]
    if not User.objects.filter(username=base).exists():
        return base
    n = 1
    while n < 10_000:
        suffix = f'_{n}'
        cand = (base[: 150 - len(suffix)] + suffix)
        if not User.objects.filter(username=cand).exists():
            return cand
        n += 1
    return f'{base[:120]}_{secrets.token_hex(4)}'


def _find_estudiante_por_documento_global(doc_raw: str):
    """
    Busca un Estudiante en cualquier sede por número de documento (cualquier tipo).
    Incluye compatibilidad con registros antiguos guardados solo como dígitos (import).
    """
    from hechos.seguimiento_import import _normalize_doc_username, _normalize_documento

    doc_norm = _normalize_numero_documento_matricula(doc_raw)
    if doc_norm:
        u = User.objects.filter(documento_identidad__iexact=doc_norm).first()
        if u and hasattr(u, 'estudiante_profile'):
            return u.estudiante_profile
        est = (
            Estudiante.objects.filter(numero_documento__iexact=doc_norm)
            .select_related('user', 'sede')
            .first()
        )
        if est:
            return est

    legacy_digits = _normalize_doc_username(doc_raw)
    if legacy_digits and legacy_digits != (doc_norm or ''):
        u = User.objects.filter(documento_identidad__iexact=legacy_digits).first()
        if u and hasattr(u, 'estudiante_profile'):
            return u.estudiante_profile
        est = (
            Estudiante.objects.filter(numero_documento__iexact=legacy_digits)
            .select_related('user', 'sede')
            .first()
        )
        if est:
            return est

    doc = _normalize_documento(doc_raw)
    if len(doc) < 3:
        return None
    qs = Estudiante.objects.select_related('user', 'sede')
    est = qs.filter(numero_documento=doc).first()
    if est:
        return est
    est = qs.filter(numero_documento__iexact=doc).first()
    if est:
        return est
    digits = ''.join(c for c in doc if c.isdigit())
    if len(digits) < 5:
        return None
    for e in qs.iterator(chunk_size=800):
        ed = ''.join(c for c in (e.numero_documento or '') if c.isdigit())
        if ed == digits:
            return e
    return None


def _safe_matricula_asistencia_return_path(path: str, escuela_id: int) -> str | None:
    """Permite volver al panel de asistencia o de notas de la misma escuela (evita redirección abierta)."""
    if not path or not isinstance(path, str):
        return None
    path = path.strip()
    if not path.startswith('/') or path.startswith('//'):
        return None
    path_only = path.split('?', 1)[0]
    try:
        match = resolve(path_only)
    except Exception:
        return None
    if match.url_name not in (
        'profesor_escuela_asistencia_panel',
        'profesor_escuela_notas_panel',
    ):
        return None
    try:
        if int(match.kwargs.get('escuela_id', -1)) != int(escuela_id):
            return None
    except (TypeError, ValueError):
        return None
    return path_only


def _matricular_embed_return_targets(request, escuela_id: int) -> tuple[bool, str | None]:
    if request.POST.get('matricular_embed') != '1':
        return False, None
    raw = (request.POST.get('matricular_return_url') or '').strip()
    safe = _safe_matricula_asistencia_return_path(raw, escuela_id)
    return True, safe


def _matricular_simple_redirect_after_error(request, escuela):
    embed_active, ret_path = _matricular_embed_return_targets(request, escuela.id)
    if embed_active and ret_path:
        q = urlencode(
            {'escuela': str(escuela.id), 'embed': '1', 'return': ret_path}
        )
        return redirect(f"{reverse('hechos:matricular_estudiante')}?{q}")
    return redirect(f"{reverse('hechos:matricular_estudiante')}?escuela={escuela.id}")


def _matricular_simple_redirect_after_success(request, escuela):
    embed_active, ret_path = _matricular_embed_return_targets(request, escuela.id)
    if embed_active and ret_path:
        return redirect(ret_path)
    return redirect(f"{reverse('hechos:mis_escuelas_profesor')}?escuela={escuela.id}")


def _matricula_linea_curso_mensaje(escuela: Escuela | None, curso: Curso | None = None) -> str:
    """
    Texto para mensajes de matrícula: escuela operativa y, si aporta contexto, el curso.
    Si el nombre del curso es solo «Módulo N», no lo repetimos.
    """
    esc_n = (escuela.nombre or '').strip() if escuela else ''
    cur_n = (curso.nombre or '').strip() if curso else ''
    mod_solo = re.compile(r'^(?i)m[oó]dulo\s*\d+\s*$')
    es_etiqueta_modulo = bool(cur_n and mod_solo.match(cur_n.strip()))
    if esc_n:
        if cur_n and not es_etiqueta_modulo:
            return f'{esc_n} · {cur_n}'
        return esc_n
    if cur_n:
        return cur_n
    return 'la escuela'


def _matricular_escuela_simple_post(request, user_sede):
    """
    Matrícula rápida desde ?escuela=: documento + nombre (si no existe el usuario).
    El tipo de documento (CC/TI/CE/etc.) es solo referencia en BD; aquí manda el número/código.
    Devuelve HttpResponse o None si no aplica este formulario.
    """
    from hechos.seguimiento_import import (
        _grant_hechos_permission,
        _normalize_doc_username,
        _split_nombre_import,
    )

    if request.POST.get('matricula_escuela_simple') != '1':
        return None
    if not hasattr(request.user, 'profesor_profile'):
        messages.error(request, 'Este formulario solo está disponible para el docente de la escuela.')
        return redirect('hechos:matricular_estudiante')

    post_escuela_id = request.POST.get('escuela_matricular')
    if not post_escuela_id:
        messages.error(request, 'Falta la escuela.')
        return redirect('hechos:matricular_estudiante')

    prof = request.user.profesor_profile
    escuela = get_object_or_404(
        Escuela,
        id=post_escuela_id,
        sede=user_sede,
        maestro=prof,
        is_active=True,
    )

    doc_raw = (request.POST.get('numero_documento') or request.POST.get('cedula') or '').strip()
    doc_norm = _normalize_numero_documento_matricula(doc_raw)
    if not doc_norm:
        messages.error(
            request,
            'Número de documento no válido. Usa entre 3 y 15 caracteres (letras y números; '
            'se ignoran espacios, puntos y guiones). Válido para CC, TI, CE, pasaporte, etc.',
        )
        return _matricular_simple_redirect_after_error(request, escuela)

    if not _profesor_puede_gestionar_escuela(prof, escuela):
        messages.error(request, 'No puedes matricular en esta escuela.')
        return _matricular_simple_redirect_after_error(request, escuela)

    if escuela.estudiantes_inscritos >= escuela.cupo_maximo:
        messages.warning(request, f'No hay cupo disponible en «{escuela.nombre}».')
        return _matricular_simple_redirect_after_error(request, escuela)

    curso_ref = (
        Curso.objects.filter(ruta_estudio__escuela=escuela, sede=user_sede, is_active=True)
        .order_by('orden', 'id')
        .first()
    )

    nombre_completo = (request.POST.get('nombre_completo') or '').strip()
    est_existente = _find_estudiante_por_documento_global(doc_raw)
    if not est_existente and not nombre_completo:
        messages.error(
            request,
            'Ese documento no está en el sistema. Escribe el nombre completo para registrar a la persona.',
        )
        return _matricular_simple_redirect_after_error(request, escuela)

    password_plain = ''
    mat_created = False

    try:
        with transaction.atomic():
            est = _find_estudiante_por_documento_global(doc_raw)
            if est:
                if not est.is_active:
                    est.is_active = True
                    est.save(update_fields=['is_active'])
                user = est.user
                if est.sede_id != user_sede.id:
                    est.sede = user_sede
                    est.save(update_fields=['sede'])
                if user.sede_id != user_sede.id:
                    user.sede = user_sede
                    user.save(update_fields=['sede'])
                if nombre_completo:
                    fn, ln = _split_nombre_import(nombre_completo)
                    user.first_name = fn
                    user.last_name = ln
                    user.save(update_fields=['first_name', 'last_name'])
                if not est.numero_documento or est.numero_documento != doc_norm:
                    est.numero_documento = doc_norm
                    est.save(update_fields=['numero_documento'])
                if not user.documento_identidad:
                    user.documento_identidad = doc_norm
                    user.save(update_fields=['documento_identidad'])
                _grant_hechos_permission(user)
            else:
                first_name, last_name = _split_nombre_import(nombre_completo)
                legacy_u = _normalize_doc_username(doc_raw)
                existing_user = User.objects.filter(documento_identidad__iexact=doc_norm).first()
                if existing_user is None:
                    existing_user = User.objects.filter(username__iexact=doc_norm).first()
                if existing_user is None and legacy_u and legacy_u != doc_norm:
                    existing_user = User.objects.filter(documento_identidad__iexact=legacy_u).first()
                if existing_user is None and legacy_u and legacy_u != doc_norm:
                    existing_user = User.objects.filter(username__iexact=legacy_u).first()

                if existing_user is not None:
                    if hasattr(existing_user, 'estudiante_profile'):
                        est = existing_user.estudiante_profile
                        est.sede_id = user_sede.pk
                        est.numero_documento = doc_norm
                        est.is_active = True
                        est.save(update_fields=['sede', 'numero_documento', 'is_active'])
                        user = existing_user
                        user.sede = user_sede
                        if not user.documento_identidad:
                            user.documento_identidad = doc_norm
                        user.first_name = first_name
                        user.last_name = last_name
                        user.save(
                            update_fields=['sede', 'documento_identidad', 'first_name', 'last_name']
                        )
                    else:
                        user = existing_user
                        user.sede = user_sede
                        if not user.documento_identidad:
                            user.documento_identidad = doc_norm
                        user.first_name = first_name
                        user.last_name = last_name
                        user.save(
                            update_fields=['sede', 'documento_identidad', 'first_name', 'last_name']
                        )
                        est = Estudiante.objects.create(
                            user=user,
                            sede=user_sede,
                            tipo_documento=Estudiante.TipoDocumento.CC,
                            numero_documento=doc_norm,
                            is_active=True,
                        )
                        password_plain = doc_norm
                        user.set_password(password_plain)
                        user.save(update_fields=['password'])
                    _grant_hechos_permission(user)
                else:
                    # Contraseña inicial = mismo número de documento (estándar de la sede)
                    password_plain = doc_norm
                    user = User(
                        username=_username_unico_desde_documento(doc_norm),
                        first_name=first_name,
                        last_name=last_name,
                        email=None,
                        role='user',
                        sede=user_sede,
                        tipo_documento=User.TipoDocumento.CC,
                        documento_identidad=doc_norm,
                        is_active=True,
                    )
                    user.set_password(password_plain)
                    user.save()
                    est = Estudiante.objects.create(
                        user=user,
                        sede=user_sede,
                        tipo_documento=Estudiante.TipoDocumento.CC,
                        numero_documento=doc_norm,
                        is_active=True,
                    )
                    _grant_hechos_permission(user)

            mat, mat_created = Matricula.objects.get_or_create(
                estudiante=est,
                escuela=escuela,
                defaults={
                    'sede': user_sede,
                    'estado': 'activa',
                    'is_active': True,
                    'periodo': get_current_period(),
                },
            )
            if not mat_created and not mat.is_active:
                mat.is_active = True
                mat.estado = 'activa'
                mat.save(update_fields=['is_active', 'estado'])
    except Exception as e:
        messages.error(request, f'No se pudo completar la matrícula: {e}')
        return _matricular_simple_redirect_after_error(request, escuela)

    if not mat_created:
        messages.info(
            request,
            f'{est.user.get_full_name()} ya estaba matriculado en '
            f'{_matricula_linea_curso_mensaje(escuela, curso_ref)}.',
        )
    else:
        msg = (
            f'{est.user.get_full_name()} quedó matriculado en '
            f'{_matricula_linea_curso_mensaje(escuela, curso_ref)}.'
        )
        if password_plain:
            msg += (
                f' Documento para iniciar sesión: {est.user.username}. '
                f'Contraseña inicial (entregar al estudiante): {password_plain}'
            )
        messages.success(request, msg)
    return _matricular_simple_redirect_after_success(request, escuela)


@login_required
def buscar_estudiante_doc_matricula(request):
    """JSON: buscar estudiante por cédula en todo el sistema (matrícula rápida por escuela)."""
    if request.method != 'GET':
        return JsonResponse({'ok': False, 'error': 'Método no permitido'}, status=405)
    if not hasattr(request.user, 'profesor_profile'):
        return JsonResponse({'ok': False, 'error': 'No autorizado'}, status=403)
    user_sede = get_user_sede(request.user)
    escuela_id = request.GET.get('escuela')
    doc_raw = request.GET.get('doc', '')
    if not user_sede or not escuela_id:
        return JsonResponse({'ok': False, 'error': 'Faltan datos'}, status=400)
    if not Escuela.objects.filter(
        id=escuela_id,
        sede=user_sede,
        maestro=request.user.profesor_profile,
        is_active=True,
    ).exists():
        return JsonResponse({'ok': False, 'error': 'Escuela no válida'}, status=403)

    doc_norm = _normalize_numero_documento_matricula(doc_raw)
    if not doc_norm:
        return JsonResponse({'ok': True, 'found': False, 'invalid_doc': True, 'nombre': None, 'sede_registro': None})

    est = _find_estudiante_por_documento_global(doc_raw)
    if est:
        sede_nom = est.sede.nombre if getattr(est, 'sede_id', None) else None
        return JsonResponse(
            {
                'ok': True,
                'found': True,
                'invalid_doc': False,
                'nombre': est.user.get_full_name(),
                'estudiante_id': est.id,
                'sede_registro': sede_nom,
            }
        )
    return JsonResponse({'ok': True, 'found': False, 'invalid_doc': False, 'nombre': None, 'sede_registro': None})


@login_required
def sugerencias_estudiante_doc_matricula(request):
    """
    JSON: lista de estudiantes cuyo documento empieza por el prefijo dado (matrícula rápida).
    Mínimo 3 caracteres alfanuméricos normalizados; máximo 20 resultados.
    """
    if request.method != 'GET':
        return JsonResponse({'ok': False, 'error': 'Método no permitido'}, status=405)
    if not hasattr(request.user, 'profesor_profile'):
        return JsonResponse({'ok': False, 'error': 'No autorizado'}, status=403)
    user_sede = get_user_sede(request.user)
    escuela_id = request.GET.get('escuela')
    q_raw = request.GET.get('q', '')
    if not user_sede or not escuela_id:
        return JsonResponse({'ok': False, 'error': 'Faltan datos'}, status=400)
    if not Escuela.objects.filter(
        id=escuela_id,
        sede=user_sede,
        maestro=request.user.profesor_profile,
        is_active=True,
    ).exists():
        return JsonResponse({'ok': False, 'error': 'Escuela no válida'}, status=403)

    q_norm = _normalize_numero_documento_matricula(q_raw)
    if not q_norm:
        return JsonResponse({'ok': True, 'q': '', 'results': [], 'truncated': False, 'invalid_q': True})

    rows = list(
        Estudiante.objects.filter(is_active=True)
        .filter(
            Q(numero_documento__istartswith=q_norm)
            | Q(user__documento_identidad__istartswith=q_norm)
        )
        .select_related('user', 'sede')
        .distinct()
        .order_by('numero_documento', 'id')[:21]
    )
    truncated = len(rows) > 20
    rows = rows[:20]
    out = []
    for est in rows:
        doc = (est.numero_documento or '').strip() or (est.user.documento_identidad or '').strip()
        sede_nom = est.sede.nombre if getattr(est, 'sede_id', None) else None
        out.append(
            {
                'estudiante_id': est.id,
                'nombre': est.user.get_full_name() or est.user.username,
                'documento': doc,
                'sede_registro': sede_nom,
            }
        )
    return JsonResponse(
        {'ok': True, 'q': q_norm, 'results': out, 'truncated': truncated, 'invalid_q': False}
    )


@login_required
def matricular_estudiante(request):
    """
    Matricular estudiante (para admins y profesores).
    Profesores: cursos cuya escuela operativa tiene a ese docente como maestro.
    ?escuela=<id> (solo docente): acota cursos a esa escuela (debe ser su maestro asignado).
    """
    allowed = (
        request.user.is_super_admin()
        or hasattr(request.user, 'profesor_profile')
        or (
            hasattr(request.user, 'admin_escuela_profile')
            and ca.has_capacidad(request.user, 'academico')
        )
    )
    if not allowed:
        messages.error(request, 'No tienes permisos para matricular estudiantes.')
        return redirect('core:dashboard')

    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada. Contacta al administrador.')
        return redirect('core:dashboard')

    escuela_matricular = None
    escuela_matricular_id = request.GET.get('escuela') or request.POST.get('escuela_matricular')
    if escuela_matricular_id and hasattr(request.user, 'profesor_profile'):
        escuela_matricular = get_object_or_404(
            Escuela,
            id=escuela_matricular_id,
            sede=user_sede,
            maestro=request.user.profesor_profile,
            is_active=True,
        )
    elif escuela_matricular_id and not hasattr(request.user, 'profesor_profile'):
        messages.error(request, 'El filtro por escuela solo aplica al docente desde «Mis escuelas».')
        return redirect('hechos:matricular_estudiante')

    estudiantes = Estudiante.objects.filter(sede=user_sede, is_active=True).select_related('user')

    if hasattr(request.user, 'profesor_profile'):
        prof = request.user.profesor_profile
        if escuela_matricular:
            cursos = Curso.objects.filter(
                ruta_estudio__escuela=escuela_matricular,
                sede=user_sede,
                is_active=True,
            )
        else:
            cursos = (
                Curso.objects.filter(sede=user_sede, is_active=True)
                .filter(ruta_estudio__escuela__maestro=prof)
                .distinct()
            )
    else:
        cursos = Curso.objects.filter(sede=user_sede, is_active=True)

    if request.method == 'POST':
        simple_resp = _matricular_escuela_simple_post(request, user_sede)
        if simple_resp is not None:
            return simple_resp
        try:
            estudiante_id = request.POST.get('estudiante')
            curso_id = request.POST.get('curso')
            post_escuela_id = request.POST.get('escuela_matricular')

            estudiante = Estudiante.objects.get(id=estudiante_id, sede=user_sede)
            curso = Curso.objects.get(id=curso_id, sede=user_sede)
            escuela_op = curso.ruta_estudio.escuela
            if escuela_op is None:
                messages.error(request, 'El curso no está vinculado a una escuela operativa.')
                return redirect('hechos:matricular_estudiante')

            if hasattr(request.user, 'profesor_profile'):
                if not _profesor_puede_gestionar_escuela(request.user.profesor_profile, escuela_op):
                    messages.error(request, 'No puedes matricular en esta escuela.')
                    return redirect('hechos:matricular_estudiante')
                if post_escuela_id:
                    if str(escuela_op.id) != str(post_escuela_id):
                        messages.error(request, 'El curso no pertenece a la escuela indicada.')
                        return redirect('hechos:matricular_estudiante')
                    esc_chk = Escuela.objects.filter(
                        id=post_escuela_id,
                        sede=user_sede,
                        maestro=request.user.profesor_profile,
                        is_active=True,
                    ).exists()
                    if not esc_chk:
                        messages.error(request, 'Escuela no válida.')
                        return redirect('hechos:matricular_estudiante')

            if Matricula.objects.filter(
                estudiante=estudiante,
                escuela=escuela_op,
                sede=user_sede,
                is_active=True,
            ).exists():
                messages.warning(request, 'El estudiante ya está matriculado en esta escuela.')
            else:
                if escuela_op.estudiantes_inscritos >= escuela_op.cupo_maximo:
                    messages.warning(
                        request,
                        f'No hay cupo disponible en «{escuela_op.nombre}».',
                    )
                else:
                    periodo = get_current_period()
                    Matricula.objects.create(
                        estudiante=estudiante,
                        escuela=escuela_op,
                        sede=user_sede,
                        periodo=periodo,
                        fecha_matricula=request.POST.get('fecha_matricula') or timezone.now(),
                    )
                    messages.success(
                        request,
                        f'Estudiante {estudiante.user.get_full_name()} matriculado en '
                        f'{_matricula_linea_curso_mensaje(escuela_op, curso)} ({get_period_display(periodo)}).',
                    )
                    if post_escuela_id and hasattr(request.user, 'profesor_profile'):
                        return redirect(f"{reverse('hechos:mis_escuelas_profesor')}?escuela={post_escuela_id}")
                    return redirect('hechos:cursos_list')

        except Exception as e:
            messages.error(request, f'Error al matricular estudiante: {str(e)}')

    cursos_con_ediciones = [{'curso': curso, 'ediciones': []} for curso in cursos]
    prof = getattr(request.user, 'profesor_profile', None)

    volver_url = (
        f"{reverse('hechos:mis_escuelas_profesor')}?escuela={escuela_matricular.id}"
        if escuela_matricular
        else reverse('hechos:cursos_list')
    )

    matricula_escuela_simple = bool(escuela_matricular and prof)
    puede_matricula_escuela_simple = bool(
        matricula_escuela_simple
        and _profesor_puede_gestionar_escuela(prof, escuela_matricular)
        and escuela_matricular.cupos_disponibles > 0
    )

    matricular_embed = False
    matricular_return_url = ''
    matricular_cancel_url = volver_url
    if (
        request.GET.get('embed') == '1'
        and matricula_escuela_simple
        and escuela_matricular
    ):
        ret_raw = (request.GET.get('return') or '').strip()
        if not ret_raw:
            ret_raw = reverse(
                'hechos:profesor_escuela_asistencia_panel',
                args=[escuela_matricular.id],
            )
        safe_ret = _safe_matricula_asistencia_return_path(ret_raw, escuela_matricular.id)
        if safe_ret:
            matricular_embed = True
            matricular_return_url = safe_ret
            matricular_cancel_url = safe_ret

    context = {
        'user_sede': user_sede,
        'estudiantes': estudiantes,
        'cursos': cursos,
        'cursos_con_ediciones': cursos_con_ediciones,
        'escuela_matricular': escuela_matricular,
        'volver_url': volver_url,
        'matricula_escuela_simple': matricula_escuela_simple,
        'puede_matricula_escuela_simple': puede_matricula_escuela_simple,
        'buscar_doc_url': reverse('hechos:buscar_estudiante_doc_matricula'),
        'sugerencias_doc_url': reverse('hechos:sugerencias_estudiante_doc_matricula'),
        'matricular_embed': matricular_embed,
        'matricular_return_url': matricular_return_url,
        'matricular_cancel_url': matricular_cancel_url,
    }

    if matricular_embed:
        return render(request, 'hechos/matricular_estudiante_embed.html', context)
    return render(request, 'hechos/matricular_estudiante.html', context)


# ===== VISTAS DE GESTIÓN PARA ADMIN DE SEDE =====

@login_required
def editar_estudiante(request, estudiante_id):
    """
    Editar estudiante existente (solo para admins)
    """
    if not request.user.is_super_admin():
        r = ca.require_academico(request)
        if r:
            return r
    
    user_sede = get_user_sede(request.user)
    if request.user.is_super_admin():
        estudiante = get_object_or_404(Estudiante, id=estudiante_id)
    else:
        if not user_sede:
            messages.warning(request, 'No tienes una sede asignada. Contacta al administrador.')
            return redirect('core:dashboard')
        estudiante = get_object_or_404(Estudiante, id=estudiante_id, sede=user_sede)

    if request.method == 'POST':
        try:
            # Actualizar usuario
            estudiante.user.first_name = request.POST.get('first_name')
            estudiante.user.last_name = request.POST.get('last_name')
            estudiante.user.email = request.POST.get('email')
            if request.POST.get('password'):
                estudiante.user.set_password(request.POST.get('password'))
            estudiante.user.save()
            
            # Actualizar perfil de estudiante
            estudiante.telefono_emergencia = request.POST.get('telefono', '') or ''
            estudiante.direccion = request.POST.get('direccion', '')
            estudiante.fecha_nacimiento = request.POST.get('fecha_nacimiento') or None
            estudiante.save()
            
            messages.success(request, f'Estudiante {estudiante.user.get_full_name()} actualizado exitosamente.')
            return redirect('hechos:estudiantes_list')
            
        except Exception as e:
            messages.error(request, f'Error al actualizar estudiante: {str(e)}')
    
    context = {
        'estudiante': estudiante,
        'user_sede': user_sede,
    }
    
    return render(request, 'hechos/editar_estudiante.html', context)


@login_required
def editar_profesor(request, profesor_id):
    """
    Editar profesor existente (solo para admins)
    """
    if not request.user.is_super_admin():
        r = ca.require_pedagogico(request)
        if r:
            return r
    
    # Obtener la sede del usuario
    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada. Contacta al administrador.')
        return redirect('core:dashboard')
    
    profesor = get_object_or_404(Profesor, id=profesor_id, sede=user_sede)
    
    if request.method == 'POST':
        try:
            # Actualizar usuario
            profesor.user.first_name = request.POST.get('first_name')
            profesor.user.last_name = request.POST.get('last_name')
            profesor.user.email = request.POST.get('email')
            if request.POST.get('password'):
                profesor.user.set_password(request.POST.get('password'))
            tel = (request.POST.get('telefono') or '').strip()
            profesor.user.phone = tel
            profesor.user.save()

            messages.success(request, f'Profesor {profesor.user.get_full_name()} actualizado exitosamente.')
            return redirect('hechos:profesores_list')
            
        except Exception as e:
            messages.error(request, f'Error al actualizar profesor: {str(e)}')
    
    context = {
        'profesor': profesor,
        'user_sede': user_sede,
    }
    
    return render(request, 'hechos/editar_profesor.html', context)


@login_required
def editar_curso(request, curso_id):
    """
    Editar curso existente (solo para admins)
    """
    if not request.user.is_super_admin():
        r = ca.require_academico_o_pedagogico(request)
        if r:
            return r
    
    # Obtener la sede del usuario
    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada. Contacta al administrador.')
        return redirect('core:dashboard')
    
    curso = get_object_or_404(Curso, id=curso_id, sede=user_sede)
    rutas = RutaEstudio.objects.filter(sede=user_sede, is_active=True)
    profesores = Profesor.objects.filter(sede=user_sede, is_active=True)
    
    if request.method == 'POST':
        try:
            ruta_id = request.POST.get('ruta_estudio')
            profesor_id = request.POST.get('profesor')
            
            ruta = RutaEstudio.objects.get(id=ruta_id, sede=user_sede)
            profesor = Profesor.objects.get(id=profesor_id, sede=user_sede)
            
            curso.nombre = request.POST.get('nombre')
            curso.descripcion = request.POST.get('descripcion', '')
            curso.ruta_estudio = ruta
            curso.profesor = profesor
            curso.orden = request.POST.get('orden', 1) or 1
            curso.fecha_inicio = request.POST.get('fecha_inicio')
            curso.fecha_fin = request.POST.get('fecha_fin')
            curso.save()
            
            messages.success(request, f'Curso "{curso.nombre}" actualizado exitosamente.')
            return redirect('hechos:cursos_list')
            
        except Exception as e:
            messages.error(request, f'Error al actualizar curso: {str(e)}')
    
    context = {
        'curso': curso,
        'rutas': rutas,
        'profesores': profesores,
        'user_sede': user_sede,
    }
    
    return render(request, 'hechos/editar_curso.html', context)


@login_required
def editar_ruta_estudio(request, ruta_id):
    """
    Editar ruta de estudio existente (solo para admins)
    """
    if not request.user.is_super_admin():
        r = ca.require_academico_o_pedagogico(request)
        if r:
            return r
    
    # Obtener la sede del usuario
    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada. Contacta al administrador.')
        return redirect('core:dashboard')
    
    ruta = get_object_or_404(RutaEstudio, id=ruta_id, sede=user_sede)
    
    if request.method == 'POST':
        try:
            ruta.nombre = request.POST.get('nombre')
            ruta.descripcion = request.POST.get('descripcion', '')
            ruta.duracion_semanas = request.POST.get('duracion_semanas', 0) or 0
            ruta.nivel = request.POST.get('nivel')
            ruta.save()
            
            messages.success(request, f'Ruta de estudio "{ruta.nombre}" actualizada exitosamente.')
            return redirect('hechos:rutas_estudio_list')
            
        except Exception as e:
            messages.error(request, f'Error al actualizar ruta de estudio: {str(e)}')
    
    context = {
        'ruta': ruta,
        'user_sede': user_sede,
        'nivel_choices': [
            ('basico', 'Básico'),
            ('intermedio', 'Intermedio'),
            ('avanzado', 'Avanzado'),
        ],
    }
    
    return render(request, 'hechos/editar_ruta_estudio.html', context)


@login_required
def eliminar_estudiante(request, estudiante_id):
    """
    Eliminar estudiante (soft delete)
    """
    if not request.user.is_super_admin():
        r = ca.require_academico(request)
        if r:
            return r
    
    user_sede = get_user_sede(request.user)
    if request.user.is_super_admin():
        estudiante = get_object_or_404(Estudiante, id=estudiante_id)
    else:
        if not user_sede:
            messages.warning(request, 'No tienes una sede asignada. Contacta al administrador.')
            return redirect('core:dashboard')
        estudiante = get_object_or_404(Estudiante, id=estudiante_id, sede=user_sede)

    if request.method == 'POST':
        estudiante.is_active = False
        estudiante.save()
        messages.success(request, f'Estudiante {estudiante.user.get_full_name()} eliminado exitosamente.')
        return redirect('hechos:estudiantes_list')
    
    context = {
        'estudiante': estudiante,
        'user_sede': user_sede,
    }
    
    return render(request, 'hechos/eliminar_estudiante.html', context)


@login_required
def eliminar_profesor(request, profesor_id):
    """
    Eliminar profesor (soft delete)
    """
    if not request.user.is_super_admin():
        r = ca.require_pedagogico(request)
        if r:
            return r
    
    # Obtener la sede del usuario
    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada. Contacta al administrador.')
        return redirect('core:dashboard')
    
    profesor = get_object_or_404(Profesor, id=profesor_id, sede=user_sede)
    
    if request.method == 'POST':
        profesor.is_active = False
        profesor.save()
        messages.success(request, f'Profesor {profesor.user.get_full_name()} eliminado exitosamente.')
        return redirect('hechos:profesores_list')
    
    context = {
        'profesor': profesor,
        'user_sede': user_sede,
    }
    
    return render(request, 'hechos/eliminar_profesor.html', context)


@login_required
def eliminar_curso(request, curso_id):
    """
    Eliminar curso (soft delete)
    """
    if not request.user.is_super_admin():
        r = ca.require_academico_o_pedagogico(request)
        if r:
            return r
    
    # Obtener la sede del usuario
    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada. Contacta al administrador.')
        return redirect('core:dashboard')
    
    curso = get_object_or_404(Curso, id=curso_id, sede=user_sede)
    
    if request.method == 'POST':
        curso.is_active = False
        curso.save()
        messages.success(request, f'Curso "{curso.nombre}" eliminado exitosamente.')
        return redirect('hechos:cursos_list')
    
    context = {
        'curso': curso,
        'user_sede': user_sede,
    }
    
    return render(request, 'hechos/eliminar_curso.html', context)


@login_required
def eliminar_ruta_estudio(request, ruta_id):
    """
    Eliminar ruta de estudio (soft delete)
    """
    if not request.user.is_super_admin():
        r = ca.require_academico_o_pedagogico(request)
        if r:
            return r
    
    # Obtener la sede del usuario
    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada. Contacta al administrador.')
        return redirect('core:dashboard')
    
    ruta = get_object_or_404(RutaEstudio, id=ruta_id, sede=user_sede)
    
    if request.method == 'POST':
        ruta.is_active = False
        ruta.save()
        messages.success(request, f'Ruta de estudio "{ruta.nombre}" eliminada exitosamente.')
        return redirect('hechos:rutas_estudio_list')
    
    context = {
        'ruta': ruta,
        'user_sede': user_sede,
    }
    
    return render(request, 'hechos/eliminar_ruta_estudio.html', context)


@login_required
def detalle_estudiante(request, estudiante_id):
    """
    Ficha del estudiante en cualquier sede: matrículas y notas en toda la plataforma.
    Mismo criterio de acceso que el directorio global (por-sede).
    """
    allowed = (
        request.user.is_super_admin()
        or hasattr(request.user, 'profesor_profile')
        or (
            hasattr(request.user, 'admin_escuela_profile')
            and ca.has_capacidad(request.user, 'academico')
        )
    )
    if not allowed:
        messages.error(request, 'No tienes permisos para ver esta ficha.')
        return redirect('core:dashboard')

    user_sede = get_user_sede(request.user)
    estudiante = get_object_or_404(
        Estudiante.objects.select_related('user', 'sede'),
        id=estudiante_id,
    )
    matriculas = (
        Matricula.objects.filter(estudiante=estudiante, is_active=True)
        .select_related(
            'sede',
            'escuela__maestro__user',
            'escuela',
            'escuela__sede',
        )
        .order_by('-fecha_matricula', 'id')
    )
    notas = (
        Nota.objects.filter(estudiante=estudiante)
        .select_related('curso', 'curso__ruta_estudio__escuela', 'curso__sede', 'sede', 'profesor__user')
        .order_by('-fecha_evaluacion')[:40]
    )

    context = {
        'estudiante': estudiante,
        'matriculas': matriculas,
        'notas': notas,
        'user_sede': user_sede,
    }

    return render(request, 'hechos/detalle_estudiante.html', context)


@login_required
def estudiante_resumen_fragment(request, estudiante_id):
    """
    HTML parcial para modal (misma política de acceso que detalle_estudiante).
    """
    allowed = (
        request.user.is_super_admin()
        or hasattr(request.user, 'profesor_profile')
        or (
            hasattr(request.user, 'admin_escuela_profile')
            and ca.has_capacidad(request.user, 'academico')
        )
    )
    if not allowed:
        return HttpResponse('No autorizado', status=403, content_type='text/plain; charset=utf-8')

    estudiante = get_object_or_404(
        Estudiante.objects.select_related('user', 'sede'),
        id=estudiante_id,
    )
    matriculas = (
        Matricula.objects.filter(estudiante=estudiante, is_active=True)
        .select_related(
            'sede',
            'escuela__maestro__user',
            'escuela',
            'escuela__sede',
        )
        .order_by('-fecha_matricula', 'id')
    )
    notas = (
        Nota.objects.filter(estudiante=estudiante)
        .select_related('escuela', 'escuela__sede', 'sede', 'profesor__user')
        .order_by('-fecha_evaluacion')[:40]
    )
    return render(
        request,
        'hechos/estudiante_resumen_fragment.html',
        {
            'estudiante': estudiante,
            'matriculas': matriculas,
            'notas': notas,
        },
    )


@login_required
def detalle_profesor(request, profesor_id):
    """
    Ver detalles completos del profesor
    """
    if not request.user.is_super_admin():
        r = ca.require_pedagogico(request)
        if r:
            return r
    
    # Obtener la sede del usuario
    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada. Contacta al administrador.')
        return redirect('core:dashboard')
    
    profesor = get_object_or_404(Profesor, id=profesor_id, sede=user_sede)
    cursos = Curso.objects.filter(
        ruta_estudio__escuela__maestro=profesor,
        sede=user_sede,
        is_active=True,
    ).distinct()
    clases = Clase.objects.filter(profesor=profesor, sede=user_sede, is_active=True).order_by('-fecha_clase')[:10]
    
    context = {
        'profesor': profesor,
        'cursos': cursos,
        'clases': clases,
        'user_sede': user_sede,
    }
    
    return render(request, 'hechos/detalle_profesor.html', context)


@login_required
def detalle_curso(request, curso_id):
    """
    Ver detalles completos del curso
    """
    if (
        hasattr(request.user, 'estudiante_profile')
        and request.resolver_match
        and request.resolver_match.url_name == 'detalle_curso'
    ):
        return redirect('hechos:detalle_mis_escuela', curso_id=curso_id)

    allowed = (
        request.user.is_super_admin()
        or hasattr(request.user, 'estudiante_profile')
        or hasattr(request.user, 'profesor_profile')
        or (
            hasattr(request.user, 'admin_escuela_profile')
            and ca.has_any_capacidad(request.user, ('academico', 'pedagogico'))
        )
    )
    if not allowed:
        messages.error(request, 'No tienes permisos para ver detalles de cursos.')
        return redirect('core:dashboard')
    
    # Obtener la sede del usuario
    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada. Contacta al administrador.')
        return redirect('core:dashboard')
    
    curso = get_object_or_404(Curso, id=curso_id, sede=user_sede, is_active=True)
    escuela_curso = getattr(curso.ruta_estudio, 'escuela', None)

    if hasattr(request.user, 'estudiante_profile'):
        estudiante = request.user.estudiante_profile
        if not escuela_curso:
            messages.error(request, 'No tienes acceso a este curso.')
            return redirect('hechos:mis_escuelas')
        matricula = Matricula.objects.filter(
            estudiante=estudiante,
            escuela=escuela_curso,
            sede=user_sede,
            is_active=True,
        ).select_related('escuela__maestro__user').first()
        if not matricula:
            messages.error(request, 'No tienes acceso a este curso.')
            return redirect('hechos:mis_escuelas')

        clases = Clase.objects.filter(
            escuela=escuela_curso,
            sede=user_sede,
            is_active=True,
        ).order_by('fecha_clase')
        notas_curso = Nota.objects.filter(
            escuela=escuela_curso,
            estudiante=estudiante,
            sede=user_sede,
        ).order_by('-fecha_evaluacion')
        promedio = notas_curso.aggregate(promedio=Avg('puntaje_obtenido'))['promedio'] or 0
        escuela = escuela_curso
        if escuela:
            actividades_qs = ActividadEscuela.objects.filter(
                escuela=escuela,
                sede=user_sede,
                is_active=True,
            ).order_by('-created_at')
            actividades_list = list(actividades_qs)
            act_ids = [a.id for a in actividades_list]
            entregas_map = {
                e.actividad_id: e
                for e in EntregaActividad.objects.filter(
                    estudiante=estudiante,
                    sede=user_sede,
                    actividad_id__in=act_ids,
                )
            }
            ahora = timezone.now()
            actividades_con_entrega = []
            for a in actividades_list:
                ent = entregas_map.get(a.id)
                plazo_vencido = bool(a.fecha_limite and ahora > a.fecha_limite)
                limite_ok = not plazo_vencido
                puede_entregar = limite_ok and (not ent or not ent.evaluado_en)
                actividades_con_entrega.append(
                    {
                        'actividad': a,
                        'entrega': ent,
                        'puede_entregar': puede_entregar,
                        'plazo_vencido': plazo_vencido,
                    }
                )
            actividades_escuela = actividades_qs
            recursos_con_archivo = actividades_qs.exclude(
                Q(material_adjunto='') | Q(material_adjunto__isnull=True)
            )
            archivos_recurso_escuela = list(
                ArchivoRecursoEscuela.objects.filter(
                    escuela=escuela,
                    sede=user_sede,
                    is_active=True,
                )
                .select_related('subido_por__user')
                .order_by('-created_at')
            )
        else:
            actividades_escuela = ActividadEscuela.objects.none()
            recursos_con_archivo = ActividadEscuela.objects.none()
            actividades_con_entrega = []
            archivos_recurso_escuela = []
        context = {
            'curso': curso,
            'matricula': matricula,
            'clases': clases,
            'notas_curso': notas_curso,
            'promedio_general': round(promedio, 2),
            'total_clases': clases.count(),
            'total_notas': notas_curso.count(),
            'user_sede': user_sede,
            'escuela': escuela,
            'actividades_escuela': actividades_escuela,
            'actividades_con_entrega': actividades_con_entrega,
            'recursos_con_archivo': recursos_con_archivo,
            'archivos_recurso_escuela': archivos_recurso_escuela,
        }
        return render(request, 'hechos/detalle_curso_estudiante.html', context)

    if not escuela_curso:
        messages.error(request, 'Este curso no está vinculado a una escuela operativa.')
        return redirect('hechos:cursos_list')

    if hasattr(request.user, 'profesor_profile'):
        if escuela_curso.maestro_id != request.user.profesor_profile.id:
            messages.error(request, 'No tienes permisos para acceder a este curso.')
            return redirect('hechos:cursos_list')

    if hasattr(request.user, 'profesor_profile'):
        matriculas = Matricula.objects.filter(
            escuela=escuela_curso,
            escuela__maestro=request.user.profesor_profile,
            sede=user_sede,
            is_active=True,
        ).select_related('estudiante__user')
        clases = Clase.objects.filter(
            escuela=escuela_curso,
            escuela__maestro=request.user.profesor_profile,
            sede=user_sede,
            is_active=True,
        ).order_by('fecha_clase')
    else:
        matriculas = Matricula.objects.filter(
            escuela=escuela_curso, sede=user_sede, is_active=True
        ).select_related('estudiante__user')
        clases = Clase.objects.filter(
            escuela=escuela_curso, sede=user_sede, is_active=True
        ).order_by('fecha_clase')

    total_estudiantes = matriculas.count()
    total_clases = clases.count()

    notas_qs = Nota.objects.filter(
        escuela=escuela_curso,
        sede=user_sede,
    ).select_related('estudiante__user').order_by('-fecha_evaluacion')

    promedio_general = notas_qs.aggregate(
        promedio=Avg('puntaje_obtenido')
    )['promedio'] or 0

    estudiantes_con_notas = []
    for matricula in matriculas:
        notas_estudiante = notas_qs.filter(estudiante=matricula.estudiante)
        
        # Calcular promedio del estudiante
        if notas_estudiante.exists():
            promedio_estudiante = notas_estudiante.aggregate(
                promedio=Avg('puntaje_obtenido')
            )['promedio'] or 0
        else:
            promedio_estudiante = 0
        
        estudiantes_con_notas.append({
            'matricula': matricula,
            'promedio': round(promedio_estudiante, 2),
            'total_notas': notas_estudiante.count(),
            'notas': notas_estudiante[:3]  # Últimas 3 notas
        })
    
    context = {
        'curso': curso,
        'matriculas': matriculas,
        'clases': clases,
        'estudiantes_con_notas': estudiantes_con_notas,
        'total_estudiantes': total_estudiantes,
        'total_clases': total_clases,
        'promedio_general': round(promedio_general, 2),
        'user_sede': user_sede,
    }

    return render(request, 'hechos/detalle_curso.html', context)


@login_required
def tomar_asistencia_curso(request, curso_id):
    """
    Tomar asistencia semanal para un curso específico
    """
    if not hasattr(request.user, 'profesor_profile'):
        messages.error(request, 'Solo los profesores pueden tomar asistencia.')
        return redirect('core:dashboard')
    
    # Obtener la sede del usuario
    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada. Contacta al administrador.')
        return redirect('core:dashboard')
    
    profesor = request.user.profesor_profile
    curso = get_object_or_404(Curso, id=curso_id, sede=user_sede, is_active=True)
    escuela_op = getattr(curso.ruta_estudio, 'escuela', None)
    if not escuela_op or escuela_op.maestro_id != profesor.id:
        messages.error(request, 'No tienes permisos para acceder a este curso.')
        return redirect('hechos:cursos_list')

    estudiantes = Estudiante.objects.filter(
        matriculas__escuela=escuela_op,
        matriculas__sede=user_sede,
        matriculas__is_active=True,
        sede=user_sede,
        is_active=True,
    ).distinct().select_related('user')

    fecha_actual = timezone.now().date()

    clase_semanal = (
        Clase.objects.filter(
            escuela=escuela_op,
            fecha_clase__date=fecha_actual,
            sede=user_sede,
            is_active=True,
        )
        .order_by('id')
        .first()
    )
    if not clase_semanal:
        mx = (
            Clase.objects.filter(escuela=escuela_op, sede=user_sede).aggregate(
                m=Max('numero_clase')
            )['m']
            or 0
        )
        clase_semanal = Clase.objects.create(
            escuela=escuela_op,
            sede=user_sede,
            numero_clase=mx + 1,
            titulo=f'Clase Semanal - {fecha_actual.strftime("%d/%m/%Y")}',
            descripcion='Asistencia semanal del curso',
            fecha_clase=timezone.now(),
            profesor=profesor,
            is_active=True,
        )

    # Obtener asistencias existentes para esta clase
    asistencias_existentes = Asistencia.objects.filter(
        clase=clase_semanal,
        sede=user_sede
    ).select_related('estudiante__user')
    
    asistencia_dict = {a.estudiante.id: a.estado for a in asistencias_existentes}
    
    if request.method == 'POST':
        # Procesar asistencia
        for key, value in request.POST.items():
            if key.startswith('estudiante_'):
                estudiante_id = key.split('_')[1]
                try:
                    estudiante = Estudiante.objects.get(id=estudiante_id, sede=user_sede)
                    asistencia, created = Asistencia.objects.get_or_create(
                        clase=clase_semanal,
                        estudiante=estudiante,
                        sede=user_sede,
                        defaults={
                            'estado': value,
                            'registrado_por': request.user
                        }
                    )
                    if not created:
                        asistencia.estado = value
                        asistencia.registrado_por = request.user
                        asistencia.save()
                except Estudiante.DoesNotExist:
                    continue
        
        messages.success(request, f'Asistencia registrada correctamente para el {fecha_actual.strftime("%d/%m/%Y")}.')
        return redirect('hechos:tomar_asistencia_curso', curso_id=curso_id)
    
    # Obtener historial de asistencias (últimas 5 fechas)
    historial_asistencias = Clase.objects.filter(
        escuela=escuela_op,
        sede=user_sede,
        is_active=True,
    ).values('fecha_clase__date').distinct().order_by('-fecha_clase__date')[:5]
    
    # Preparar datos de estudiantes con asistencia
    estudiantes_con_asistencia = []
    for estudiante in estudiantes:
        estado_asistencia = asistencia_dict.get(estudiante.id, 'presente')  # Default a presente
        estudiantes_con_asistencia.append({
            'estudiante': estudiante,
            'estado_asistencia': estado_asistencia
        })
    
    context = {
        'curso': curso,
        'estudiantes': estudiantes,
        'estudiantes_con_asistencia': estudiantes_con_asistencia,
        'asistencia_dict': asistencia_dict,
        'fecha_actual': fecha_actual,
        'historial_asistencias': historial_asistencias,
        'user_sede': user_sede,
    }
    
    return render(request, 'hechos/tomar_asistencia_curso.html', context)




@login_required
def matriculas_list(request):
    """
    Lista de matrículas para administradores de sede con filtros y agrupación
    """
    if not request.user.is_super_admin():
        r = ca.require_academico(request)
        if r:
            return r
    
    # Obtener la sede del usuario
    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada. Contacta al administrador.')
        return redirect('core:dashboard')
    
    # Obtener parámetros de filtro
    filtro_curso = request.GET.get('curso', '')
    filtro_estado = request.GET.get('estado', '')
    filtro_ruta = request.GET.get('ruta', '')
    filtro_periodo = request.GET.get('periodo', '')
    agrupar_por = request.GET.get('agrupar', 'fecha')  # fecha, curso, ruta, estudiante, periodo
    
    # Obtener todas las matrículas de la sede
    matriculas = Matricula.objects.filter(
        sede=user_sede
    ).select_related(
        'estudiante__user', 
        'escuela', 
        'escuela__maestro__user'
    ).order_by('-periodo', '-fecha_matricula')
    
    # Aplicar filtros
    if filtro_curso:
        matriculas = matriculas.filter(escuela__id=filtro_curso)
    
    if filtro_estado:
        if filtro_estado == 'activa':
            matriculas = matriculas.filter(is_active=True)
        elif filtro_estado == 'inactiva':
            matriculas = matriculas.filter(is_active=False)
    
    if filtro_ruta:
        matriculas = matriculas.filter(escuela__id=filtro_ruta)
    
    if filtro_periodo:
        matriculas = matriculas.filter(periodo=filtro_periodo)
    
    # Ordenar según agrupación
    if agrupar_por == 'curso':
        matriculas = matriculas.order_by('escuela__nombre', '-fecha_matricula')
    elif agrupar_por == 'ruta':
        matriculas = matriculas.order_by('escuela__nombre', '-fecha_matricula')
    elif agrupar_por == 'estudiante':
        matriculas = matriculas.order_by('estudiante__user__first_name', 'estudiante__user__last_name', '-fecha_matricula')
    elif agrupar_por == 'periodo':
        matriculas = matriculas.order_by('-periodo', 'estudiante__user__first_name')
    else:  # fecha
        matriculas = matriculas.order_by('-fecha_matricula')
    
    # Estadísticas
    matriculas_activas = matriculas.filter(is_active=True)
    hoy = timezone.localdate()
    matriculas_en_progreso = matriculas_activas.filter(
        Q(escuela__fecha_inicio__isnull=True) | Q(escuela__fecha_inicio__lte=hoy),
        Q(escuela__fecha_fin__isnull=True) | Q(escuela__fecha_fin__gte=hoy),
    )
    esc_ids = matriculas_activas.values_list('escuela_id', flat=True).distinct()
    cursos_unicos = Curso.objects.filter(
        ruta_estudio__escuela_id__in=esc_ids,
        sede=user_sede,
    ).distinct()
    
    # Obtener opciones para filtros
    cursos_disponibles = Curso.objects.filter(sede=user_sede, is_active=True).order_by('nombre')
    rutas_disponibles = RutaEstudio.objects.filter(sede=user_sede, is_active=True).order_by('nombre')
    periodos_disponibles = Matricula.objects.filter(sede=user_sede).values_list('periodo', flat=True).distinct().order_by('-periodo')
    
    # Agrupar matrículas según el parámetro
    matriculas_agrupadas = {}
    if agrupar_por == 'curso':
        for matricula in matriculas:
            c = matricula.curso
            curso_nombre = (c.nombre if c else matricula.escuela.nombre) or 'Sin curso'
            if curso_nombre not in matriculas_agrupadas:
                matriculas_agrupadas[curso_nombre] = []
            matriculas_agrupadas[curso_nombre].append(matricula)
    elif agrupar_por == 'ruta':
        for matricula in matriculas:
            c = matricula.curso
            ruta = getattr(c, 'ruta_estudio', None) if c else None
            ruta_nombre = ruta.nombre if ruta else 'Sin ruta'
            if ruta_nombre not in matriculas_agrupadas:
                matriculas_agrupadas[ruta_nombre] = []
            matriculas_agrupadas[ruta_nombre].append(matricula)
    elif agrupar_por == 'estudiante':
        for matricula in matriculas:
            estudiante_nombre = matricula.estudiante.user.get_full_name()
            if estudiante_nombre not in matriculas_agrupadas:
                matriculas_agrupadas[estudiante_nombre] = []
            matriculas_agrupadas[estudiante_nombre].append(matricula)
    elif agrupar_por == 'periodo':
        for matricula in matriculas:
            periodo_display = get_period_display(matricula.periodo)
            if periodo_display not in matriculas_agrupadas:
                matriculas_agrupadas[periodo_display] = []
            matriculas_agrupadas[periodo_display].append(matricula)
    else:  # fecha
        for matricula in matriculas:
            fecha_str = matricula.fecha_matricula.strftime('%Y-%m-%d')
            if fecha_str not in matriculas_agrupadas:
                matriculas_agrupadas[fecha_str] = []
            matriculas_agrupadas[fecha_str].append(matricula)
    
    context = {
        'matriculas': matriculas,
        'matriculas_agrupadas': matriculas_agrupadas,
        'matriculas_activas': matriculas_activas,
        'matriculas_en_progreso': matriculas_en_progreso,
        'cursos_unicos': cursos_unicos,
        'cursos_disponibles': cursos_disponibles,
        'rutas_disponibles': rutas_disponibles,
        'periodos_disponibles': periodos_disponibles,
        'user_sede': user_sede,
        'filtro_curso': filtro_curso,
        'filtro_estado': filtro_estado,
        'filtro_ruta': filtro_ruta,
        'filtro_periodo': filtro_periodo,
        'agrupar_por': agrupar_por,
    }
    
    return render(request, 'hechos/matriculas_list.html', context)


@login_required
def desmatricular_estudiante(request, matricula_id):
    """
    Desmatricular estudiante de un curso
    """
    if not request.user.is_super_admin():
        r = ca.require_academico(request)
        if r:
            return r
    
    # Obtener la sede del usuario
    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada. Contacta al administrador.')
        return redirect('core:dashboard')
    
    try:
        matricula = Matricula.objects.get(id=matricula_id, sede=user_sede)
    except Matricula.DoesNotExist:
        messages.error(request, 'Matrícula no encontrada.')
        return redirect('hechos:matriculas_list')
    
    if request.method == 'POST':
        try:
            matricula.is_active = False
            matricula.save()
            messages.success(request, f'Estudiante {matricula.estudiante.user.get_full_name()} desmatriculado exitosamente del curso {matricula.curso.nombre}.')
            return redirect('hechos:matriculas_list')
        except Exception as e:
            messages.error(request, f'Error al desmatricular estudiante: {str(e)}')
    
    context = {
        'matricula': matricula,
        'user_sede': user_sede,
    }
    
    return render(request, 'hechos/desmatricular_estudiante.html', context)


# ===== VISTAS PARA GESTIONAR EDICIONES DE CURSOS =====

@login_required
def ediciones_curso_list(request, curso_id):
    """Compatibilidad: ya no hay listado de ediciones; la oferta vive en la escuela."""
    if not request.user.is_super_admin():
        r = ca.require_academico_o_pedagogico(request)
        if r:
            return r

    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada. Contacta al administrador.')
        return redirect('core:dashboard')

    curso = get_object_or_404(Curso, id=curso_id, sede=user_sede, is_active=True)
    esc = getattr(curso.ruta_estudio, 'escuela', None)
    if esc:
        messages.info(
            request,
            'La oferta (fechas, cupo, horario, docente) se configura en la escuela. '
            'Usa el asistente de estructura.',
        )
        return redirect('hechos:estructura_configurar', escuela_id=esc.id, step=1)
    messages.warning(request, 'Este curso no tiene una escuela operativa vinculada.')
    return redirect('hechos:cursos_list')


@login_required
def crear_edicion_curso(request, curso_id):
    """Compatibilidad: redirige a configuración por escuela."""
    return ediciones_curso_list(request, curso_id)


@login_required
def editar_edicion_curso(request, edicion_id):
    """Compatibilidad: rutas antiguas por id de edición."""
    if not request.user.is_super_admin():
        r = ca.require_academico_o_pedagogico(request)
        if r:
            return r
    messages.info(
        request,
        'Las rutas por «edición» ya no se usan. Abre la escuela desde Estructura y edita la oferta allí.',
    )
    return redirect('hechos:estructura_ediciones')


def _coordinador_financiero_gate(request):
    if request.user.is_super_admin():
        return None
    if ca.has_capacidad(request.user, 'financiero'):
        return None
    messages.error(request, 'Necesitas la función de recursos y finanzas para acceder a esta sección.')
    return redirect('core:dashboard')


@login_required
def coordinador_recursos(request):
    gate = _coordinador_financiero_gate(request)
    if gate:
        return gate
    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada.')
        return redirect('core:dashboard')

    ofrendas_qs = (
        Ofrenda.objects.filter(sede=user_sede)
        .select_related('escuela', 'registrado_por')
        .order_by('-fecha', '-id')[:1500]
    )
    recaudos_qs = (
        RecaudoOcasional.objects.filter(sede=user_sede)
        .select_related('registrado_por')
        .order_by('-fecha', '-id')[:1500]
    )
    agg_of = Ofrenda.objects.filter(sede=user_sede).aggregate(
        total=Sum('valor'),
        n=Count('id'),
    )
    agg_rec = RecaudoOcasional.objects.filter(sede=user_sede).aggregate(
        total=Sum('monto'),
        n=Count('id'),
    )
    total_of = agg_of['total'] or Decimal('0')
    total_rec = agg_rec['total'] or Decimal('0')
    n_of = agg_of['n'] or 0
    n_rec = agg_rec['n'] or 0

    movimientos = []
    for o in ofrendas_qs:
        movimientos.append(
            {
                'kind': 'ofrenda',
                'fecha': o.fecha,
                'pk': o.id,
                'origen_label': 'Ofrenda',
                'detalle': o.escuela.nombre,
                'valor': o.valor,
                'registrado_por': o.registrado_por,
            }
        )
    for r in recaudos_qs:
        movimientos.append(
            {
                'kind': 'recaudo',
                'fecha': r.fecha,
                'pk': r.id,
                'origen_label': 'Recaudo ocasional',
                'detalle': r.concepto,
                'valor': r.monto,
                'registrado_por': r.registrado_por,
            }
        )
    movimientos.sort(key=lambda m: (m['fecha'], m['pk']), reverse=True)
    movimientos = movimientos[:500]

    return render(
        request,
        'hechos/coordinador_recursos.html',
        {
            'user_sede': user_sede,
            'movimientos': movimientos,
            'total_ofrendado': total_of,
            'total_recaudos_ocasionales': total_rec,
            'total_ingresos_sede': total_of + total_rec,
            'num_ofrendas': n_of,
            'num_recaudos': n_rec,
            'num_movimientos': n_of + n_rec,
        },
    )


@login_required
def coordinador_eventos_presupuesto(request):
    """
    Eventos o partidas y monto que la sede destina (planificación), aparte del registro de ofrendas.
    """
    gate = _coordinador_financiero_gate(request)
    if gate:
        return gate
    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada.')
        return redirect('core:dashboard')

    if request.method == 'POST':
        action = (request.POST.get('action') or '').strip()
        if action == 'eliminar_evento':
            try:
                eid = int(request.POST.get('evento_id'))
            except (TypeError, ValueError):
                messages.error(request, 'Solicitud inválida.')
                return redirect('hechos:coordinador_eventos_presupuesto')
            evento = get_object_or_404(PresupuestoEvento, id=eid, sede=user_sede)
            evento.delete()
            messages.success(request, 'Registro de evento eliminado.')
            return redirect('hechos:coordinador_eventos_presupuesto')

        form = PresupuestoEventoForm(request.POST, sede=user_sede)
        if form.is_valid():
            obj = form.save(commit=False)
            obj.sede = user_sede
            obj.creado_por = request.user
            obj.save()
            messages.success(
                request,
                'Listo: quedó registrado el evento y el monto que planeas destinar.',
            )
            return redirect('hechos:coordinador_eventos_presupuesto')
        messages.error(request, 'Revisa los datos del formulario.')
    else:
        form = PresupuestoEventoForm(sede=user_sede)

    eventos = PresupuestoEvento.objects.filter(sede=user_sede).order_by(
        '-fecha_evento', '-created_at', '-id'
    )
    agg = PresupuestoEvento.objects.filter(sede=user_sede).aggregate(total=Sum('monto_destinado'))
    td = agg['total'] or Decimal('0')
    ti = ingresos_totales_sede(user_sede)
    return render(
        request,
        'hechos/coordinador_eventos_presupuesto.html',
        {
            'user_sede': user_sede,
            'eventos': eventos,
            'form': form,
            'total_destinado_eventos': td,
            'total_ingresos_sede': ti,
            'disponible_para_eventos': max(ti - td, Decimal('0')),
        },
    )


@login_required
def coordinador_recaudos_ocasionales(request):
    """
    Ingresos puntuales que registra el financiero: extras, sobrantes de evento, donaciones ocasionales, etc.
    """
    gate = _coordinador_financiero_gate(request)
    if gate:
        return gate
    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada.')
        return redirect('core:dashboard')

    if request.method == 'POST':
        action = (request.POST.get('action') or '').strip()
        if action == 'eliminar_recaudo':
            try:
                rid = int(request.POST.get('recaudo_id'))
            except (TypeError, ValueError):
                messages.error(request, 'Solicitud inválida.')
                return redirect('hechos:coordinador_recaudos_ocasionales')
            rec = get_object_or_404(RecaudoOcasional, id=rid, sede=user_sede)
            rec.delete()
            messages.success(request, 'Recaudo ocasional eliminado.')
            return redirect('hechos:coordinador_recaudos_ocasionales')

        form = RecaudoOcasionalForm(request.POST)
        if form.is_valid():
            obj = form.save(commit=False)
            obj.sede = user_sede
            obj.registrado_por = request.user
            obj.save()
            messages.success(request, 'Recaudo ocasional registrado correctamente.')
            return redirect('hechos:coordinador_recaudos_ocasionales')
        messages.error(request, 'Revisa los datos del formulario.')
    else:
        form = RecaudoOcasionalForm(initial={'fecha': timezone.localdate()})

    recaudos = RecaudoOcasional.objects.filter(sede=user_sede).select_related('registrado_por').order_by(
        '-fecha', '-created_at', '-id'
    )
    agg = RecaudoOcasional.objects.filter(sede=user_sede).aggregate(total=Sum('monto'))
    return render(
        request,
        'hechos/coordinador_recaudos_ocasionales.html',
        {
            'user_sede': user_sede,
            'recaudos': recaudos,
            'form': form,
            'total_recaudado_ocasional': agg['total'] or 0,
        },
    )


@login_required
def coordinador_logistica(request):
    if not request.user.is_super_admin():
        p = getattr(request.user, 'admin_escuela_profile', None)
        if not p or not p.is_active or not ca.has_capacidad(request.user, 'logistico'):
            messages.error(request, 'Necesitas la función logística para acceder a esta sección.')
            return redirect('core:dashboard')
    user_sede = get_user_sede(request.user)
    if not user_sede:
        messages.warning(request, 'No tienes una sede asignada.')
        return redirect('core:dashboard')

    raw_tab = (request.GET.get('tab') or '').strip().lower()
    if raw_tab in ('salones', 'escuelas'):
        logistica_tab = raw_tab
    elif request.GET.get('escuela'):
        logistica_tab = 'escuelas'
    else:
        logistica_tab = 'salones'

    salon_form = SalonForm(prefix='nuevo')
    salon_editando = None
    salon_edit_form = None

    if request.method == 'POST':
        action = (request.POST.get('action') or '').strip()

        if action == 'crear_salon':
            salon_form = SalonForm(request.POST, prefix='nuevo')
            if salon_form.is_valid():
                salon = salon_form.save(commit=False)
                salon.sede = user_sede
                try:
                    salon.save()
                    messages.success(
                        request,
                        f'Salón «{salon.nombre}» creado con {salon.capacidad_plazas} plazas.',
                    )
                except IntegrityError:
                    messages.error(request, 'Ya existe un salón con ese nombre en esta sede.')
                return _redirect_coordinador_logistica(tab='salones')
            messages.error(request, 'Revisa los datos del salón.')

        elif action == 'actualizar_salon':
            salon = get_object_or_404(
                Salon,
                id=request.POST.get('salon_id'),
                sede=user_sede,
            )
            salon_edit_form = SalonForm(request.POST, instance=salon, prefix='edit')
            if salon_edit_form.is_valid():
                try:
                    salon_edit_form.save()
                    s = salon_edit_form.instance
                    if s.escuela_id:
                        Escuela.objects.filter(
                            pk=s.escuela_id,
                            sede=user_sede,
                            is_active=True,
                        ).update(cupo_maximo=s.capacidad_plazas)
                    messages.success(
                        request,
                        f'Salón «{salon_edit_form.instance.nombre}» actualizado.',
                    )
                    return _redirect_coordinador_logistica(tab='salones')
                except IntegrityError:
                    messages.error(request, 'Ya existe un salón con ese nombre en esta sede.')
            else:
                messages.error(request, 'Revisa los datos del salón.')
            salon_editando = salon

        elif action == 'toggle_salon':
            salon = get_object_or_404(
                Salon,
                id=request.POST.get('salon_id'),
                sede=user_sede,
            )
            salon.is_active = not salon.is_active
            salon.save(update_fields=['is_active', 'updated_at'])
            messages.success(request, 'Estado del salón actualizado.')
            return _redirect_coordinador_logistica(tab='salones')

        elif action == 'asignar_salon_escuela':
            escuela = get_object_or_404(
                Escuela,
                id=request.POST.get('escuela_id'),
                sede=user_sede,
                is_active=True,
            )
            raw_sid = (request.POST.get('salon_id') or '').strip()
            if not raw_sid:
                messages.error(request, 'Elige un salón en el menú desplegable.')
                return _redirect_coordinador_logistica(tab='escuelas', escuela_id=escuela.id)
            salon = get_object_or_404(Salon, id=raw_sid, sede=user_sede)
            if salon.escuela_id == escuela.id:
                messages.info(request, 'Ese salón ya está asignado a esta escuela.')
                return _redirect_coordinador_logistica(tab='escuelas', escuela_id=escuela.id)
            if not salon.is_active:
                messages.error(request, 'Activa el salón antes de asignarlo.')
            elif escuela.sede_id != user_sede.id:
                messages.error(request, 'La escuela no pertenece a tu sede.')
            else:
                otros_en_escuela = list(
                    Salon.objects.filter(escuela=escuela).exclude(pk=salon.pk),
                )
                if otros_en_escuela:
                    Salon.objects.filter(escuela=escuela).exclude(pk=salon.pk).update(
                        escuela=None,
                        updated_at=timezone.now(),
                    )
                prev = salon.escuela
                salon.escuela = escuela
                salon.save(update_fields=['escuela', 'updated_at'])
                Escuela.objects.filter(pk=escuela.pk, sede=user_sede, is_active=True).update(
                    cupo_maximo=salon.capacidad_plazas
                )

                frases = []
                if otros_en_escuela:
                    nombres = ', '.join(f'«{o.nombre}»' for o in otros_en_escuela)
                    frases.append(
                        f'Cada escuela solo puede tener un salón; se quitó {nombres} de esta escuela.',
                    )
                if prev and prev.id != escuela.id:
                    frases.append(
                        f'«{salon.nombre}» pasó de «{prev.nombre}» a «{escuela.nombre}».',
                    )
                elif not prev:
                    frases.append(f'«{salon.nombre}» quedó asignado a «{escuela.nombre}».')
                if frases:
                    messages.success(request, ' '.join(frases))
            return _redirect_coordinador_logistica(tab='escuelas', escuela_id=escuela.id)

        elif action == 'quitar_salon_escuela':
            escuela = get_object_or_404(
                Escuela,
                id=request.POST.get('escuela_id'),
                sede=user_sede,
                is_active=True,
            )
            salon = get_object_or_404(
                Salon,
                id=request.POST.get('salon_id'),
                sede=user_sede,
                escuela=escuela,
            )
            salon.escuela = None
            salon.save(update_fields=['escuela', 'updated_at'])
            Escuela.objects.filter(pk=escuela.pk, sede=user_sede, is_active=True).update(
                cupo_maximo=0
            )
            messages.success(request, f'«{salon.nombre}» ya no está asignado a «{escuela.nombre}».')
            return _redirect_coordinador_logistica(tab='escuelas', escuela_id=escuela.id)

    salones = (
        Salon.objects.filter(sede=user_sede)
        .select_related('escuela')
        .order_by('-is_active', 'nombre')
    )
    escuelas = Escuela.objects.filter(sede=user_sede, is_active=True).order_by('-anio', 'nombre')

    escuela_focus = None
    eid = request.GET.get('escuela')
    if eid:
        try:
            escuela_focus = Escuela.objects.get(
                id=int(eid),
                sede=user_sede,
                is_active=True,
            )
        except (ValueError, Escuela.DoesNotExist):
            escuela_focus = None

    if logistica_tab == 'salones' and salon_editando is None:
        edit_get = request.GET.get('editar')
        if edit_get:
            try:
                salon_editando = Salon.objects.get(id=int(edit_get), sede=user_sede)
                salon_edit_form = SalonForm(instance=salon_editando, prefix='edit')
            except (ValueError, Salon.DoesNotExist):
                messages.warning(request, 'No se encontró ese salón en tu sede.')

    salon_asignado_escuela = None
    salones_elegibles_asignar = []
    if escuela_focus:
        salon_asignado_escuela = (
            Salon.objects.filter(sede=user_sede, escuela_id=escuela_focus.id)
            .select_related('escuela')
            .order_by('-is_active', 'nombre')
            .first()
        )
        salones_elegibles_asignar = list(
            Salon.objects.filter(sede=user_sede, is_active=True)
            .exclude(escuela_id=escuela_focus.id)
            .select_related('escuela')
            .order_by('nombre'),
        )

    return render(
        request,
        'hechos/coordinador_logistica.html',
        {
            'user_sede': user_sede,
            'salon_form': salon_form,
            'salon_editando': salon_editando,
            'salon_edit_form': salon_edit_form,
            'salones': salones,
            'escuelas': escuelas,
            'escuela_focus': escuela_focus,
            'logistica_tab': logistica_tab,
            'salon_asignado_escuela': salon_asignado_escuela,
            'salones_elegibles_asignar': salones_elegibles_asignar,
        },
    )
